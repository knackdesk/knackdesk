"""Build the Online Seller Profit Tracker workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
PROD_ROWS = (5, 54); ORDER_ROWS = (5, 504); MONTHS = 12


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Online Seller Profit Tracker", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: enter your platform fee, payment fee and fixed fee per order, and a default packaging cost. These apply to every order.",
        "2. Products: one row per SKU with its cost and selling price. Orders look these up by SKU.",
        "3. Orders: one row per order: date, SKU, quantity, shipping charged to the customer and shipping you paid. Price, cost, fees, revenue, profit and margin are calculated. Overwrite the price or cost on a row if that order differed.",
        "4. Summary: revenue, fees, shipping, profit and margin by month (set the start month), and profit by product.",
        "5. Ads: monthly ad spend against revenue gives ROAS, and the break-even ROAS from the month's real margin shows whether the spend paid off.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "Blue cells that already contain a formula (price and cost on order rows) are defaults you may overwrite.",
        "", "FORMULAS",
        "Revenue = price × quantity + shipping charged. Fees = revenue × (platform % + payment %) + fixed fee. Profit = revenue − cost × quantity − packaging − shipping paid − fees.",
        "Break-even ROAS = revenue ÷ profit before ads for the month (1 ÷ contribution margin).",
        "Arithmetic only; not tax or financial advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 60
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [("Platform or marketplace fee (% of revenue)", 10, "0.00", "Charged on price plus shipping on most marketplaces."),
            ("Payment processing fee (%)", 3, "0.00", None),
            ("Payment processing fixed fee per order", 0.30, MONEY, None),
            ("Default packaging cost per order", 1, MONEY, "Box, filler, label, inserts."),
            ("Currency label", "USD", None, "Text only.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True, note=note)
        if note: label(ws, f"C{i}", note, muted=True)


def build_products(ws):
    lo, hi = PROD_ROWS
    label(ws, "A1", "Products", bold=True).font = TITLE
    head(ws, 4, ["SKU", "Name", "Cost per unit", "Price per unit", "Unit margin", "Unit margin %"], [12, 30, 14, 14, 14, 14])
    for r in range(lo, hi + 1):
        for col in "ABCD": ws[f"{col}{r}"].font = BLUE
        ws[f"C{r}"].number_format = MONEY; ws[f"D{r}"].number_format = MONEY
        fx(ws, f"E{r}", f'=IF(A{r}="","",D{r}-C{r})', MONEY)
        fx(ws, f"F{r}", f'=IF(OR(A{r}="",D{r}=0),"",ROUND((D{r}-C{r})/D{r}*100,2))', PCT)
    inp(ws, f"A{lo}", "SKU-001", key=True); inp(ws, f"B{lo}", "Example candle, large"); inp(ws, f"C{lo}", 15, MONEY); inp(ws, f"D{lo}", 40, MONEY)
    inp(ws, f"A{lo+1}", "SKU-002"); inp(ws, f"B{lo+1}", "Example candle, small"); inp(ws, f"C{lo+1}", 8, MONEY); inp(ws, f"D{lo+1}", 22, MONEY)
    ws.freeze_panes = "A5"


def build_orders(ws):
    lo, hi = ORDER_ROWS; plo, phi = PROD_ROWS
    label(ws, "A1", "Orders", bold=True).font = TITLE
    label(ws, "A2", "Total profit", bold=True); fx(ws, "B2", f"=SUM(J{lo}:J{hi})", MONEY, bold=True)
    label(ws, "C2", "Orders", bold=True); fx(ws, "D2", f'=COUNTA(A{lo}:A{hi})', "0", bold=True)
    head(ws, 4, ["Date", "SKU", "Quantity", "Price per unit", "Cost per unit", "Shipping charged", "Shipping paid", "Packaging", "Fees", "Profit", "Revenue", "Margin %", "Month"],
         [12, 11, 9, 12, 12, 12, 12, 11, 11, 12, 12, 10, 10])
    dv = DataValidation(type="list", formula1=f"=Products!$A${plo}:$A${phi}", allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        for col in "ABCFG": ws[f"{col}{r}"].font = BLUE
        ws[f"A{r}"].number_format = DATE; dv.add(f"B{r}")
        for col in "DEFGHIJK": ws[f"{col}{r}"].number_format = MONEY
        inp(ws, f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX(Products!$D${plo}:$D${phi},MATCH(B{r},Products!$A${plo}:$A${phi},0)),0))', MONEY)
        inp(ws, f"E{r}", f'=IF(B{r}="","",IFERROR(INDEX(Products!$C${plo}:$C${phi},MATCH(B{r},Products!$A${plo}:$A${phi},0)),0))', MONEY)
        inp(ws, f"H{r}", f'=IF(A{r}="","",Settings!$B$5)', MONEY)
        fx(ws, f"K{r}", f'=IF(A{r}="","",D{r}*C{r}+F{r})', MONEY)
        fx(ws, f"I{r}", f'=IF(A{r}="","",ROUND(K{r}*(Settings!$B$2+Settings!$B$3)/100+Settings!$B$4,2))', MONEY)
        fx(ws, f"J{r}", f'=IF(A{r}="","",ROUND(K{r}-E{r}*C{r}-H{r}-G{r}-I{r},2))', MONEY)
        fx(ws, f"L{r}", f'=IF(OR(A{r}="",K{r}=0),"",ROUND(J{r}/K{r}*100,2))', PCT)
        fx(ws, f"M{r}", f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
    today = dt.date.today()
    inp(ws, f"A{lo}", today - dt.timedelta(days=9), DATE, key=True); inp(ws, f"B{lo}", "SKU-001", key=True); inp(ws, f"C{lo}", 1, "0"); inp(ws, f"F{lo}", 5, MONEY); inp(ws, f"G{lo}", 6, MONEY)
    inp(ws, f"A{lo+1}", today - dt.timedelta(days=4), DATE); inp(ws, f"B{lo+1}", "SKU-002"); inp(ws, f"C{lo+1}", 2, "0"); inp(ws, f"F{lo+1}", 0, MONEY); inp(ws, f"G{lo+1}", 4.5, MONEY)
    ws.freeze_panes = "A5"


def build_summary(ws):
    lo, hi = ORDER_ROWS; plo, phi = PROD_ROWS
    ws.column_dimensions["A"].width = 12
    for col in "BCDEFG": ws.column_dimensions[col].width = 14
    label(ws, "A1", "Summary", bold=True).font = TITLE
    label(ws, "A2", "Start month"); inp(ws, "B2", (dt.date.today() - dt.timedelta(days=9)).replace(day=1), DATE, key=True, note="First month to show; twelve months follow.")
    head(ws, 4, ["Month", "Orders", "Revenue", "Fees", "Shipping paid", "Profit", "Margin %"])
    for i in range(MONTHS):
        r = 5 + i
        fx(ws, f"A{r}", f'=TEXT(EDATE($B$2,{i}),"yyyy-mm")')
        fx(ws, f"B{r}", f'=COUNTIF(Orders!$M${lo}:$M${hi},A{r})', "0")
        fx(ws, f"C{r}", f'=SUMIF(Orders!$M${lo}:$M${hi},A{r},Orders!$K${lo}:$K${hi})', MONEY)
        fx(ws, f"D{r}", f'=SUMIF(Orders!$M${lo}:$M${hi},A{r},Orders!$I${lo}:$I${hi})', MONEY)
        fx(ws, f"E{r}", f'=SUMIF(Orders!$M${lo}:$M${hi},A{r},Orders!$G${lo}:$G${hi})', MONEY)
        fx(ws, f"F{r}", f'=SUMIF(Orders!$M${lo}:$M${hi},A{r},Orders!$J${lo}:$J${hi})', MONEY)
        fx(ws, f"G{r}", f'=IF(C{r}=0,"",ROUND(F{r}/C{r}*100,2))', PCT)
    label(ws, "A17", "12 months", bold=True)
    for col in "BCDEF": fx(ws, f"{col}17", f"=SUM({col}5:{col}16)", MONEY if col != "B" else "0", bold=True)
    fx(ws, "G17", '=IF(C17=0,"",ROUND(F17/C17*100,2))', PCT, bold=True)
    head(ws, 20, ["SKU", "Name", "Units sold", "Revenue", "Profit", "Margin %"])
    for i, pr in enumerate(range(plo, plo + 20)):
        r = 21 + i
        fx(ws, f"A{r}", f'=IF(Products!A{pr}="","",Products!A{pr})')
        fx(ws, f"B{r}", f'=IF(Products!A{pr}="","",Products!B{pr})')
        fx(ws, f"C{r}", f'=IF(A{r}="","",SUMIF(Orders!$B${lo}:$B${hi},A{r},Orders!$C${lo}:$C${hi}))', "0")
        fx(ws, f"D{r}", f'=IF(A{r}="","",SUMIF(Orders!$B${lo}:$B${hi},A{r},Orders!$K${lo}:$K${hi}))', MONEY)
        fx(ws, f"E{r}", f'=IF(A{r}="","",SUMIF(Orders!$B${lo}:$B${hi},A{r},Orders!$J${lo}:$J${hi}))', MONEY)
        fx(ws, f"F{r}", f'=IF(OR(A{r}="",D{r}=0),"",ROUND(E{r}/D{r}*100,2))', PCT)
    label(ws, "A42", "Profit by product covers the first 20 SKUs on the Products sheet.", muted=True)


def build_ads(ws):
    ws.column_dimensions["A"].width = 12
    for col in "BCDEFGH": ws.column_dimensions[col].width = 15
    label(ws, "A1", "Ads", bold=True).font = TITLE
    label(ws, "A2", "Enter monthly ad spend. Revenue and profit come from the Summary sheet; break-even ROAS is revenue ÷ profit before ads.", muted=True)
    head(ws, 4, ["Month", "Ad spend", "Revenue", "Profit before ads", "ROAS", "Break-even ROAS", "Profit after ads", "Verdict"])
    for i in range(MONTHS):
        r = 5 + i
        fx(ws, f"A{r}", f"=Summary!A{r}")
        inp(ws, f"B{r}", 0, MONEY)
        fx(ws, f"C{r}", f"=Summary!C{r}", MONEY)
        fx(ws, f"D{r}", f"=Summary!F{r}", MONEY)
        fx(ws, f"E{r}", f'=IF(B{r}=0,"",ROUND(C{r}/B{r},2))', "0.00")
        fx(ws, f"F{r}", f'=IF(OR(C{r}=0,D{r}<=0),"",ROUND(C{r}/D{r},2))', "0.00")
        fx(ws, f"G{r}", f"=D{r}-B{r}", MONEY)
        fx(ws, f"H{r}", f'=IF(B{r}=0,"",IF(G{r}>=0,"Profitable","Loss after ads"))')
    inp(ws, "B5", 150, MONEY, key=True)


README = """# Online Seller Profit Tracker

Thank you for buying the tracker. Open `online-seller-profit-tracker.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** your platform fee, payment fee, fixed fee and default packaging cost.
2. **Products:** one row per SKU with cost and price.
3. **Orders:** one row per order; price and cost are looked up from the SKU, fees and profit are calculated.
4. **Summary** and **Ads** update from the orders you log.

Blue text = type here. Black text = formula, leave it alone. Formulas are explained on the Start Here sheet.

No macros, no tracking. Arithmetic only, not tax or financial advice. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_products(wb.create_sheet("Products")); build_orders(wb.create_sheet("Orders")); build_summary(wb.create_sheet("Summary")); build_ads(wb.create_sheet("Ads"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "online-seller-profit-tracker.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "online-seller-profit-tracker.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "online-seller-profit-tracker.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
