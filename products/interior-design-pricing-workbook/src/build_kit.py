"""Build the Interior Design Pricing Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "interior-design-pricing-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; INT0 = "0"; HRS = "0.0"
FEE_ROWS = (5, 24); PR_ROWS = (5, 34); ROOM_ROWS = (5, 16); RP_ROWS = (20, 25); BP_ROWS = (5, 10)
S = "Settings!$B$"  # settings cell prefix
FB = "'Fee Builder'!"; PR = "Procurement!"; RF = "'Room Fees'!"; BP = "'Budget Plan'!"
PROJECTS = f"{FB}$A${FEE_ROWS[0]}:$A${FEE_ROWS[1]}"
# Settings rows (column B): referenced by every other sheet.
R_INCOME = 2; R_OH_FIRST = 3; R_OH_LAST = 8; R_OVERHEAD = 9; R_HOURS = 10; R_CONT = 11; R_PCTFEE = 12
R_MARKUP = 13; R_TRADE = 14; R_VISIT = 15; R_MIN = 16; R_UPLIFT = 17; R_NEEDED = 18; R_OWN = 19; R_RATE = 20
PCT_MORE = "percentage fee"; MIN_FLAG = "minimum applies"; OVER = "over-allocated"


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def note(ws, ref, text):
    ws[ref].comment = Comment(text, "Knackdesk")


def project_list(ws, cells):
    """Project dropdown fed by the project names typed on Fee Builder."""
    dv = DataValidation(type="list", formula1=PROJECTS, allow_blank=True); ws.add_data_validation(dv)
    for ref in cells: dv.add(ref)


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Interior Design Pricing Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: your annual income goal, six annual overhead lines, billable hours per year, default contingency %, percentage fee %, procurement markup %, trade discount %, hours per site visit, minimum project fee and complexity uplift %. The hourly rate you need follows; type your own rate to override it, and every other sheet prices at the rate used.",
        "2. Fee Builder: one row per project or quote with client, concept, development and documentation hours, site visits, hours per visit, contingency % and percentage fee % (defaults from Settings, type your own to override) and the project budget. Total hours, base fee, contingency, design fee, fee as % of budget, the percentage fee, the difference, which pays more, the percentage fee's implied hourly rate and a minimum fee flag are calculated.",
        "3. Procurement: one row per item you specify and sell, with project, supplier, retail price, trade discount % and markup % (defaults from Settings, type your own to override), an optional client price, freight and handling and quantity. Your cost, the client price used, line cost, revenue, gross profit, margin % and the client's saving against retail are calculated, with totals.",
        "4. Room Fees: one row per room with the hours it takes, complexity uplift % and rate (defaults from Settings, type your own to override); hours with uplift and the room fee follow. The block below totals each project's room fees, adds site visit hours, applies your minimum project fee and shows the fee per room.",
        "5. Budget Plan: one row per project with the total budget and the share for your design fee (default from Settings), furniture, construction and contingency. Amounts, allocated %, the amount left and a status (over-allocated, unallocated or fully allocated) are calculated.",
        "6. Summary: hourly rate needed and used, projects quoted, design fees quoted, projects where a percentage fee pays more, projects at the minimum, procurement revenue, gross profit and margin %, items priced above retail, the highest-fee project and budgets over-allocated.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Hourly rate from income and overhead: the rate that pays your income goal and your overhead from the hours you can actually bill. Hourly rate needed = (annual income goal + annual overhead) ÷ billable hours per year.",
        "Design phases: concept (brief, mood boards, layouts), development (selections, drawings, specifications) and documentation (drawings, schedules and specifications for the trades). Site visits are counted separately at hours per visit.",
        "Contingency: extra hours priced into a quote for revisions and the unexpected, as a % of the base fee. Design fee = base fee + contingency.",
        "Percentage fee vs hourly: a percentage fee charges a % of the project budget; an hourly fee charges the hours you expect at your rate. The workbook shows both, which pays more and the hourly rate the percentage fee works out to.",
        "Procurement and trade discount: buying furniture, lighting and finishes for the client. The trade discount is what the supplier takes off retail for you; your cost per unit = retail × (1 − trade discount %).",
        "Markup on cost vs margin on price: markup is added to your cost; margin is the share of the client price you keep. A 25% markup on a 100.00 cost is a 125.00 client price and a 20% margin, not 25%.",
        "Per-room flat fee: a fixed fee for each room, from the hours the room takes with a complexity uplift, priced at your rate. Never below your minimum project fee for the project as a whole.",
        "Budget allocation: splitting a client's total budget into your design fee, furniture, construction and contingency so the plan adds up before anything is bought. Negative remaining = over-allocated.",
        "The targets in this workbook are your own; it does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


SETTINGS_INPUTS = [
    (R_INCOME, "Annual income goal (your own)", 60000, MONEY, True, "Example default 60,000 a year. Your own placeholder, not a recommendation: what you want to pay yourself before tax."),
    (3, "Overhead: studio or office (annual)", 6000, MONEY, False, "Example default 6,000 a year. Rent, utilities and a share of home-office costs."),
    (4, "Overhead: software and subscriptions (annual)", 1800, MONEY, False, "Example default 1,800 a year. CAD, rendering, project management and accounting software."),
    (5, "Overhead: insurance (annual)", 1500, MONEY, False, "Example default 1,500 a year. Professional indemnity and liability insurance."),
    (6, "Overhead: marketing and portfolio (annual)", 2400, MONEY, False, "Example default 2,400 a year. Website, photography, directories and print."),
    (7, "Overhead: samples and travel (annual)", 1800, MONEY, False, "Example default 1,800 a year. Samples, showroom visits and unbilled travel."),
    (8, "Overhead: other (annual)", 1500, MONEY, False, "Example default 1,500 a year. Accounting, memberships, training and everything else."),
    (R_HOURS, "Billable hours per year", 1100, HRS, True, "Example default 1,100. Hours you can bill in a year after admin, marketing, holidays and quiet weeks; your own placeholder."),
    (R_CONT, "Default contingency %", 10, PCT, False, "Default 10. Added to the base fee on Fee Builder for revisions; each project can override it."),
    (R_PCTFEE, "Default percentage fee %", 15, PCT, True, "Default 15. Your own placeholder, not a recommendation: the % of project budget you would charge as a percentage fee."),
    (R_MARKUP, "Default markup on procurement %", 25, PCT, True, "Default 25. Your own placeholder, not a recommendation: added to your cost on Procurement; each item can override it."),
    (R_TRADE, "Default trade discount %", 30, PCT, False, "Default 30. Placeholder; use what your suppliers actually give you. Each item can override it."),
    (R_VISIT, "Default hours per site visit", 2, HRS, False, "Default 2. Including travel; each project can override it on Fee Builder."),
    (R_MIN, "Minimum project fee", 2500, MONEY, True, "Example default 2,500. Your own placeholder: the smallest fee you take a project on for."),
    (R_UPLIFT, "Default complexity uplift %", 10, PCT, False, "Default 10. Extra hours for a room's complexity on Room Fees; each room can override it.")]


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    for r, name, val, fmt, key, txt in SETTINGS_INPUTS:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", txt, muted=True)
    note(ws, f"B{R_INCOME}", "Your own placeholder, not a recommendation and not an industry benchmark.")
    note(ws, f"B{R_HOURS}", "Assumption: 1,100 is a placeholder. Count the hours you actually billed last year to set it.")
    note(ws, f"B{R_PCTFEE}", "Your own placeholder, not a recommendation and not an industry benchmark.")
    note(ws, f"B{R_MARKUP}", "Your own placeholder, not a recommendation and not an industry benchmark. Markup is on your cost, not a margin on price.")
    note(ws, f"B{R_TRADE}", "Assumption: 30% is a placeholder. Trade discounts vary by supplier and order size.")
    note(ws, f"B{R_MIN}", "Your own placeholder. Fee Builder flags quotes below it; Room Fees raises a project's fee to it.")
    rows_block(ws, [(R_OVERHEAD, "Total annual overhead", f"=SUM(B{R_OH_FIRST}:B{R_OH_LAST})", MONEY,
                     "formula: sum of the six overhead lines above"),
                    (R_NEEDED, "Hourly rate needed", f"=IFERROR(ROUND((B{R_INCOME}+B{R_OVERHEAD})/B{R_HOURS},2),0)", MONEY,
                     "formula: (annual income goal + total annual overhead) ÷ billable hours per year")],
               bold_rows=(R_OVERHEAD, R_NEEDED))
    label(ws, f"A{R_OWN}", "or type your own rate")
    inp(ws, f"B{R_OWN}", None, MONEY, key=True)
    label(ws, f"C{R_OWN}", "Leave blank to price at the hourly rate needed; type a rate to price every sheet at your own rate instead.", muted=True)
    rows_block(ws, [(R_RATE, "Rate used for pricing", f'=IF(B{R_OWN}<>"",B{R_OWN},B{R_NEEDED})', MONEY,
                     "formula: your own rate if typed, otherwise the hourly rate needed"),
                    (R_RATE + 1, "Total to cover per year", f"=B{R_INCOME}+B{R_OVERHEAD}", MONEY,
                     "formula: annual income goal + total annual overhead"),
                    (R_RATE + 2, "Overhead per billable hour", f"=IFERROR(ROUND(B{R_OVERHEAD}/B{R_HOURS},2),0)", MONEY,
                     "formula: total annual overhead ÷ billable hours per year (the part of the rate that pays overhead)"),
                    (R_RATE + 3, "Minimum project fee in hours", f"=IFERROR(ROUND(B{R_MIN}/B{R_RATE},2),0)", HRS,
                     "formula: minimum project fee ÷ rate used: the hours of work your minimum fee pays for")],
               bold_rows=(R_RATE,))
    note(ws, f"B{R_RATE}", "Fee Builder, Room Fees and the Summary all price at this rate.")


def build_fees(ws):
    lo, hi = FEE_ROWS; tot = hi + 1
    label(ws, "A1", "Fee Builder", bold=True).font = TITLE
    label(ws, "A2", "One row per project or quote. Hours per visit, contingency % and percentage fee % default from Settings; type your own to override them. The fee is priced at the rate used on Settings.", muted=True)
    head(ws, 4, ["Project", "Client", "Concept hours", "Development hours", "Documentation hours", "Site visits",
                 "Default hours per visit (Settings)", "or type your own hours per visit", "Hours per visit used",
                 "Default contingency % (Settings)", "or type your own contingency %", "Contingency % used",
                 "Project budget", "Default percentage fee % (Settings)", "or type your own percentage fee %",
                 "Percentage fee % used", "Total hours", "Base fee", "Contingency amount", "Design fee",
                 "Fee as % of budget", "Percentage fee amount", "Difference (percentage − design fee)",
                 "Which pays more", "Implied hourly rate of percentage fee", "Minimum fee check"],
         [20, 14, 10, 10, 10, 9, 10, 10, 9, 10, 10, 9, 12, 10, 10, 9, 9, 11, 11, 11, 10, 11, 12, 14, 11, 15])
    fm = {"C": HRS, "D": HRS, "E": HRS, "F": INT0, "H": HRS, "K": PCT, "M": MONEY, "O": PCT}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFHKMO", fm); fee_row(ws, r)
    sample = [("Hart residence", "J. Hart", 20, 30, 25, 6, None, None, 120000, None),
              ("Lake Street loft", "M. Okafor", 14, 18, 12, 4, 3, None, 60000, None),
              ("Corner cafe", "Bean & Co.", 10, 16, 14, 3, None, 15, 45000, 12),
              ("Powder room refresh", "S. Lind", 4, 6, 3, 1, None, None, 6000, None),
              ("Mill house", "R. & T. Vance", 30, 45, 40, 10, None, None, 250000, None),
              ("Office reception", "Northgate Ltd", 12, 15, 10, 3, None, None, 30000, None)]
    for i, (proj, client, con, dev, doc, visits, hpv, cont, budget, pct) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", proj); inp(ws, f"B{r}", client); inp(ws, f"C{r}", con, HRS, key=True)
        inp(ws, f"D{r}", dev, HRS, key=True); inp(ws, f"E{r}", doc, HRS, key=True); inp(ws, f"F{r}", visits, INT0)
        inp(ws, f"H{r}", hpv, HRS); inp(ws, f"K{r}", cont, PCT); inp(ws, f"M{r}", budget, MONEY, key=True)
        inp(ws, f"O{r}", pct, PCT)
    fee_notes(ws, lo)
    fee_totals(ws, lo, hi, tot)
    ws.freeze_panes = "C5"


def fee_row(ws, r):
    g = f'=IF(A{r}="","",{{}})'
    fx(ws, f"G{r}", g.format(f"{S}{R_VISIT}"), HRS)
    fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), HRS)
    fx(ws, f"J{r}", g.format(f"{S}{R_CONT}"), PCT)
    fx(ws, f"L{r}", g.format(f'IF(K{r}<>"",K{r},J{r})'), PCT)
    fx(ws, f"N{r}", g.format(f"{S}{R_PCTFEE}"), PCT)
    fx(ws, f"P{r}", g.format(f'IF(O{r}<>"",O{r},N{r})'), PCT)
    fx(ws, f"Q{r}", g.format(f"C{r}+D{r}+E{r}+F{r}*I{r}"), HRS)
    fx(ws, f"R{r}", g.format(f"ROUND(Q{r}*{S}{R_RATE},2)"), MONEY)
    fx(ws, f"S{r}", g.format(f"ROUND(R{r}*L{r}/100,2)"), MONEY)
    fx(ws, f"T{r}", g.format(f"ROUND(R{r}+S{r},2)"), MONEY)
    fx(ws, f"U{r}", g.format(f"IFERROR(ROUND(T{r}/M{r}*100,2),0)"), PCT)
    fx(ws, f"V{r}", g.format(f"ROUND(M{r}*P{r}/100,2)"), MONEY)
    fx(ws, f"W{r}", g.format(f"ROUND(V{r}-T{r},2)"), MONEY)
    fx(ws, f"X{r}", g.format(f'IF(V{r}>T{r},"{PCT_MORE}",IF(V{r}<T{r},"hourly fee","same"))'))
    fx(ws, f"Y{r}", g.format(f"IFERROR(ROUND(V{r}/Q{r},2),0)"), MONEY)
    fx(ws, f"Z{r}", g.format(f'IF(T{r}<{S}{R_MIN},"{MIN_FLAG}","")'))


def fee_notes(ws, lo):
    note(ws, f"C{lo}", "Hours you expect for the concept phase: brief, mood boards and layouts.")
    note(ws, f"H{lo + 1}", "Example override: site visits to this loft take longer because of travel.")
    note(ws, f"K{lo + 2}", "Example override: a higher contingency for a commercial fit-out with more revisions.")
    note(ws, f"O{lo + 2}", "Example override: a lower percentage fee agreed for this client.")
    note(ws, f"A{lo + 3}", "Example: a small project whose design fee falls below the minimum project fee on Settings.")
    note(ws, f"M{lo}", "The client's total project budget, used only to compare a percentage fee with your hourly fee.")
    note(ws, f"R{lo}", "Assumption: total hours × the rate used on Settings.")
    note(ws, f"U{lo}", "Design fee ÷ project budget. 0 when no budget is typed.")
    note(ws, f"W{lo}", "Positive = a percentage fee would pay you more than your hourly fee for this project.")
    note(ws, f"Y{lo}", "Percentage fee ÷ total hours: what each hour earns if you charge the percentage fee. 0 when there are no hours.")
    note(ws, f"Z{lo}", "Flags a design fee below the minimum project fee on Settings. The fee itself is not changed here.")


def fee_totals(ws, lo, hi, tot):
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" projects"', bold=True)
    for col, fmt in (("C", HRS), ("D", HRS), ("E", HRS), ("F", INT0), ("M", MONEY), ("Q", HRS), ("R", MONEY),
                     ("S", MONEY), ("T", MONEY), ("V", MONEY), ("W", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"U{tot}", f"=IFERROR(ROUND(T{tot}/M{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"Y{tot}", f"=IFERROR(ROUND(V{tot}/Q{tot},2),0)", MONEY, bold=True)
    note(ws, f"U{tot}", "From the totals, not an average of the rows.")


def build_procurement(ws):
    lo, hi = PR_ROWS; tot = hi + 1
    label(ws, "A1", "Procurement", bold=True).font = TITLE
    label(ws, "A2", "One row per item you buy for a client. Trade discount % and markup % default from Settings; type your own to override them, or type a client price to set it directly. Freight and handling is for the whole line and is passed to the client at cost.", muted=True)
    head(ws, 4, ["Project", "Item", "Supplier", "Retail price (per unit)", "Default trade discount % (Settings)",
                 "or type your own trade discount %", "Trade discount % used", "Default markup % (Settings)",
                 "or type your own markup %", "Markup % used", "or type your own client price (per unit)",
                 "Freight and handling (line)", "Quantity", "Your cost per unit", "Client price used (per unit)",
                 "Line cost", "Line revenue", "Gross profit", "Margin %", "Client saving vs retail"],
         [18, 20, 14, 11, 10, 10, 9, 10, 10, 9, 12, 11, 8, 11, 11, 11, 11, 11, 9, 12])
    fm = {"D": MONEY, "F": PCT, "I": PCT, "K": MONEY, "L": MONEY, "M": INT0}
    rows = range(lo, hi + 1)
    for r in rows:
        blank(ws, r, "ABCDFIKLM", fm); procurement_row(ws, r)
    project_list(ws, [f"A{r}" for r in rows])
    sample = [("Hart residence", "Sofa", "Supplier A", 3200, None, None, None, 150, 1),
              ("Hart residence", "Dining chairs", "Supplier B", 420, None, None, None, 80, 6),
              ("Hart residence", "Pendant lights", "Supplier C", 380, None, None, None, 40, 3),
              ("Hart residence", "Rug", "Supplier A", 1800, None, 20, None, 60, 1),
              ("Lake Street loft", "Bed frame", "Supplier D", 2100, None, None, None, 120, 1),
              ("Lake Street loft", "Bedside tables", "Supplier D", 450, None, None, None, 30, 2),
              ("Lake Street loft", "Curtains (made to measure)", "Workroom E", 1600, 10, None, None, 0, 1),
              ("Corner cafe", "Cafe chairs", "Supplier F", 160, None, None, None, 200, 24),
              ("Corner cafe", "Bar stools", "Supplier F", 240, None, None, None, 90, 8),
              ("Corner cafe", "Wall lights", "Supplier C", 210, None, None, 175, 25, 6)]
    for i, (proj, item, sup, retail, disc, markup, price, freight, qty) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", proj); inp(ws, f"B{r}", item); inp(ws, f"C{r}", sup)
        inp(ws, f"D{r}", retail, MONEY, key=True); inp(ws, f"F{r}", disc, PCT); inp(ws, f"I{r}", markup, PCT)
        inp(ws, f"K{r}", price, MONEY); inp(ws, f"L{r}", freight, MONEY); inp(ws, f"M{r}", qty, INT0, key=True)
    procurement_notes(ws, lo)
    procurement_totals(ws, lo, hi, tot)
    ws.freeze_panes = "C5"


def procurement_row(ws, r):
    g = f'=IF(A{r}="","",{{}})'
    fx(ws, f"E{r}", g.format(f"{S}{R_TRADE}"), PCT)
    fx(ws, f"G{r}", g.format(f'IF(F{r}<>"",F{r},E{r})'), PCT)
    fx(ws, f"H{r}", g.format(f"{S}{R_MARKUP}"), PCT)
    fx(ws, f"J{r}", g.format(f'IF(I{r}<>"",I{r},H{r})'), PCT)
    fx(ws, f"N{r}", g.format(f"ROUND(D{r}*(1-G{r}/100),2)"), MONEY)
    fx(ws, f"O{r}", g.format(f'IF(K{r}<>"",K{r},ROUND(N{r}*(1+J{r}/100),2))'), MONEY)
    fx(ws, f"P{r}", g.format(f"ROUND(N{r}*M{r}+L{r},2)"), MONEY)
    fx(ws, f"Q{r}", g.format(f"ROUND(O{r}*M{r}+L{r},2)"), MONEY)
    fx(ws, f"R{r}", g.format(f"ROUND(Q{r}-P{r},2)"), MONEY)
    fx(ws, f"S{r}", g.format(f"IFERROR(ROUND(R{r}/Q{r}*100,2),0)"), PCT)
    fx(ws, f"T{r}", g.format(f"ROUND((D{r}-O{r})*M{r},2)"), MONEY)


def procurement_notes(ws, lo):
    note(ws, f"D{lo}", "The supplier's retail or list price per unit, the price the client could pay without you.")
    note(ws, f"I{lo + 3}", "Example override: a lower markup on a large rug to stay below retail.")
    note(ws, f"F{lo + 6}", "Example override: a workroom that gives only a 10% trade discount, so the client price at your markup ends up above retail and the saving is negative.")
    note(ws, f"K{lo + 9}", "Example: a client price typed directly instead of cost plus markup.")
    note(ws, f"L{lo}", "Assumption: freight, delivery and handling for the whole line, paid by you and passed to the client at cost (in both line cost and line revenue).")
    note(ws, f"O{lo}", "Your own client price if typed, otherwise your cost per unit × (1 + markup %).")
    note(ws, f"S{lo}", "Margin % = gross profit ÷ line revenue: the share of the price you keep, not the markup on cost.")
    note(ws, f"T{lo}", "(Retail price − client price) × quantity. Negative = the client pays more than retail through you.")


def procurement_totals(ws, lo, hi, tot):
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(B{lo}:B{hi})", '0" items"', bold=True)
    for col in "LPQRT":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"M{tot}", f"=SUM(M{lo}:M{hi})", INT0, bold=True)
    fx(ws, f"S{tot}", f"=IFERROR(ROUND(R{tot}/Q{tot}*100,2),0)", PCT, bold=True)
    note(ws, f"S{tot}", "Overall margin % from the totals, not an average of the rows.")


def build_rooms(ws):
    lo, hi = ROOM_ROWS; tot = hi + 1
    label(ws, "A1", "Room Fees", bold=True).font = TITLE
    label(ws, "A2", "One row per room. Complexity uplift % and rate default from Settings; type your own to override them. The block below totals each project and applies your minimum project fee.", muted=True)
    head(ws, 4, ["Project", "Room", "Hours for the room", "Default complexity uplift % (Settings)",
                 "or type your own uplift %", "Uplift % used", "Default rate (Settings)", "or type your own rate",
                 "Rate used", "Hours with uplift", "Room fee"],
         [20, 20, 10, 11, 10, 9, 11, 10, 10, 10, 12])
    fm = {"C": HRS, "E": PCT, "H": MONEY}
    rows = range(lo, hi + 1)
    for r in rows:
        blank(ws, r, "ABCEH", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_UPLIFT}"), PCT)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), PCT)
        fx(ws, f"G{r}", g.format(f"{S}{R_RATE}"), MONEY)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), MONEY)
        fx(ws, f"J{r}", g.format(f"ROUND(C{r}*(1+F{r}/100),2)"), HRS)
        fx(ws, f"K{r}", g.format(f"ROUND(J{r}*I{r},2)"), MONEY)
    sample = [("Hart residence", "Living room", 18, None, None), ("Hart residence", "Kitchen", 24, 20, None),
              ("Hart residence", "Primary bedroom", 14, None, None), ("Lake Street loft", "Open-plan living", 20, None, None),
              ("Lake Street loft", "Bedroom", 10, None, None), ("Lake Street loft", "Bathroom", 12, None, 80),
              ("Corner cafe", "Dining area", 22, None, None), ("Corner cafe", "Restroom", 6, None, None)]
    for i, (proj, room, hrs, uplift, rate) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", proj); inp(ws, f"B{r}", room); inp(ws, f"C{r}", hrs, HRS, key=True)
        inp(ws, f"E{r}", uplift, PCT); inp(ws, f"H{r}", rate, MONEY)
    note(ws, f"C{lo}", "Design hours you expect this room to take before any complexity uplift.")
    note(ws, f"E{lo + 1}", "Example override: a kitchen with joinery and services is more complex than the default.")
    note(ws, f"H{lo + 5}", "Example override: a higher rate for a bathroom with detailed tiling drawings.")
    note(ws, f"J{lo}", "Assumption: hours × (1 + complexity uplift %).")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(B{lo}:B{hi})", '0" rooms"', bold=True)
    for col, fmt in (("C", HRS), ("J", HRS), ("K", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    build_room_projects(ws)
    project_list(ws, [f"A{r}" for r in rows] + [f"A{r}" for r in range(RP_ROWS[0], RP_ROWS[1] + 1)])
    ws.freeze_panes = "C5"


def build_room_projects(ws):
    lo, hi = ROOM_ROWS; plo, phi = RP_ROWS
    label(ws, f"A{plo - 2}", "PER-PROJECT FLAT FEE", bold=True)
    labels = ["Project", "Rooms", "Room fees", "Site visit hours", "Site visit fee", "Fee before minimum",
              "Project fee (at least the minimum)", "Minimum check", "Fee per room"]
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=plo - 1, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    names = f"$A${lo}:$A${hi}"
    for r in range(plo, phi + 1):
        blank(ws, r, "AD", {"D": HRS})
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"B{r}", g.format(f"COUNTIF({names},A{r})"), INT0)
        fx(ws, f"C{r}", g.format(f"SUMIF({names},A{r},$K${lo}:$K${hi})"), MONEY)
        fx(ws, f"E{r}", g.format(f"ROUND(D{r}*{S}{R_RATE},2)"), MONEY)
        fx(ws, f"F{r}", g.format(f"ROUND(C{r}+E{r},2)"), MONEY)
        fx(ws, f"G{r}", g.format(f"MAX(F{r},{S}{R_MIN})"), MONEY)
        fx(ws, f"H{r}", g.format(f'IF(F{r}<{S}{R_MIN},"{MIN_FLAG}","")'))
        fx(ws, f"I{r}", g.format(f"IFERROR(ROUND(G{r}/B{r},2),0)"), MONEY)
    for i, (proj, visits) in enumerate([("Hart residence", 8), ("Lake Street loft", 6), ("Corner cafe", 4)]):
        inp(ws, f"A{plo + i}", proj); inp(ws, f"D{plo + i}", visits, HRS, key=True)
    note(ws, f"A{plo}", "Type each project once; its rooms above are counted and summed by name.")
    note(ws, f"D{plo}", "Assumption: site visit hours for the whole project, priced at the rate used on Settings.")
    note(ws, f"G{plo}", "The larger of the fee before minimum and the minimum project fee on Settings.")
    note(ws, f"I{plo}", "Project fee ÷ rooms. 0 when the project has no rooms above.")


def build_budget(ws):
    lo, hi = BP_ROWS; tot = hi + 1
    label(ws, "A1", "Budget Plan", bold=True).font = TITLE
    label(ws, "A2", "One row per project. The design fee % defaults from the percentage fee % on Settings; type your own to override it. Type the furniture, construction and contingency shares as % of the total budget.", muted=True)
    head(ws, 4, ["Project", "Total budget", "Default design fee % (Settings)", "or type your own design fee %",
                 "Design fee % used", "Furniture %", "Construction %", "Contingency %", "Design fee amount",
                 "Furniture amount", "Construction amount", "Contingency amount", "Allocated %", "Allocated amount",
                 "Remaining", "Status"],
         [20, 13, 11, 11, 10, 10, 11, 11, 12, 12, 12, 12, 10, 12, 12, 16])
    fm = {"B": MONEY, "D": PCT, "F": PCT, "G": PCT, "H": PCT}
    rows = range(lo, hi + 1)
    for r in rows:
        blank(ws, r, "ABDFGH", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"C{r}", g.format(f"{S}{R_PCTFEE}"), PCT)
        fx(ws, f"E{r}", g.format(f'IF(D{r}<>"",D{r},C{r})'), PCT)
        for src, dst in zip("EFGH", "IJKL"):
            fx(ws, f"{dst}{r}", g.format(f"ROUND(B{r}*{src}{r}/100,2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"E{r}+F{r}+G{r}+H{r}"), PCT)
        fx(ws, f"N{r}", g.format(f"ROUND(I{r}+J{r}+K{r}+L{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"ROUND(B{r}-N{r},2)"), MONEY)
        fx(ws, f"P{r}", g.format(f'IF(O{r}<0,"{OVER}",IF(O{r}>0,"unallocated","fully allocated"))'))
    sample = [("Hart residence", 120000, None, 40, 30, 10), ("Lake Street loft", 60000, None, 45, 30, 10),
              ("Corner cafe", 45000, 12, 35, 50, 10), ("Mill house", 250000, None, 30, 40, 10)]
    for i, (proj, budget, fee, furn, cons, cont) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", proj); inp(ws, f"B{r}", budget, MONEY, key=True); inp(ws, f"D{r}", fee, PCT)
        inp(ws, f"F{r}", furn, PCT); inp(ws, f"G{r}", cons, PCT); inp(ws, f"H{r}", cont, PCT)
    project_list(ws, [f"A{r}" for r in rows])
    note(ws, f"D{lo + 2}", "Example override: the lower percentage fee agreed with this client.")
    note(ws, f"G{lo + 2}", "Example: construction set high, so the plan adds up to more than the budget and shows as over-allocated.")
    note(ws, f"H{lo}", "Assumption: a reserve for the unexpected, as a % of the total budget. Your own placeholder.")
    note(ws, f"O{lo}", "Total budget − allocated amount. Negative = over-allocated; positive = budget not yet assigned.")
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "BIJKLNO":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"M{tot}", f"=IFERROR(ROUND(N{tot}/B{tot}*100,2),0)", PCT, bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    flo, fhi = FEE_ROWS; ftot = fhi + 1; plo, phi = PR_ROWS; ptot = phi + 1; blo, bhi = BP_ROWS
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 80
    label(ws, "A1", "Summary", bold=True).font = TITLE
    fees = f"{FB}$T${flo}:$T${fhi}"
    rows = [(2, "Hourly rate needed", f"={S}{R_NEEDED}", MONEY, "from Settings: (income goal + overhead) ÷ billable hours"),
            (3, "Rate used for pricing", f"={S}{R_RATE}", MONEY, "from Settings: your own rate if typed, otherwise the rate needed"),
            (4, "Projects quoted", f"=COUNTA({FB}$A${flo}:$A${fhi})", INT0, "rows with a project name on Fee Builder"),
            (5, "Total design fees quoted", f"={FB}T{ftot}", MONEY, "Fee Builder totals"),
            (6, "Projects where the percentage fee pays more", f'=COUNTIF({FB}$X${flo}:$X${fhi},"{PCT_MORE}")', INT0, "budget × percentage fee % is above the hourly design fee"),
            (7, "Projects at the minimum fee", f'=COUNTIF({FB}$Z${flo}:$Z${fhi},"{MIN_FLAG}")', INT0, "design fee below the minimum project fee on Settings"),
            (8, "Procurement revenue", f"={PR}Q{ptot}", MONEY, "Procurement totals, freight included at cost"),
            (9, "Procurement gross profit", f"={PR}R{ptot}", MONEY, "line revenue − line cost"),
            (10, "Procurement margin %", f"={PR}S{ptot}", PCT, "gross profit ÷ revenue, from the totals"),
            (11, "Items priced above retail", f'=COUNTIF({PR}$T${plo}:$T${phi},"<0")', INT0, "items where the client saving vs retail is negative"),
            (12, "Highest-fee project", f'=IFERROR(INDEX({FB}$A${flo}:$A${fhi},MATCH(MAX({fees}),{fees},0)),"")', None, "largest design fee on Fee Builder"),
            (13, "Its design fee", f"=IFERROR(MAX({fees}),0)", MONEY, "highest design fee on Fee Builder"),
            (14, "Budgets over-allocated", f'=COUNTIF({BP}$P${blo}:$P${bhi},"{OVER}")', INT0, "Budget Plan rows where the shares add up to more than the budget")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))
    note(ws, "B6", "Count only. Whether to charge a percentage fee is your decision; this is arithmetic, not advice.")
    note(ws, "B11", "Items where the client would pay more than retail through you. Check the trade discount and markup on those rows.")


README = """# Interior Design Pricing Workbook

Thank you for buying the workbook. Open `interior-design-pricing-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** your annual income goal, six annual overhead lines, billable hours per year, default contingency %, percentage fee %, procurement markup %, trade discount %, hours per site visit, minimum project fee and complexity uplift %; the hourly rate you need follows, and you can type your own rate to price at instead.
2. **Fee Builder:** each project or quote with concept, development and documentation hours, site visits, hours per visit, contingency % and percentage fee % (or your own) and the project budget; total hours, design fee, fee as % of budget, the percentage fee, which pays more, its implied hourly rate and a minimum fee check follow.
3. **Procurement:** each item with project, supplier, retail price, trade discount % and markup % (or your own), an optional client price, freight and quantity; your cost, client price, line cost and revenue, gross profit, margin % and the client's saving against retail follow.
4. **Room Fees:** each room with hours, complexity uplift % and rate (or your own) gives a room fee; the block below totals each project with site visits, applies your minimum project fee and shows the fee per room.
5. **Budget Plan:** each project's total budget split into design fee, furniture, construction and contingency; amounts, allocated %, remaining and an over-allocated or unallocated status follow.
6. **Summary:** hourly rate needed and used, projects and design fees quoted, projects where a percentage fee pays more, projects at the minimum, procurement revenue, profit and margin, items above retail, the highest-fee project and budgets over-allocated.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_fees(wb.create_sheet("Fee Builder"))
    build_procurement(wb.create_sheet("Procurement")); build_rooms(wb.create_sheet("Room Fees"))
    build_budget(wb.create_sheet("Budget Plan")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
