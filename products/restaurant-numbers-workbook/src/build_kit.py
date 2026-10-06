"""Build the Restaurant Numbers Workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "restaurant-numbers-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; QTY = "0.000"
RECIPE_ROWS = (5, 44); MENU_ROWS = (5, 34); WEEK_ROWS = (5, 16); STAFF_ROWS = (5, 24)
S = "Settings!$B$"  # settings cell prefix
ROLES = '"kitchen,front of house,bar,management"'


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Restaurant Numbers Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: your target food cost %, target labour cost %, payroll taxes and benefits, waste allowance, seats, hours open, dining minutes and days open. Every other sheet reads them.",
        "2. Recipes: one row per ingredient line, with the recipe name repeated on each line. Enter portions on the first line of each recipe to get the cost per portion, waste allowance included.",
        "3. Menu: one row per dish. Cost per portion is looked up from Recipes by dish name (type over it to use your own figure). Add the menu price and weekly units to get food cost %, gross profit and the price at your target food cost.",
        "4. Weekly P&L: one row per week with sales, food, beverage and labour cost and other operating costs. Prime cost, food cost %, labour % and operating margin are calculated, with totals in row 17.",
        "5. Labour: one row per staff member for a typical week with role, hours and wage. Loaded weekly cost, labour % against projected sales and the hours over or under your target are calculated.",
        "6. Summary: menu items, blended food cost %, best margin dish, highest food cost dish, weekly menu revenue, projected labour %, latest week prime cost % and seat capacity.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Food cost % = cost per portion ÷ menu price. Price at target food cost = cost per portion ÷ target food cost %.",
        "Prime cost = food cost + beverage cost + labour cost. Prime cost % = prime cost ÷ total sales.",
        "Labour % = labour cost ÷ sales. Labour cost on the Labour sheet includes your payroll taxes and benefits % from Settings.",
        "Waste allowance adds a percentage to every recipe's ingredient cost for trim, spoilage and over-portioning.",
        "The targets in Settings are your own; this workbook does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [("Target food cost %", 30, "0.0", "Your own target: ingredient cost as a share of the menu price."),
            ("Target labour cost %", 28, "0.0", "Your own target: loaded labour cost as a share of sales."),
            ("Payroll taxes and benefits %", 12, "0.0", "Added on top of wages for the true cost of each hour."),
            ("Waste allowance %", 5, "0.0", "Added to recipe ingredient cost for trim, spoilage and over-portioning."),
            ("Seats", 60, "0", "Seats available for diners."),
            ("Hours open per day", 8, "0.0", "Service hours per day."),
            ("Average dining minutes", 75, "0", "How long a table is typically occupied."),
            ("Days open per week", 6, "0", "Used for covers per week on the Summary.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True); label(ws, f"C{i}", note, muted=True)


def build_recipes(ws):
    lo, hi = RECIPE_ROWS
    label(ws, "A1", "Recipes", bold=True).font = TITLE
    label(ws, "A2", "One row per ingredient line. Repeat the recipe name on each line; enter portions on the first line only.", muted=True)
    head(ws, 4, ["Recipe", "Ingredient", "Quantity", "Unit cost", "Line cost", "Portions", "Cost per portion"],
         [22, 22, 11, 12, 12, 10, 15])
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDF", {"C": QTY, "D": MONEY, "F": "0"})
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(C{r}*D{r},2))', MONEY)
        fx(ws, f"G{r}", f'=IF(OR(A{r}="",F{r}=""),"",ROUND(SUMIF($A${lo}:$A${hi},A{r},$E${lo}:$E${hi})*(1+{S}5/100)/F{r},2))', MONEY)
    sample = [("Pasta pomodoro", "Pasta (kg)", 0.12, 2.40, 1), ("Pasta pomodoro", "Tomato sauce (kg)", 0.15, 3.00, None),
              ("Pasta pomodoro", "Cheese (kg)", 0.03, 12.00, None), ("Steak frites", "Steak (kg)", 0.25, 28.00, 1),
              ("Steak frites", "Potatoes (kg)", 0.3, 1.20, None), ("Steak frites", "Butter (kg)", 0.03, 9.00, None)]
    for i, (rec, ing, q, uc, por) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", rec); inp(ws, f"B{r}", ing); inp(ws, f"C{r}", q, QTY); inp(ws, f"D{r}", uc, MONEY, key=True)
        if por is not None: inp(ws, f"F{r}", por, "0", key=True)
    ws.freeze_panes = "A5"


def build_menu(ws):
    lo, hi = MENU_ROWS; tot = hi + 1
    rl, rh = RECIPE_ROWS
    label(ws, "A1", "Menu", bold=True).font = TITLE
    label(ws, "A2", "Cost per portion is looked up from Recipes by dish name; it is blue because you may type over it with your own figure.", muted=True)
    head(ws, 4, ["Dish", "Cost per portion", "Menu price", "Food cost %", "Gross profit", "Price at target food cost",
                 "Units sold per week", "Weekly food cost", "Weekly revenue", "Weekly gross profit"],
         [22, 13, 12, 12, 12, 15, 12, 14, 14, 15])
    look = 'IF(A{r}="","",IFERROR(INDEX(Recipes!$G${rl}:$G${rh},MATCH(A{r},Recipes!$A${rl}:$A${rh},0)),""))'
    for r in range(lo, hi + 1):
        blank(ws, r, "ACG", {"C": MONEY, "G": "0"})
        inp(ws, f"B{r}", "=" + look.format(r=r, rl=rl, rh=rh), MONEY)
        fx(ws, f"D{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}/C{r}*100,2),""))', PCT)
        fx(ws, f"E{r}", f'=IF(A{r}="","",IFERROR(C{r}-B{r},""))', MONEY)
        fx(ws, f"F{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}/({S}2/100),2),""))', MONEY)
        fx(ws, f"H{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}*G{r},2),""))', MONEY)
        fx(ws, f"I{r}", f'=IF(A{r}="","",IFERROR(ROUND(C{r}*G{r},2),""))', MONEY)
        fx(ws, f"J{r}", f'=IF(A{r}="","",IFERROR(I{r}-H{r},""))', MONEY)
    sample = [("Pasta pomodoro", None, 12.50, 180), ("Steak frites", None, 24.00, 60), ("House salad", 1.80, 7.50, 90)]
    for i, (dish, cost, price, units) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dish); inp(ws, f"C{r}", price, MONEY, key=True); inp(ws, f"G{r}", units, "0", key=True)
        if cost is not None: inp(ws, f"B{r}", cost, MONEY, note="Typed cost: this dish has no recipe on the Recipes sheet.")
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "GHIJ": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", "0" if col == "G" else MONEY, bold=True)
    label(ws, f"A{tot+1}", "Blended food cost %", bold=True)
    fx(ws, f"D{tot+1}", f"=IFERROR(ROUND(H{tot}/I{tot}*100,2),0)", PCT, bold=True)
    label(ws, f"E{tot+1}", "weekly food cost ÷ weekly revenue, weighted by units sold", muted=True)
    ws.freeze_panes = "B5"


def build_pnl(ws):
    lo, hi = WEEK_ROWS; tot = hi + 1
    label(ws, "A1", "Weekly P&L", bold=True).font = TITLE
    label(ws, "A2", "One row per week. Labour cost is typed from your payroll; the Labour sheet gives a typical week for comparison.", muted=True)
    head(ws, 4, ["Week ending", "Food sales", "Beverage sales", "Total sales", "Food cost", "Beverage cost", "Labour cost",
                 "Prime cost", "Prime cost %", "Food cost %", "Labour %", "Other operating costs", "Operating profit", "Margin %"],
         [13, 12, 12, 12, 12, 12, 12, 12, 11, 11, 11, 13, 13, 11])
    fm = {c: MONEY for c in "BCEFGL"}; fm["A"] = DATE
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEFGL", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"B{r}+C{r}"), MONEY)
        fx(ws, f"H{r}", g.format(f"E{r}+F{r}+G{r}"), MONEY)
        fx(ws, f"I{r}", g.format(f'IFERROR(ROUND(H{r}/D{r}*100,2),"")'), PCT)
        fx(ws, f"J{r}", g.format(f'IFERROR(ROUND(E{r}/B{r}*100,2),"")'), PCT)
        fx(ws, f"K{r}", g.format(f'IFERROR(ROUND(G{r}/D{r}*100,2),"")'), PCT)
        fx(ws, f"M{r}", g.format(f"D{r}-H{r}-L{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f'IFERROR(ROUND(M{r}/D{r}*100,2),"")'), PCT)
    ws[f"G{lo}"].comment = Comment("Type the week's labour cost from payroll, including taxes and benefits.", "Knackdesk")
    y = dt.date.today().year
    sample = [(dt.date(y, 9, 6), 4400, 1500, 1350, 380, 1650, 1200),
              (dt.date(y, 9, 13), 4650, 1620, 1420, 410, 1700, 1200),
              (dt.date(y, 9, 20), 4300, 1480, 1390, 395, 1720, 1250)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCEFGL", row): inp(ws, f"{col}{r}", val, fm.get(col), key=(i == 0 and col in "BG"))
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "BCDEFGHLM": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    for col, num, den in [("I", "H", "D"), ("J", "E", "B"), ("K", "G", "D"), ("N", "M", "D")]:
        fx(ws, f"{col}{tot}", f"=IFERROR(ROUND({num}{tot}/{den}{tot}*100,2),0)", PCT, bold=True)
    ws.freeze_panes = "B5"


def build_labour(ws):
    lo, hi = STAFF_ROWS; tot = hi + 1; m_tot = MENU_ROWS[1] + 1
    label(ws, "A1", "Labour", bold=True).font = TITLE
    label(ws, "A2", "One row per staff member for a typical week. Weekly cost includes payroll taxes and benefits from Settings.", muted=True)
    head(ws, 4, ["Name", "Role", "Hours", "Hourly wage", "Weekly cost"], [30, 16, 10, 12, 14])
    ws.column_dimensions["C"].width = 10
    dv = DataValidation(type="list", formula1=ROLES, allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCD", {"C": "0.0", "D": MONEY}); dv.add(f"B{r}")
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(C{r}*D{r}*(1+{S}4/100),2))', MONEY)
    sample = [("Head chef", "kitchen", 24, 18), ("Line cook", "kitchen", 20, 14), ("Server A", "front of house", 16, 12),
              ("Server B", "front of house", 12, 12), ("Bartender", "bar", 10, 13), ("Manager", "management", 8, 20)]
    for i, (name, role, hours, wage) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", role); inp(ws, f"C{r}", hours, "0.0", key=True); inp(ws, f"D{r}", wage, MONEY, key=True)
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"C{tot}", f"=SUM(C{lo}:C{hi})", "0.0", bold=True); fx(ws, f"E{tot}", f"=SUM(E{lo}:E{hi})", MONEY, bold=True)
    label(ws, "A27", "Projected weekly sales", bold=True)
    inp(ws, "B27", f"=Menu!I{m_tot}", MONEY, key=True, note="Defaults to weekly revenue from the Menu sheet; type over it with your own forecast.")
    label(ws, "C27", "defaults to Menu weekly revenue; type over it", muted=True)
    rows = [(28, "Labour %", f"=IFERROR(ROUND(E{tot}/B27*100,2),0)", PCT, "weekly cost ÷ projected sales"),
            (29, "Allowed labour cost at target", f"=ROUND(B27*{S}3/100,2)", MONEY, "projected sales × target labour %"),
            (30, "Hours over (+) or under (−) target", f"=IFERROR(ROUND((E{tot}-B29)/(AVERAGE(D{lo}:D{hi})*(1+{S}4/100)),1),0)", "0.0",
             "(weekly cost − allowed) ÷ average loaded hourly wage")]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True); label(ws, f"C{r}", note, muted=True)
    ws.freeze_panes = "A5"


def build_summary(ws):
    ml, mh = MENU_ROWS; m_tot = mh + 1; wl, wh = WEEK_ROWS
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 55
    label(ws, "A1", "Summary", bold=True).font = TITLE
    rows = [(2, "Menu items", f"=COUNTA(Menu!A{ml}:A{mh})", "0", None),
            (3, "Blended food cost %", f"=Menu!D{m_tot+1}", PCT, "weighted by units sold"),
            (4, "Best margin dish", f'=IFERROR(INDEX(Menu!A{ml}:A{mh},MATCH(MAX(Menu!E{ml}:E{mh}),Menu!E{ml}:E{mh},0)),"")', None, "highest gross profit per portion"),
            (5, "Highest food cost % dish", f'=IFERROR(INDEX(Menu!A{ml}:A{mh},MATCH(MAX(Menu!D{ml}:D{mh}),Menu!D{ml}:D{mh},0)),"")', None, "the dish to re-price or re-cost first"),
            (6, "Weekly revenue from menu", f"=Menu!I{m_tot}", MONEY, None),
            (7, "Projected labour %", "=Labour!B28", PCT, "from the Labour sheet"),
            (8, "Latest week prime cost %", f'=IF(COUNT(\'Weekly P&L\'!D{wl}:D{wh})=0,"",INDEX(\'Weekly P&L\'!I{wl}:I{wh},COUNT(\'Weekly P&L\'!D{wl}:D{wh})))', PCT,
             "last filled week on Weekly P&L (fill weeks top down)")]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True)
        if note: label(ws, f"C{r}", note, muted=True)
    label(ws, "A10", "Seat capacity", bold=True).font = TITLE
    label(ws, "A11", "Occupancy %"); inp(ws, "B11", 70, "0.0", key=True); label(ws, "C11", "share of seats filled on an average turn", muted=True)
    cap = [(12, "Turns per day", f"=IFERROR(ROUND({S}7*60/{S}8,2),0)", "0.00", "hours open × 60 ÷ average dining minutes"),
           (13, "Covers per day", f"=ROUND({S}6*B12*B11/100,0)", "0", "seats × turns × occupancy"),
           (14, "Covers per week", f"=B13*{S}9", "0", "covers per day × days open")]
    for r, name, f, fmt, note in cap:
        label(ws, f"A{r}", name, bold=(r == 14)); fx(ws, f"B{r}", f, fmt, bold=(r == 14)); label(ws, f"C{r}", note, muted=True)


README = """# Restaurant Numbers Workbook

Thank you for buying the workbook. Open `restaurant-numbers-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** your target food cost % and labour %, payroll taxes and benefits, waste allowance, seats, hours, dining minutes and days open.
2. **Recipes:** ingredient lines per recipe with quantity and unit cost; cost per portion includes your waste allowance.
3. **Menu:** dishes with cost per portion (looked up from Recipes), menu price and weekly units; food cost %, gross profit and the price at your target food cost follow.
4. **Weekly P&L:** sales and costs per week with prime cost %, food cost %, labour % and operating margin.
5. **Labour:** staff hours and wages with loaded weekly cost, labour % and the hours over or under your target.
6. **Summary:** blended food cost %, best margin dish, highest food cost dish, projected labour %, latest prime cost % and seat capacity.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_recipes(wb.create_sheet("Recipes")); build_menu(wb.create_sheet("Menu"))
    build_pnl(wb.create_sheet("Weekly P&L")); build_labour(wb.create_sheet("Labour")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
