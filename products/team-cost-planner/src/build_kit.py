"""Build the Team Cost & Headcount Planner workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
TEAM_ROWS = (5, 40); RAISE_ROWS = (5, 20); PTO_ROWS = (5, 40); HIRE_ROWS = (5, 16)
S = "Settings!$B$"  # settings cell prefix


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Team Cost & Headcount Planner", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: your defaults for employer contributions, hours, weeks worked, productive share, working days and PTO. Every other sheet reads them.",
        "2. Team: one row per person with salary, benefits and overheads. Employer contributions, total annual and monthly cost, cost per productive hour and the cost multiplier are calculated, with team totals in row 41.",
        "3. Raises: name a person, choose percent or amount and enter the raise. New salary, monthly difference and the yearly cost to the business (raise plus contributions) are calculated.",
        "4. PTO: one row per person with entitlement, pay frequency and pay periods completed this year, plus carried-over and taken days. Accrual per period, accrued to date, balance and days still to earn are calculated.",
        "5. Hiring & Turnover: planned hires with their cost in the plan year, and the cost of turnover per leaver and per year using your average cost per employee.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "Blue cells that already contain a formula (defaults pulled from Settings or Team) may be overwritten with your own figure.",
        "", "DEFINITIONS",
        "Total annual cost = salary + salary × employer contribution rate + benefits + overheads. Monthly cost = annual ÷ 12.",
        "Productive hours = hours per week × weeks worked × productive share. Cost per productive hour = total annual cost ÷ productive hours.",
        "PTO per period = annual days ÷ pay periods per year (weekly 52, biweekly 26, semimonthly 24, monthly 12). Balance = carried over + accrued − taken.",
        "Turnover cost per leaver = recruiting + vacancy (daily cost × vacant days) + onboarding + ramp-up (daily cost × ramp days × (1 − ramp productivity)).",
        "Rates for employer contributions and leave rules differ by country and contract; enter your own. Arithmetic only; not legal, tax or HR advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [("Employer contributions (% of salary)", 12, "0.00", "Social security, pension, health or unemployment contributions the employer pays on top of gross salary."),
            ("Hours per week", 40, "0.0", "Contracted hours; overridable per person on the Team sheet."),
            ("Weeks worked per year", 46, "0.0", "52 minus holiday, public holidays and typical sick leave."),
            ("Productive share of hours (%)", 80, "0.0", "Share of contracted hours spent on billable or output work."),
            ("Working days per year", 240, "0", "Used for the daily cost in turnover calculations."),
            ("Ramp-up productivity for new hires (%)", 50, "0.0", "Average output during the ramp-up period as a share of full productivity."),
            ("Default paid time off per year (days)", 20, "0.0", "Overridable per person on the PTO sheet."),
            ("Default pay frequency", "biweekly", None, "weekly, biweekly, semimonthly or monthly."),
            ("Plan year", dt.date.today().year, "0", "Hires starting in this year are prorated by the months remaining.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True); label(ws, f"C{i}", note, muted=True)
    dv = DataValidation(type="list", formula1='"weekly,biweekly,semimonthly,monthly"', allow_blank=True); ws.add_data_validation(dv); dv.add("B9")


def build_team(ws):
    lo, hi = TEAM_ROWS; tot = hi + 1
    label(ws, "A1", "Team", bold=True).font = TITLE
    label(ws, "A2", "One row per person. Contribution rate and hours default to Settings; overwrite them per person if needed.", muted=True)
    head(ws, 4, ["Name", "Role", "Start date", "Annual salary", "Employer contrib. %", "Benefits / yr", "Overheads / yr", "Employer contributions", "Total annual cost", "Monthly cost", "Hours / week", "Cost per productive hour", "Cost multiple of salary"],
         [18, 18, 12, 14, 12, 13, 13, 14, 15, 13, 11, 14, 13])
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDFG", {"C": DATE, "D": MONEY, "F": MONEY, "G": MONEY})
        fx(ws, f"E{r}", f'=IF(A{r}="","",{S}2)', "0.00").font = BLUE
        fx(ws, f"K{r}", f'=IF(A{r}="","",{S}3)', "0.0").font = BLUE
        fx(ws, f"H{r}", f'=IF(A{r}="","",ROUND(D{r}*E{r}/100,2))', MONEY)
        fx(ws, f"I{r}", f'=IF(A{r}="","",D{r}+H{r}+F{r}+G{r})', MONEY)
        fx(ws, f"J{r}", f'=IF(A{r}="","",ROUND(I{r}/12,2))', MONEY)
        fx(ws, f"L{r}", f'=IF(A{r}="","",IFERROR(ROUND(I{r}/(K{r}*{S}4*{S}5/100),2),""))', MONEY)
        fx(ws, f"M{r}", f'=IF(A{r}="","",IFERROR(ROUND(I{r}/D{r},2),""))', "0.00")
    y = dt.date.today().year
    sample = [("Ana Example", "Designer", dt.date(y - 2, 3, 1), 50000, 4000, 3000), ("Ben Example", "Developer", dt.date(y - 1, 9, 15), 64000, 4500, 3200), ("Cara Example", "Support", dt.date(y, 1, 10), 38000, 3000, 2000)]
    for i, (n, role, start, sal, ben, over) in enumerate(sample):
        r = lo + i; inp(ws, f"A{r}", n, key=(i == 0)); inp(ws, f"B{r}", role); inp(ws, f"C{r}", start, DATE); inp(ws, f"D{r}", sal, MONEY, key=(i == 0)); inp(ws, f"F{r}", ben, MONEY); inp(ws, f"G{r}", over, MONEY)
    label(ws, f"A{tot}", "Team total", bold=True); label(ws, f"B{tot}", "Headcount", bold=True)
    fx(ws, f"C{tot}", f"=COUNTA(A{lo}:A{hi})", "0", bold=True)
    for col in "DFGHIJ": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    label(ws, f"A{tot+1}", "Average annual cost per person", bold=True); fx(ws, f"I{tot+1}", f'=IFERROR(ROUND(I{tot}/C{tot},2),0)', MONEY, bold=True)
    ws.freeze_panes = "B5"


def build_raises(ws):
    lo, hi = RAISE_ROWS; tot = hi + 1
    label(ws, "A1", "Raises", bold=True).font = TITLE
    label(ws, "A2", "Type a name from the Team sheet; the current salary is looked up. Choose percent or amount and enter the raise.", muted=True)
    head(ws, 4, ["Employee", "Current salary", "Raise type", "Raise (% or amount)", "Raise per year", "New salary", "Raise %", "Monthly difference", "Employer contrib. %", "Yearly cost to business"],
         [18, 14, 11, 14, 13, 14, 10, 13, 12, 15])
    dv = DataValidation(type="list", formula1='"percent,amount"', allow_blank=True); ws.add_data_validation(dv)
    t_lo, t_hi = TEAM_ROWS
    for r in range(lo, hi + 1):
        blank(ws, r, "ACD", {"D": "0.00"}); dv.add(f"C{r}")
        fx(ws, f"B{r}", f'=IF(A{r}="","",IFERROR(INDEX(Team!$D${t_lo}:$D${t_hi},MATCH(A{r},Team!$A${t_lo}:$A${t_hi},0)),""))', MONEY).font = BLUE
        fx(ws, f"I{r}", f'=IF(A{r}="","",{S}2)', "0.00").font = BLUE
        fx(ws, f"E{r}", f'=IF(A{r}="","",IF(C{r}="percent",ROUND(B{r}*D{r}/100,2),D{r}))', MONEY)
        fx(ws, f"F{r}", f'=IF(A{r}="","",B{r}+E{r})', MONEY)
        fx(ws, f"G{r}", f'=IF(A{r}="","",IFERROR(ROUND(E{r}/B{r}*100,2),""))', PCT)
        fx(ws, f"H{r}", f'=IF(A{r}="","",ROUND(E{r}/12,2))', MONEY)
        fx(ws, f"J{r}", f'=IF(A{r}="","",ROUND(E{r}*(1+I{r}/100),2))', MONEY)
    inp(ws, f"A{lo}", "Ana Example", key=True); inp(ws, f"C{lo}", "percent"); inp(ws, f"D{lo}", 5, "0.00", key=True)
    inp(ws, f"A{lo+1}", "Ben Example"); inp(ws, f"C{lo+1}", "amount"); inp(ws, f"D{lo+1}", 3000, "0.00")
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "EHJ": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    ws.freeze_panes = "B5"


def build_pto(ws):
    lo, hi = PTO_ROWS
    label(ws, "A1", "PTO", bold=True).font = TITLE
    label(ws, "A2", "One row per person. Entitlement and pay frequency default to Settings. Periods elapsed = pay periods already completed this year.", muted=True)
    head(ws, 4, ["Employee", "PTO per year (days)", "Pay frequency", "Periods per year", "Periods elapsed", "Accrual per period", "Accrued to date", "Carried over", "Taken", "Balance", "Still to accrue"],
         [18, 12, 13, 11, 11, 12, 12, 11, 9, 11, 12])
    dv = DataValidation(type="list", formula1='"weekly,biweekly,semimonthly,monthly"', allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        blank(ws, r, "AEHI", {"E": "0", "H": "0.00", "I": "0.00"}); dv.add(f"C{r}")
        fx(ws, f"B{r}", f'=IF(A{r}="","",{S}8)', "0.00").font = BLUE
        fx(ws, f"C{r}", f'=IF(A{r}="","",{S}9)').font = BLUE
        fx(ws, f"D{r}", f'=IF(A{r}="","",IF(C{r}="weekly",52,IF(C{r}="biweekly",26,IF(C{r}="semimonthly",24,IF(C{r}="monthly",12,"")))))', "0")
        fx(ws, f"F{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}/D{r},2),""))', "0.00")
        fx(ws, f"G{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}/D{r}*E{r},2),""))', "0.00")
        fx(ws, f"J{r}", f'=IF(A{r}="","",IFERROR(ROUND(H{r}+G{r}-I{r},2),""))', "0.00")
        fx(ws, f"K{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}-G{r},2),""))', "0.00")
    for i, (n, el, carry, taken) in enumerate([("Ana Example", 13, 3, 5), ("Ben Example", 13, 0, 2), ("Cara Example", 13, 0, 0)]):
        r = lo + i; inp(ws, f"A{r}", n, key=(i == 0)); inp(ws, f"E{r}", el, "0", key=(i == 0)); inp(ws, f"H{r}", carry, "0.00"); inp(ws, f"I{r}", taken, "0.00")
    ws.freeze_panes = "B5"


def build_hiring(ws):
    lo, hi = HIRE_ROWS; tot = hi + 1; t_lo, t_hi = TEAM_ROWS
    label(ws, "A1", "Hiring & Turnover", bold=True).font = TITLE
    label(ws, "A2", "Planned hires are prorated by the months left in the plan year (Settings). The turnover block uses your average cost per person from the Team sheet.", muted=True)
    head(ws, 4, ["Role", "Start month", "Annual salary", "Employer contrib. %", "Benefits + overheads / yr", "Total annual cost", "Months in plan year", "Cost in plan year"], [40, 12, 14, 12, 16, 15, 12, 15])
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCE", {"B": "mmm yyyy", "C": MONEY, "E": MONEY})
        fx(ws, f"D{r}", f'=IF(A{r}="","",{S}2)', "0.00").font = BLUE
        fx(ws, f"F{r}", f'=IF(A{r}="","",ROUND(C{r}*(1+D{r}/100)+E{r},2))', MONEY)
        fx(ws, f"G{r}", f'=IF(OR(A{r}="",B{r}=""),"",IF(YEAR(B{r})<{S}10,12,IF(YEAR(B{r})>{S}10,0,13-MONTH(B{r}))))', "0")
        fx(ws, f"H{r}", f'=IF(A{r}="","",IFERROR(ROUND(F{r}*G{r}/12,2),""))', MONEY)
    y = dt.date.today().year
    inp(ws, f"A{lo}", "Account manager", key=True); inp(ws, f"B{lo}", dt.date(y, 9, 1), "mmm yyyy", key=True); inp(ws, f"C{lo}", 55000, MONEY); inp(ws, f"E{lo}", 7000, MONEY)
    inp(ws, f"A{lo+1}", "Junior developer"); inp(ws, f"B{lo+1}", dt.date(y, 11, 1), "mmm yyyy"); inp(ws, f"C{lo+1}", 42000, MONEY); inp(ws, f"E{lo+1}", 6000, MONEY)
    label(ws, f"A{tot}", "Total planned hires", bold=True)
    for col in "FH": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    label(ws, "A20", "Turnover cost", bold=True).font = TITLE
    rows = [(21, "Average annual cost per employee", f"=Team!I{t_hi+2}", MONEY, True, "Defaults to the Team sheet average; overwrite for a specific role."),
            (22, "Working days per year", f"={S}6", "0", True, None),
            (23, "Recruiting and hiring cost per leaver", 5000, MONEY, False, "Ads, agency fees, interview time."),
            (24, "Days the role stays vacant", 30, "0", False, None),
            (25, "Onboarding cost per hire", 2000, MONEY, False, "Training, equipment, admin."),
            (26, "Ramp-up period (days)", 60, "0", False, None),
            (27, "Productivity during ramp-up (%)", f"={S}7", "0.0", True, None),
            (28, "Headcount", f"=Team!C{t_hi+1}", "0", True, None),
            (29, "Annual turnover rate (%)", 15, "0.00", False, "Leavers in a year ÷ average headcount.")]
    for r, name, val, fmt, default, note in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=not default)
        if note: label(ws, f"C{r}", note, muted=True)
    calc = [(30, "Cost per leaver", "=ROUND(B23+B21/B22*B24+B25+B21/B22*B26*(1-B27/100),2)", MONEY, "recruiting + vacancy + onboarding + ramp-up"),
            (31, "Expected leavers per year", "=ROUND(B28*B29/100,2)", "0.00", "headcount × turnover rate"),
            (32, "Turnover cost per year", "=ROUND(B30*B31,2)", MONEY, "per leaver × leavers"),
            (33, "Cost per leaver as % of annual cost", '=IFERROR(ROUND(B30/B21*100,2),"")', PCT, None),
            (34, "Daily cost of the role", '=IFERROR(ROUND(B21/B22,2),"")', MONEY, "annual cost ÷ working days")]
    for r, name, f, fmt, note in calc:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True)
        if note: label(ws, f"C{r}", note, muted=True)


README = """# Team Cost & Headcount Planner

Thank you for buying the planner. Open `team-cost-planner.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** your defaults for employer contributions, hours, weeks worked, productive share, working days, PTO and the plan year.
2. **Team:** one row per person. Total annual and monthly cost, cost per productive hour and the cost multiple are calculated, with team totals.
3. **Raises:** name a person, choose percent or amount, enter the raise. New salary, monthly difference and yearly cost to the business follow.
4. **PTO:** entitlement, pay frequency and periods completed give accrual, balance and days still to earn per person.
5. **Hiring & Turnover:** planned hires prorated for the plan year, and the cost of turnover per leaver and per year.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not legal, tax or HR advice; contribution rates and leave rules differ by country. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_team(wb.create_sheet("Team")); build_raises(wb.create_sheet("Raises")); build_pto(wb.create_sheet("PTO")); build_hiring(wb.create_sheet("Hiring & Turnover"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "team-cost-planner.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "team-cost-planner.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "team-cost-planner.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
