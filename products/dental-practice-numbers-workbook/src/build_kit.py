"""Build the Dental Practice Numbers Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "dental-practice-numbers-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"; RATIO = '0.00"x"'
MONTH_ROWS = (5, 16); PROVIDER_ROWS = (5, 24); NEW_PT_ROWS = (22, 33)
S = "Settings!$B$"  # settings cell prefix
PL = "'Monthly P&L'!"; PP = "'Provider Production'!"; HY = "Hygiene!"; CN = "'Chairs & New Patients'!"
# Settings rows (column B): referenced by every other sheet.
R_CHAIRS = 2; R_HOURS = 3; R_DAYS = 4; R_PROVIDERS = 5; R_HYGIENISTS = 6; R_TAXES = 7; R_OWNER = 8
R_TARGET_OH = 9; R_TARGET_HYG = 10; R_RETENTION = 11; R_YEARS = 12; R_FIRST_YEAR = 13
R_CHAIR_HRS = 14; R_PROV_HRS = 15; R_HYG_HRS = 16; R_OPEN_HRS = 17; R_LTV = 18
NOTE = "Knackdesk"


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, NOTE)
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def totals(ws, tot, lo, hi, cols):
    """Bold SUM of each (column, format) over rows lo..hi, written on the total row."""
    for col, fmt in cols:
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)


def lookup(col, month_ref):
    """INDEX/MATCH a Monthly P&L column on the month in month_ref."""
    lo, hi = MONTH_ROWS
    return f"INDEX({PL}${col}${lo}:${col}${hi},MATCH({month_ref},{PL}$A${lo}:$A${hi},0))"


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Dental Practice Numbers Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: chairs, opening hours and days, dentists and hygienists, payroll taxes and benefits, owner compensation, your target overhead % and hygiene margin %, and new-patient retention and first-year value. Available chair hours and provider hours follow and every other sheet reads them.",
        "2. Monthly P&L: one row per month with gross production, adjustments, collections, staff wages, rent and facilities, supplies and lab, marketing and other overhead. Net production, collection rate, loaded payroll, total overhead, overhead % of collections, profit before and after owner pay, profit margin and the gap to your target overhead are calculated, with totals.",
        "3. Provider Production: one row per dentist per month with hours, days, gross production and adjustments. Net production, production per hour and per day, overhead per provider hour and profit per hour are calculated.",
        "4. Hygiene: one row per month with hygiene production, hygienist wages, supplies, hygienist hours and your own overhead allocation. Loaded wages, total cost, department profit, margin %, production and cost per hour and the gap to your target margin are calculated, with totals.",
        "5. Chairs & New Patients: scheduled chair hours per month against available chair hours give utilization %, idle hours and the production left unrealised (production per hour defaults from the Monthly P&L; type your own to override). Below, marketing spend (defaults from the Monthly P&L) and new patients give cost per new patient, lifetime value, value-to-cost ratio and break-even new patients.",
        "6. Summary: year-to-date collections, net production, collection rate, overhead % against target, profit before and after owner pay, best and worst month, production per provider hour, hygiene margin, chair utilization and unrealised production, cost per new patient and months above your overhead target.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Production vs collections: gross production is the fee value of the dentistry done at your full fees; net production = gross production − adjustments; collections is the money actually received. Collection rate % = collections ÷ net production.",
        "Adjustments: the part of gross production you will never collect, such as insurance or plan write-offs, discounts, courtesy adjustments and refunds.",
        "Overhead percentage: total overhead (loaded payroll + rent and facilities + supplies and lab + marketing + other) ÷ collections. Owner compensation is not in overhead; it is shown separately as profit before owner pay − owner compensation = profit after owner pay.",
        "Production per hour: a provider's net production ÷ hours worked. Overhead per hour = the month's total overhead ÷ provider hours available (providers × hours open × days open). Profit per hour = production per hour − overhead per hour.",
        "Hygiene department profit: hygiene production − (hygienist wages loaded with payroll taxes and benefits + hygiene supplies + your own overhead allocation). Margin % = department profit ÷ hygiene production.",
        "Chair utilization: scheduled chair hours ÷ available chair hours (chairs × hours open × days open). Unrealised production = idle chair hours × production per scheduled hour: what the empty hours would have produced at your current rate.",
        "Cost per new patient: marketing spend ÷ new patients. Lifetime value = first-year value + first-year value × retention % × (years retained − 1). Break-even new patients = marketing spend ÷ first-year value.",
        "The targets in Settings are your own; this workbook does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [(R_CHAIRS, "Chairs (operatories)", 4, INT0, True, "Example default 4. Treatment chairs you can book, dentist and hygiene chairs together."),
            (R_HOURS, "Hours open per day", 8, HRS, True, "Example default 8. Clinical hours a chair can be booked each day you are open."),
            (R_DAYS, "Days open per month", 20, INT0, True, "Example default 20. Clinical days open in a typical month."),
            (R_PROVIDERS, "Providers (dentists) count", 2, INT0, True, "Example default 2. Dentists producing chairside, full-time equivalents."),
            (R_HYGIENISTS, "Hygienists count", 2, INT0, False, "Example default 2. Hygienists, full-time equivalents; gives hygienist hours available below."),
            (R_TAXES, "Payroll taxes and benefits %", 15, PCT, True, "Default 15. Employer payroll taxes, pension, health cover and other benefits on top of wages."),
            (R_OWNER, "Owner compensation per month", 12000, MONEY, True, "Example default 12,000 a month. What the owner dentist pays themselves; kept out of overhead and shown separately."),
            (R_TARGET_OH, "Target overhead % of collections", 60, PCT, True, "Default 60. A placeholder for your own target, not a recommendation or an industry benchmark."),
            (R_TARGET_HYG, "Target hygiene margin %", 30, PCT, False, "Default 30. A placeholder for your own target, not a recommendation or an industry benchmark."),
            (R_RETENTION, "New-patient retention %", 60, PCT, False, "Default 60. Share of new patients still coming each year after the first; use your own recall data."),
            (R_YEARS, "Years retained", 3, INT0, False, "Default 3. How many years a retained patient typically stays, including the first."),
            (R_FIRST_YEAR, "Average first-year value per new patient", 900, MONEY, True, "Example default 900. Average collections from a new patient in their first year.")]
    for r, name, val, fmt, key, note in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", note, muted=True)
    pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                         showErrorMessage=True, error="Type a percentage as a whole number between 0 and 100 (for example 15).")
    count = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                           showErrorMessage=True, error="Type a whole number of 0 or more.")
    ws.add_data_validation(pct); ws.add_data_validation(count)
    for r in (R_TAXES, R_TARGET_OH, R_TARGET_HYG, R_RETENTION): pct.add(f"B{r}")
    for r in (R_CHAIRS, R_DAYS, R_PROVIDERS, R_HYGIENISTS, R_YEARS): count.add(f"B{r}")
    ws[f"B{R_TAXES}"].comment = Comment("Assumption: 15% is a placeholder for employer taxes and benefits; use your payroll figures.", NOTE)
    ws[f"B{R_TARGET_OH}"].comment = Comment("Your own target, not an industry benchmark. Overhead excludes owner compensation.", NOTE)
    ws[f"B{R_TARGET_HYG}"].comment = Comment("Your own target, not an industry benchmark. Margin is department profit ÷ hygiene production.", NOTE)
    ws[f"B{R_RETENTION}"].comment = Comment("Assumption: retention is applied flat to each year after the first. Set it from your own recall and reactivation data.", NOTE)
    ws[f"B{R_FIRST_YEAR}"].comment = Comment("Assumption: later years are worth the same as the first year for a retained patient.", NOTE)
    hrs = f"B{R_HOURS}*B{R_DAYS}"
    derived = [(R_CHAIR_HRS, "Available chair hours per month", f"=ROUND(B{R_CHAIRS}*{hrs},2)", "0.00", "formula: chairs × hours open per day × days open per month"),
               (R_PROV_HRS, "Provider hours per month", f"=ROUND(B{R_PROVIDERS}*{hrs},2)", "0.00", "formula: providers × hours open per day × days open per month"),
               (R_HYG_HRS, "Hygienist hours available per month", f"=ROUND(B{R_HYGIENISTS}*{hrs},2)", "0.00", "formula: hygienists × hours open per day × days open per month"),
               (R_OPEN_HRS, "Clinical hours open per month", f"=ROUND({hrs},2)", "0.00", "formula: hours open per day × days open per month (one chair or one person)"),
               (R_LTV, "Lifetime value per new patient", f"=ROUND(B{R_FIRST_YEAR}+B{R_FIRST_YEAR}*B{R_RETENTION}/100*(B{R_YEARS}-1),2)", MONEY,
                "formula: first-year value + first-year value × retention % × (years retained − 1); collections, not profit")]
    rows_block(ws, derived, bold_rows=(R_CHAIR_HRS, R_PROV_HRS, R_HYG_HRS, R_LTV))


def build_pnl(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1
    label(ws, "A1", "Monthly P&L", bold=True).font = TITLE
    label(ws, "A2", "One row per month, filled top down. Owner compensation comes from Settings and is kept out of overhead.", muted=True)
    head(ws, 4, ["Month", "Gross production", "Adjustments", "Collections", "Staff wages", "Payroll loaded",
                 "Rent and facilities", "Supplies and lab", "Marketing", "Other overhead", "Net production",
                 "Collection rate %", "Total overhead", "Overhead % of collections", "Profit before owner",
                 "Profit after owner", "Profit margin after owner %", "Overhead % vs target"],
         [12, 13, 12, 13, 12, 12, 12, 12, 11, 12, 13, 11, 13, 12, 13, 13, 12, 11])
    fm = {"A": MONTH, "B": MONEY, "C": MONEY, "D": MONEY, "E": MONEY, "G": MONEY, "H": MONEY, "I": MONEY, "J": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEGHIJ", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"F{r}", g.format(f"ROUND(E{r}*(1+{S}{R_TAXES}/100),2)"), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND(B{r}-C{r},2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"IFERROR(ROUND(D{r}/K{r}*100,2),0)"), PCT)
        fx(ws, f"M{r}", g.format(f"ROUND(F{r}+G{r}+H{r}+I{r}+J{r},2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(M{r}/D{r}*100,2),0)"), PCT)
        fx(ws, f"O{r}", g.format(f"ROUND(D{r}-M{r},2)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(O{r}-{S}{R_OWNER},2)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(P{r}/D{r}*100,2),0)"), PCT)
        fx(ws, f"R{r}", g.format(f"ROUND(N{r}-{S}{R_TARGET_OH},2)"), PCT)
    year = dt.date.today().year
    sample = [(120000, 18000, 98000, 26000, 7500, 7200, 3000, 4200), (112000, 16500, 92000, 26000, 7500, 6800, 3500, 4000),
              (128000, 19000, 104000, 27500, 7500, 7900, 3000, 4300), (105000, 15800, 84500, 27500, 7500, 6900, 4500, 4600),
              (124000, 18500, 101000, 27500, 7500, 7400, 3200, 4100), (131000, 19500, 108000, 28000, 7500, 8100, 3000, 4400)]
    for i, vals in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        for col, val in zip("BCDEGHIJ", vals): inp(ws, f"{col}{r}", val, MONEY, key=(i == 0 and col in "BD"))
    ws[f"C{lo}"].comment = Comment("Write-offs, plan or insurance fee reductions, discounts and refunds against gross production.", NOTE)
    ws[f"E{lo}"].comment = Comment("Gross wages for all employed staff (associates, hygienists, nurses, front desk). Not the owner's pay.", NOTE)
    ws[f"F{lo}"].comment = Comment("Assumption: payroll taxes and benefits are a flat % of wages from Settings.", NOTE)
    label(ws, f"A{tot}", "Total", bold=True)
    totals(ws, tot, lo, hi, [(c, MONEY) for c in "BCDEFGHIJKMOP"])
    fx(ws, f"L{tot}", f"=IFERROR(ROUND(D{tot}/K{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(M{tot}/D{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"Q{tot}", f"=IFERROR(ROUND(P{tot}/D{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"R{tot}", f"=ROUND(N{tot}-{S}{R_TARGET_OH},2)", PCT, bold=True)
    ws[f"N{tot}"].comment = Comment("Totals recompute from the column totals, not an average of the monthly percentages.", NOTE)
    label(ws, f"A{tot + 1}", "Overhead % vs target: positive = overhead above your own target in Settings; negative = below it.", muted=True)
    ws.freeze_panes = "B5"


def build_providers(ws):
    lo, hi = PROVIDER_ROWS; tot = hi + 1
    label(ws, "A1", "Provider Production", bold=True).font = TITLE
    label(ws, "A2", "One row per dentist per month. Overhead per hour = that month's total overhead on the Monthly P&L ÷ provider hours per month in Settings.", muted=True)
    head(ws, 4, ["Month", "Provider", "Hours worked", "Days worked", "Gross production", "Adjustments", "Net production",
                 "Production per hour", "Production per day", "Overhead per provider hour", "Profit per hour"],
         [12, 18, 10, 10, 13, 12, 13, 12, 12, 13, 12])
    fm = {"A": MONTH, "C": HRS, "D": INT0, "E": MONEY, "F": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEF", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"ROUND(E{r}-F{r},2)"), MONEY)
        fx(ws, f"H{r}", g.format(f"IFERROR(ROUND(G{r}/C{r},2),0)"), MONEY)
        fx(ws, f"I{r}", g.format(f"IFERROR(ROUND(G{r}/D{r},2),0)"), MONEY)
        fx(ws, f"J{r}", g.format(f"IFERROR(ROUND({lookup('M', f'A{r}')}/{S}{R_PROV_HRS},2),0)"), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND(H{r}-J{r},2)"), MONEY)
    year = dt.date.today().year
    sample = [(1, "Dr. Patel", 150, 19, 62000, 9000), (1, "Dr. Moreno", 140, 18, 40000, 6500),
              (2, "Dr. Patel", 148, 19, 58000, 8500), (2, "Dr. Moreno", 136, 18, 37000, 6000),
              (3, "Dr. Patel", 156, 20, 66000, 9500), (3, "Dr. Moreno", 144, 19, 42000, 6800),
              (4, "Dr. Patel", 142, 18, 54000, 8200), (4, "Dr. Moreno", 130, 17, 34000, 5600)]
    for i, (m, name, hrs, days, gross, adj) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, m, 1), MONTH); inp(ws, f"B{r}", name)
        inp(ws, f"C{r}", hrs, HRS, key=True); inp(ws, f"D{r}", days, INT0)
        inp(ws, f"E{r}", gross, MONEY, key=True); inp(ws, f"F{r}", adj, MONEY)
    ws[f"C{lo}"].comment = Comment("Chairside hours actually worked this month, from your schedule or practice management system.", NOTE)
    ws[f"J{lo}"].comment = Comment("Assumption: all practice overhead is carried by provider hours available (Settings), not hours worked. The month must also appear on the Monthly P&L, or this shows 0.", NOTE)
    label(ws, f"A{tot}", "Total", bold=True)
    totals(ws, tot, lo, hi, [("C", HRS), ("D", INT0), ("E", MONEY), ("F", MONEY), ("G", MONEY)])
    fx(ws, f"H{tot}", f"=IFERROR(ROUND(G{tot}/C{tot},2),0)", MONEY, bold=True)
    fx(ws, f"I{tot}", f"=IFERROR(ROUND(G{tot}/D{tot},2),0)", MONEY, bold=True)
    ws[f"H{tot}"].comment = Comment("Average production per provider hour = total net production ÷ total hours worked.", NOTE)
    ws.freeze_panes = "C5"


def build_hygiene(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1
    label(ws, "A1", "Hygiene", bold=True).font = TITLE
    label(ws, "A2", "One row per month. Wages are loaded with payroll taxes and benefits from Settings; the overhead allocation is your own share for the hygiene department.", muted=True)
    head(ws, 4, ["Month", "Hygiene production", "Hygienist wages", "Hygiene supplies", "Hygienist hours", "Overhead allocation",
                 "Loaded wages", "Total cost", "Department profit", "Margin %", "Production per hour", "Cost per hour", "Margin % vs target"],
         [12, 13, 12, 12, 11, 12, 12, 12, 13, 10, 12, 11, 11])
    fm = {"A": MONTH, "B": MONEY, "C": MONEY, "D": MONEY, "E": HRS, "F": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEF", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"ROUND(C{r}*(1+{S}{R_TAXES}/100),2)"), MONEY)
        fx(ws, f"H{r}", g.format(f"ROUND(G{r}+D{r}+F{r},2)"), MONEY)
        fx(ws, f"I{r}", g.format(f"ROUND(B{r}-H{r},2)"), MONEY)
        fx(ws, f"J{r}", g.format(f"IFERROR(ROUND(I{r}/B{r}*100,2),0)"), PCT)
        fx(ws, f"K{r}", g.format(f"IFERROR(ROUND(B{r}/E{r},2),0)"), MONEY)
        fx(ws, f"L{r}", g.format(f"IFERROR(ROUND(H{r}/E{r},2),0)"), MONEY)
        fx(ws, f"M{r}", g.format(f"ROUND(J{r}-{S}{R_TARGET_HYG},2)"), PCT)
    year = dt.date.today().year
    sample = [(18000, 7800, 900, 300, 2500), (17000, 7800, 850, 290, 2500), (20000, 8200, 950, 310, 2500),
              (17000, 8200, 900, 300, 2500), (19500, 8200, 950, 305, 2500), (21000, 8400, 1000, 320, 2500)]
    for i, vals in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        for col, val in zip("BCDEF", vals): inp(ws, f"{col}{r}", val, fm[col], key=(i == 0 and col in "BC"))
    ws[f"F{lo}"].comment = Comment("Your own share of rent, front desk and other overhead for the hygiene chairs (for example by chair count or hours). Example 2,500 only.", NOTE)
    ws[f"E{lo}"].comment = Comment("Hygienist hours actually worked this month. Compare with hygienist hours available in Settings.", NOTE)
    label(ws, f"A{tot}", "Total", bold=True)
    totals(ws, tot, lo, hi, [("B", MONEY), ("C", MONEY), ("D", MONEY), ("E", HRS), ("F", MONEY), ("G", MONEY), ("H", MONEY), ("I", MONEY)])
    fx(ws, f"J{tot}", f"=IFERROR(ROUND(I{tot}/B{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"K{tot}", f"=IFERROR(ROUND(B{tot}/E{tot},2),0)", MONEY, bold=True)
    fx(ws, f"L{tot}", f"=IFERROR(ROUND(H{tot}/E{tot},2),0)", MONEY, bold=True)
    fx(ws, f"M{tot}", f"=ROUND(J{tot}-{S}{R_TARGET_HYG},2)", PCT, bold=True)
    ws[f"J{tot}"].comment = Comment("Totals recompute from the column totals, not an average of the monthly margins.", NOTE)
    ws.freeze_panes = "B5"


def build_chairs(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1
    label(ws, "A1", "Chairs & New Patients", bold=True).font = TITLE
    label(ws, "A2", "Chair utilization: production per scheduled hour defaults to that month's net production on the Monthly P&L ÷ scheduled hours; type your own in column D to override it.", muted=True)
    head(ws, 4, ["Month", "Scheduled chair hours", "Production per scheduled hour (Monthly P&L)", "or type your own",
                 "Production per hour used", "Available chair hours", "Utilization %", "Idle chair hours",
                 "Production at full chairs", "Unrealised production"],
         [12, 12, 16, 12, 13, 12, 11, 11, 14, 14])
    fm = {"A": MONTH, "B": HRS, "D": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABD", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"C{r}", g.format(f"IFERROR(ROUND({lookup('K', f'A{r}')}/B{r},2),0)"), MONEY)
        fx(ws, f"E{r}", g.format(f'IF(D{r}<>"",D{r},C{r})'), MONEY)
        fx(ws, f"F{r}", g.format(f"{S}{R_CHAIR_HRS}"), "0.00")
        fx(ws, f"G{r}", g.format(f"IFERROR(ROUND(B{r}/F{r}*100,2),0)"), PCT)
        fx(ws, f"H{r}", g.format(f"ROUND(F{r}-B{r},2)"), "0.00;-0.00;0.00")
        fx(ws, f"I{r}", g.format(f"ROUND(F{r}*E{r},2)"), MONEY)
        fx(ws, f"J{r}", g.format(f"ROUND(H{r}*E{r},2)"), MONEY)
    year = dt.date.today().year
    for i, (sched, override) in enumerate([(470, None), (455, None), (500, None), (430, 230), (480, None), (510, None)]):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH); inp(ws, f"B{r}", sched, HRS, key=(i == 0))
        if override: inp(ws, f"D{r}", override, MONEY)
    ws[f"B{lo}"].comment = Comment("Chair hours with a patient booked, all chairs together, from your schedule.", NOTE)
    ws[f"D{lo + 3}"].comment = Comment("Example override: your own production per scheduled hour. Leave empty to use the Monthly P&L figure.", NOTE)
    ws[f"J{lo}"].comment = Comment("Assumption: idle hours would produce at the same rate per hour as the booked hours.", NOTE)
    label(ws, f"A{tot}", "Total", bold=True)
    totals(ws, tot, lo, hi, [("B", HRS), ("F", "0.00"), ("H", "0.00;-0.00;0.00"), ("I", MONEY), ("J", MONEY)])
    fx(ws, f"G{tot}", f"=IFERROR(ROUND(B{tot}/F{tot}*100,2),0)", PCT, bold=True)
    ws[f"G{tot}"].comment = Comment("Utilization recomputed from the totals: total scheduled ÷ total available hours.", NOTE)

    nlo, nhi = NEW_PT_ROWS; ntot = nhi + 1
    label(ws, f"A{nlo - 3}", "New patients", bold=True)
    label(ws, f"A{nlo - 2}", "Marketing spend defaults to that month's marketing on the Monthly P&L; type your own in column C to override it. Lifetime value uses retention, years and first-year value from Settings.", muted=True)
    head(ws, nlo - 1, ["Month", "Marketing spend (Monthly P&L)", "or type your own", "Marketing spend used", "New patients",
                       "Cost per new patient", "Lifetime value per patient", "Value-to-cost ratio", "Net value per patient",
                       "Break-even new patients"])
    fm = {"A": MONTH, "C": MONEY, "E": INT0}
    fy = f"{S}{R_FIRST_YEAR}"
    for r in range(nlo, nhi + 1):
        blank(ws, r, "ACE", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"B{r}", g.format(f"IFERROR({lookup('I', f'A{r}')},0)"), MONEY)
        fx(ws, f"D{r}", g.format(f'IF(C{r}<>"",C{r},B{r})'), MONEY)
        fx(ws, f"F{r}", g.format(f"IFERROR(ROUND(D{r}/E{r},2),0)"), MONEY)
        fx(ws, f"G{r}", g.format(f"ROUND({fy}+{fy}*{S}{R_RETENTION}/100*({S}{R_YEARS}-1),2)"), MONEY)
        fx(ws, f"H{r}", g.format(f"IFERROR(ROUND(G{r}/F{r},2),0)"), RATIO)
        fx(ws, f"I{r}", g.format(f"ROUND(G{r}-F{r},2)"), MONEY)
        fx(ws, f"J{r}", g.format(f"IFERROR(ROUND(D{r}/{fy},2),0)"), "0.00")
    for i, (override, patients) in enumerate([(None, 40), (4200, 38), (None, 45), (None, 52), (None, 42), (None, 47)]):
        r = nlo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH); inp(ws, f"E{r}", patients, INT0, key=(i == 0))
        if override: inp(ws, f"C{r}", override, MONEY)
    ws[f"C{nlo + 1}"].comment = Comment("Example override: spend that sits outside the P&L marketing line, such as referral gifts. Leave empty to use the Monthly P&L figure.", NOTE)
    ws[f"G{nlo}"].comment = Comment("Assumption: lifetime value = first-year value + first-year value × retention % × (years retained − 1), all from Settings. Collections, not profit.", NOTE)
    ws[f"J{nlo}"].comment = Comment("New patients needed for their first-year value to repay the month's marketing spend.", NOTE)
    label(ws, f"A{ntot}", "Total", bold=True)
    totals(ws, ntot, nlo, nhi, [("D", MONEY), ("E", INT0)])
    fx(ws, f"F{ntot}", f"=IFERROR(ROUND(D{ntot}/E{ntot},2),0)", MONEY, bold=True)
    fx(ws, f"J{ntot}", f"=IFERROR(ROUND(D{ntot}/{fy},2),0)", "0.00", bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1; ptot = PROVIDER_ROWS[1] + 1; ntot = NEW_PT_ROWS[1] + 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 64
    label(ws, "A1", "Summary", bold=True).font = TITLE
    m = lambda col: f"{PL}${col}${lo}:${col}${hi}"
    rows = [(2, "Year-to-date collections", f"={PL}D{tot}", MONEY, "Monthly P&L totals"),
            (3, "Year-to-date net production", f"={PL}K{tot}", MONEY, "gross production − adjustments"),
            (4, "Collection rate %", f"={PL}L{tot}", PCT, "collections ÷ net production"),
            (5, "Overhead % of collections", f"={PL}N{tot}", PCT, "total overhead ÷ collections, owner pay excluded"),
            (6, "Your target overhead %", f"={S}{R_TARGET_OH}", PCT, "from Settings; your own target, not a benchmark"),
            (7, "Overhead % vs target", f"={PL}R{tot}", PCT, "positive = above your target"),
            (8, "Profit before owner pay", f"={PL}O{tot}", MONEY, "collections − overhead"),
            (9, "Profit after owner pay", f"={PL}P{tot}", MONEY, "profit before owner pay − owner compensation"),
            (10, "Best month by profit", f'=IFERROR(INDEX({m("A")},MATCH(MAX({m("P")}),{m("P")},0)),"")', MONTH, "highest profit after owner pay"),
            (11, "Worst month by profit", f'=IFERROR(INDEX({m("A")},MATCH(MIN({m("P")}),{m("P")},0)),"")', MONTH, "lowest profit after owner pay"),
            (12, "Average production per provider hour", f"={PP}H{ptot}", MONEY, "Provider Production: total net production ÷ total hours"),
            (13, "Hygiene year-to-date margin %", f"={HY}J{tot}", PCT, "department profit ÷ hygiene production"),
            (14, "Hygiene margin % vs target", f"={HY}M{tot}", PCT, "positive = above your target"),
            (15, "Average chair utilization %", f"={CN}G{tot}", PCT, "total scheduled ÷ total available chair hours"),
            (16, "Total unrealised production", f"={CN}J{tot}", MONEY, "idle chair hours × production per hour"),
            (17, "Year-to-date cost per new patient", f"={CN}F{ntot}", MONEY, "total marketing spend ÷ total new patients"),
            (18, "Total new patients", f"={CN}E{ntot}", INT0, None),
            (19, "Months with overhead above target", f'=COUNTIF({m("R")},">0")', INT0, "overhead % above your target in Settings"),
            (20, "Lifetime value per new patient", f"={S}{R_LTV}", MONEY, "from Settings"),
            (21, "Hygienist hours available per month", f"={S}{R_HYG_HRS}", "0.00", "from Settings")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))


README = """# Dental Practice Numbers Workbook

Thank you for buying the workbook. Open `dental-practice-numbers-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** chairs, hours and days open, dentists and hygienists, payroll taxes and benefits, owner compensation, your target overhead % and hygiene margin %, new-patient retention, years retained and first-year value; available chair hours and provider hours follow.
2. **Monthly P&L:** gross production, adjustments, collections, wages and overhead lines per month; net production, collection rate, loaded payroll, total overhead, overhead % of collections, profit before and after owner pay, profit margin and overhead vs your target follow.
3. **Provider Production:** hours, days, gross production and adjustments per dentist per month; net production, production per hour and per day, overhead per provider hour and profit per hour follow.
4. **Hygiene:** hygiene production, hygienist wages, supplies, hours and your own overhead allocation per month; loaded wages, total cost, department profit, margin %, production and cost per hour and margin vs target follow.
5. **Chairs & New Patients:** scheduled chair hours per month give utilization %, idle hours and unrealised production; marketing spend and new patients give cost per new patient, lifetime value, value-to-cost ratio and break-even new patients.
6. **Summary:** year-to-date collections, net production, collection rate, overhead % against target, profit before and after owner pay, best and worst month, production per provider hour, hygiene margin, chair utilization, cost per new patient and months above target.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_pnl(wb.create_sheet("Monthly P&L"))
    build_providers(wb.create_sheet("Provider Production")); build_hygiene(wb.create_sheet("Hygiene"))
    build_chairs(wb.create_sheet("Chairs & New Patients")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
