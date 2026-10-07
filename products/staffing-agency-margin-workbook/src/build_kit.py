"""Build the Staffing Agency Margin Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "staffing-agency-margin-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"
RATIO = "0.00"
PERM_ROWS = (5, 34); TEMP_ROWS = (5, 34); DESK_ROWS = (5, 16); RES_ROWS = (5, 24)
S = "Settings!$B$"  # settings cell prefix
YES_NO = '"yes,no"'
OUTCOMES = '"in guarantee,retained,fell off"'
PP = "'Perm Placements'!"; TA = "'Temp Assignments'!"; DP = "'Desk Performance'!"; RR = "'Rebate Reserve'!"
# Settings rows (column B): referenced by every other sheet.
R_COUNT = 2; R_SALARY = 3; R_BENEFITS = 4; R_TOOLS = 5; R_OVERHEAD = 6; R_HOURS = 7; R_FEE = 8; R_ONCOST = 9
R_MARKUP = 10; R_GUARANTEE = 11; R_FALLOFF = 12; R_REBATE = 13; R_DEF_FEE = 14
R_LOADED = 15; R_DESK = 16; R_DESK_EACH = 17; R_COST_HOUR = 18; R_RESERVE = 19; R_RESERVE_DEF = 20


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def note(ws, ref, text):
    ws[ref].comment = Comment(text, "Knackdesk")


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Staffing Agency Margin Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: number of recruiters, average recruiter salary, benefits and payroll taxes %, tools and job boards, office and other overhead, working hours per recruiter, and your default permanent fee %, temp on-costs %, temp markup %, guarantee days, expected fall-off % and rebate %. Desk cost per recruiter, recruiter cost per hour and the reserve % of fee follow and every other sheet reads them.",
        "2. Perm Placements: one row per permanent placement with date, client, role, annual salary, fee % (default from Settings, type your own to override), split %, recruiter hours, advertising cost, invoiced and collected. Gross fee, split fee, net fee, delivery cost, margin, margin %, net fee per recruiter hour, expected rebate reserve and net fee after reserve are calculated, with totals.",
        "3. Temp Assignments: one row per temp or contractor with client, worker, pay rate, on-costs % and markup % (defaults from Settings, type your own to override), an optional bill rate you have agreed, hours per week and weeks. Bill rate used, loaded cost per hour, gross margin per hour and %, and weekly and assignment revenue and margin are calculated, with totals.",
        "4. Desk Performance: one row per month with job orders received and filled, candidates submitted, interviews, total days to fill and temp margin earned. Perm placements and net perm fees for the month are counted from the Perm Placements sheet by date (type your own to override). Gross margin, desk cost for the month, desk profit, fill rate, average days to fill, submittals per fill, interview to fill % and margin per recruiter are calculated.",
        "5. Rebate Reserve: one row per placement still inside or recently out of its guarantee, with start date, net fee, guarantee days, fall-off % and rebate % (defaults from Settings, type your own to override) and the outcome. Guarantee end date, expected reserve, reserve released, rebate paid and reserve still held are calculated, with totals.",
        "6. Summary: recruiters, desk cost per recruiter, recruiter cost per hour, year-to-date perm net fees, temp margin, gross margin, desk cost and desk profit, break-even placements per recruiter, placements below your margin % threshold, the thinnest temp assignment, year-to-date fill rate and days to fill, and the reserve position.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Placement: a candidate you introduced who starts a permanent job with your client, for which the client pays a fee.",
        "Job order: a vacancy a client has asked you to fill (also called a requisition or a req). Filled = a candidate you placed started in it.",
        "Net fee after split: the fee your agency keeps after paying any share to a split partner (another agency or recruiter who supplied the candidate or the client). Net fee = gross fee × (1 − split %).",
        "Markup on pay vs margin on bill: markup is added on top of the pay rate (bill rate = pay rate × (1 + markup %)). Gross margin % is profit as a share of the bill rate after on-costs. A 55% markup on pay with 18% on-costs is a 23.87% margin on bill, not 55%.",
        "On-costs: what a temp worker costs you on top of their pay rate: employer payroll taxes, pension, holiday pay, workers' compensation or liability insurance, and payroll fees. Loaded cost per hour = pay rate × (1 + on-costs %).",
        "Desk cost: what it costs to run a recruiter's desk for a year: loaded salary plus an equal share of tools, job boards, office and other overhead. Recruiter cost per hour = desk cost per recruiter ÷ working hours per recruiter per year.",
        "Fill rate: job orders filled ÷ job orders received for the same period, as a %. The year-to-date fill rate is computed from the totals, not averaged across months.",
        "Fall-off and rebate reserve: a fall-off is a placed candidate who leaves or is let go inside the guarantee period, so you owe the client a rebate (a refund or a free replacement). The reserve sets aside, for each placement, net fee × rebate % × expected fall-off % until the guarantee ends.",
        "The targets and thresholds in this workbook are your own; it does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [(R_COUNT, "Number of recruiters", 4, INT0, True, "Example default 4. Fee-earning recruiters whose desks carry the agency's cost; owners who bill count too."),
            (R_SALARY, "Average recruiter salary (annual)", 55000, MONEY, True, "Example default 55,000 a year. Base salary plus the commission you expect to pay, averaged across recruiters."),
            (R_BENEFITS, "Benefits and payroll taxes % for staff", 20, PCT, False, "Default 20. Employer payroll taxes, pension, health cover and other benefits on top of recruiter salary."),
            (R_TOOLS, "Tools and job boards (annual)", 36000, MONEY, True, "Example default 36,000 a year. Applicant tracking system, CRM, job board subscriptions, LinkedIn seats, background checks."),
            (R_OVERHEAD, "Office and other overhead (annual)", 72000, MONEY, True, "Example default 72,000 a year. Rent, admin staff, insurance, accounting, marketing and everything else."),
            (R_HOURS, "Working hours per recruiter per year", 1800, INT0, True, "Default 1,800. Hours a recruiter is at work in a year, used to cost an hour of desk time."),
            (R_FEE, "Default permanent fee %", 20, PCT, True, "Default 20. Fee as a % of the candidate's first-year salary; each placement can override it."),
            (R_ONCOST, "Default on-costs % for temps", 18, PCT, False, "Default 18. Employer taxes, holiday pay, insurance and payroll fees on top of a temp's pay rate; each assignment can override it."),
            (R_MARKUP, "Default temp markup %", 55, PCT, True, "Default 55. Markup on the pay rate to get the bill rate; each assignment can override it or set its own bill rate."),
            (R_GUARANTEE, "Default guarantee days", 90, INT0, False, "Default 90. Days after the start date during which a fall-off triggers a rebate; each placement can override it."),
            (R_FALLOFF, "Expected fall-off %", 8, PCT, True, "Default 8. Share of placements you expect to fall off inside the guarantee; your own placeholder, set it from your history."),
            (R_REBATE, "Rebate % of fee", 100, PCT, False, "Default 100. Share of the fee refunded or credited on a fall-off (100 = full refund or free replacement)."),
            (R_DEF_FEE, "Default placement fee for the reserve", 12000, MONEY, False, "Example default 12,000. A typical net fee, used only to show the reserve per placement below.")]
    for r, name, val, fmt, key, txt in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", txt, muted=True)
    note(ws, f"B{R_SALARY}", "Assumption: one average salary for all recruiters. If commission is a big part of pay, include what you expect to pay out this year.")
    note(ws, f"B{R_BENEFITS}", "Assumption: 20% is a placeholder for employer taxes and benefits; use your payroll figures.")
    note(ws, f"B{R_HOURS}", "Assumption: 1,800 hours is a placeholder for a full-time year after holidays. Use your own.")
    note(ws, f"B{R_ONCOST}", "Assumption: 18% is a placeholder. On-costs depend on your country, state and the worker's contract; use your payroll provider's figures.")
    note(ws, f"B{R_MARKUP}", "Your own placeholder markup, not a recommendation or an industry benchmark.")
    note(ws, f"B{R_FALLOFF}", "Your own placeholder, not an industry benchmark. Count fall-offs inside the guarantee over your last 20 or more placements.")
    note(ws, f"B{R_REBATE}", "Check your terms of business: some guarantees refund on a sliding scale, others offer one free replacement.")
    derived = [(R_LOADED, "Loaded recruiter cost per year (each)", f"=ROUND(B{R_SALARY}*(1+B{R_BENEFITS}/100),2)", MONEY,
                "formula: average salary × (1 + benefits and payroll taxes %)"),
               (R_DESK, "Total desk cost per year", f"=ROUND(B{R_COUNT}*B{R_LOADED}+B{R_TOOLS}+B{R_OVERHEAD},2)", MONEY,
                "formula: recruiters × loaded cost + tools and job boards + office and other overhead"),
               (R_DESK_EACH, "Desk cost per recruiter (annual)", f"=IFERROR(ROUND(B{R_DESK}/B{R_COUNT},2),0)", MONEY,
                "formula: total desk cost ÷ number of recruiters"),
               (R_COST_HOUR, "Recruiter cost per hour", f"=IFERROR(ROUND(B{R_DESK_EACH}/B{R_HOURS},2),0)", MONEY,
                "formula: desk cost per recruiter ÷ working hours per recruiter per year"),
               (R_RESERVE, "Reserve % of fee", f"=ROUND(B{R_FALLOFF}*B{R_REBATE}/100,2)", PCT,
                "formula: expected fall-off % × rebate % ÷ 100; the share of each net fee to hold back until the guarantee ends"),
               (R_RESERVE_DEF, "Reserve per placement at the default fee", f"=ROUND(B{R_DEF_FEE}*B{R_RESERVE}/100,2)", MONEY,
                "formula: default placement fee × reserve % of fee")]
    rows_block(ws, derived, bold_rows=(R_DESK, R_DESK_EACH, R_COST_HOUR, R_RESERVE))


def build_perm(ws):
    lo, hi = PERM_ROWS; tot = hi + 1
    label(ws, "A1", "Perm Placements", bold=True).font = TITLE
    label(ws, "A2", "One row per permanent placement. Fee % defaults from Settings; type your own in 'or type your own fee %' to override it. Delivery cost uses the recruiter cost per hour from Settings.", muted=True)
    head(ws, 4, ["Date", "Client", "Role", "Annual salary", "Default fee % (Settings)", "or type your own fee %", "Fee % used",
                 "Split %", "Recruiter hours", "Advertising cost", "Invoiced", "Collected", "Gross fee", "Split fee", "Net fee",
                 "Delivery cost", "Margin on placement", "Margin %", "Net fee per recruiter hour", "Expected rebate reserve",
                 "Net fee after reserve"],
         [12, 20, 24, 13, 11, 11, 10, 9, 10, 12, 9, 12, 12, 12, 12, 12, 13, 10, 12, 12, 13])
    dv = DataValidation(type="list", formula1=YES_NO, allow_blank=True); ws.add_data_validation(dv)
    fm = {"A": DATE, "D": MONEY, "F": PCT, "H": PCT, "I": HRS, "J": MONEY, "L": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDFHIJKL", fm); dv.add(f"K{r}")
        g = f'=IF(B{r}="","",{{}})'
        fx(ws, f"E{r}", g.format(f"{S}{R_FEE}"), PCT)
        fx(ws, f"G{r}", g.format(f'IF(F{r}<>"",F{r},E{r})'), PCT)
        fx(ws, f"M{r}", g.format(f"ROUND(D{r}*G{r}/100,2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(M{r}*H{r}/100,2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"M{r}-N{r}"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(I{r}*{S}{R_COST_HOUR}+J{r},2)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"ROUND(O{r}-P{r},2)"), MONEY)
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(Q{r}/O{r}*100,2),0)"), PCT)
        fx(ws, f"S{r}", g.format(f"IFERROR(ROUND(O{r}/I{r},2),0)"), MONEY)
        fx(ws, f"T{r}", g.format(f"ROUND(O{r}*{S}{R_RESERVE}/100,2)"), MONEY)
        fx(ws, f"U{r}", g.format(f"ROUND(O{r}-T{r},2)"), MONEY)
    year = dt.date.today().year
    sample = [((1, 12), "Northwind Logistics", "Warehouse manager", 65000, None, 25, 40, 400, "yes", 13000),
              ((1, 26), "Brightline Dental", "Practice manager", 48000, None, 0, 35, 250, "yes", 9600),
              ((2, 9), "Acme Engineering", "Senior mechanical engineer", 85000, 22, 0, 60, 1200, "yes", 18700),
              ((3, 3), "Harbor Foods", "Finance assistant", 32000, 15, 0, 45, 300, "yes", 4800),
              ((4, 7), "Delta Software", "Product designer", 72000, None, 50, 30, 0, "yes", 14400),
              ((4, 22), "Okafor & Co", "Office administrator", 28000, None, 0, 50, 500, "yes", 5600),
              ((5, 14), "Northwind Logistics", "Transport planner", 52000, None, 0, 25, 200, "yes", 0),
              ((6, 2), "Garcia Retail", "Store manager", 45000, None, 0, 30, 350, "no", 0)]
    for i, ((m, d), client, role, salary, fee, split, hrs, adv, invoiced, collected) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, m, d), DATE); inp(ws, f"B{r}", client); inp(ws, f"C{r}", role)
        inp(ws, f"D{r}", salary, MONEY, key=True); inp(ws, f"F{r}", fee, PCT); inp(ws, f"H{r}", split, PCT)
        inp(ws, f"I{r}", hrs, HRS, key=True); inp(ws, f"J{r}", adv, MONEY); inp(ws, f"K{r}", invoiced); inp(ws, f"L{r}", collected, MONEY)
    note(ws, f"D{lo}", "The candidate's guaranteed first-year salary, as your fee is quoted on it.")
    note(ws, f"F{lo + 2}", "Example override: a retained search at a higher fee %.")
    note(ws, f"H{lo}", "Share of the gross fee paid to a split partner (another agency or recruiter). 0 if the placement was all yours.")
    note(ws, f"I{lo}", "Assumption: recruiter hours spent on this job order (intake, sourcing, screening, interviews, offer). Estimate if you do not track time.")
    note(ws, f"P{lo}", "Assumption: delivery cost = recruiter hours × recruiter cost per hour (Settings) + advertising. Desk overhead is inside the hourly cost.")
    note(ws, f"S{lo}", "Net fee ÷ recruiter hours: what each hour on this placement brought in, before costs.")
    note(ws, f"T{lo}", "Net fee × reserve % of fee (Settings). Hold this back until the guarantee ends; see the Rebate Reserve sheet.")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(B{lo}:B{hi})", '0" placements"', bold=True)
    for col, fmt in (("D", MONEY), ("I", HRS), ("J", MONEY), ("L", MONEY), ("M", MONEY), ("N", MONEY), ("O", MONEY),
                     ("P", MONEY), ("Q", MONEY), ("T", MONEY), ("U", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"R{tot}", f"=IFERROR(ROUND(Q{tot}/O{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"S{tot}", f"=IFERROR(ROUND(O{tot}/I{tot},2),0)", MONEY, bold=True)
    note(ws, f"R{tot}", "Totals recompute from the column totals, not an average of the placement rows.")
    ws.freeze_panes = "C5"


def build_temp(ws):
    lo, hi = TEMP_ROWS; tot = hi + 1
    label(ws, "A1", "Temp Assignments", bold=True).font = TITLE
    label(ws, "A2", "One row per temp or contractor assignment. On-costs % and markup % default from Settings; type your own to override them. Leave 'Bill rate override' empty to bill at pay rate × (1 + markup %).", muted=True)
    head(ws, 4, ["Client", "Worker", "Pay rate", "Default on-costs % (Settings)", "or type your own on-costs %",
                 "On-costs % used", "Default markup % (Settings)", "or type your own markup %", "Markup % used",
                 "Bill rate override", "Hours per week", "Weeks", "Bill rate used", "Loaded cost per hour",
                 "Gross margin per hour", "Gross margin %", "Bill per week", "Margin per week",
                 "Margin for assignment", "Revenue for assignment"],
         [20, 18, 10, 11, 11, 10, 11, 11, 10, 11, 10, 8, 11, 11, 11, 10, 12, 12, 13, 13])
    fm = {"C": MONEY, "E": PCT, "H": PCT, "J": MONEY, "K": HRS, "L": HRS}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEHJKL", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_ONCOST}"), PCT)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), PCT)
        fx(ws, f"G{r}", g.format(f"{S}{R_MARKUP}"), PCT)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), PCT)
        fx(ws, f"M{r}", g.format(f"IF(N(J{r})>0,J{r},ROUND(C{r}*(1+I{r}/100),2))"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(C{r}*(1+F{r}/100),2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"ROUND(M{r}-N{r},2)"), MONEY)
        fx(ws, f"P{r}", g.format(f"IFERROR(ROUND(O{r}/M{r}*100,2),0)"), PCT)
        fx(ws, f"Q{r}", g.format(f"ROUND(M{r}*K{r},2)"), MONEY)
        fx(ws, f"R{r}", g.format(f"ROUND(O{r}*K{r},2)"), MONEY)
        fx(ws, f"S{r}", g.format(f"ROUND(R{r}*L{r},2)"), MONEY)
        fx(ws, f"T{r}", g.format(f"ROUND(Q{r}*L{r},2)"), MONEY)
    sample = [("Northwind Logistics", "Warehouse operative A", 20, None, None, None, 40, 12),
              ("Northwind Logistics", "Forklift driver B", 22, None, None, None, 40, 26),
              ("Brightline Dental", "Dental receptionist C", 19, None, None, None, 30, 8),
              ("Acme Engineering", "CAD technician D", 30, None, None, 44, 37.5, 20),
              ("Harbor Foods", "Production line E", 18, None, None, None, 40, 16),
              ("Delta Software", "QA contractor F", 45, 12, 40, None, 37.5, 24),
              ("Garcia Retail", "Seasonal sales G", 25, None, 25, None, 35, 10),
              ("Okafor & Co", "Payroll clerk H", 24, None, None, None, 37.5, 6)]
    for i, (client, worker, pay, oncost, markup, bill, hpw, weeks) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", client); inp(ws, f"B{r}", worker); inp(ws, f"C{r}", pay, MONEY, key=True)
        inp(ws, f"E{r}", oncost, PCT); inp(ws, f"H{r}", markup, PCT); inp(ws, f"J{r}", bill, MONEY)
        inp(ws, f"K{r}", hpw, HRS); inp(ws, f"L{r}", weeks, HRS)
    note(ws, f"C{lo}", "The hourly rate the worker is paid, before employer on-costs.")
    note(ws, f"J{lo + 3}", "Example override: the client agreed a fixed bill rate, so the markup % is not used for this row.")
    note(ws, f"E{lo + 5}", "Example override: a contractor paid through their own company carries lower on-costs. Check the rules where you operate.")
    note(ws, f"H{lo + 6}", "Example override: a seasonal deal at a low markup. This is the thin-margin row the Summary flags.")
    note(ws, f"P{lo}", "Gross margin % = margin per hour ÷ bill rate, not the markup on pay.")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" assignments"', bold=True)
    for col in "QRST":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"P{tot}", f"=IFERROR(ROUND(R{tot}/Q{tot}*100,2),0)", PCT, bold=True)
    note(ws, f"P{tot}", "Weekly margin total ÷ weekly bill total, not an average of the rows.")
    ws.freeze_panes = "C5"


def build_desk(ws):
    lo, hi = DESK_ROWS; tot = hi + 1; plo, phi = PERM_ROWS
    label(ws, "A1", "Desk Performance", bold=True).font = TITLE
    label(ws, "A2", "One row per month, filled top down, with the first day of the month in column A. Perm placements and net perm fees are counted from Perm Placements by date; type your own to override them. Temp margin is typed in because assignments have no dates.", muted=True)
    head(ws, 4, ["Month", "Perm placements (from Perm Placements)", "or type your own placements", "Perm placements used",
                 "Net perm fees (from Perm Placements)", "or type your own net fees", "Net perm fees used",
                 "Temp margin earned", "Job orders received", "Job orders filled", "Candidates submitted", "Interviews",
                 "Total days to fill", "Total gross margin", "Desk cost for the month", "Desk profit", "Fill rate %",
                 "Average days to fill", "Submittals per fill", "Interview to fill %", "Margin per recruiter"],
         [12, 13, 12, 11, 13, 12, 12, 12, 10, 10, 11, 10, 10, 13, 12, 12, 10, 10, 10, 10, 12])
    fm = {"A": MONTH, "C": INT0, "F": MONEY, "H": MONEY, "I": INT0, "J": INT0, "K": INT0, "L": INT0, "M": INT0}
    dates = f"{PP}$A${plo}:$A${phi}"
    for r in range(lo, hi + 1):
        blank(ws, r, "ACFHIJKLM", fm)
        g = f'=IF(A{r}="","",{{}})'
        nxt = f'"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1)'
        fx(ws, f"B{r}", g.format(f'COUNTIFS({dates},">="&A{r},{dates},{nxt})'), INT0)
        fx(ws, f"D{r}", g.format(f'IF(C{r}<>"",C{r},B{r})'), INT0)
        fx(ws, f"E{r}", g.format(f'ROUND(SUMIFS({PP}$O${plo}:$O${phi},{dates},">="&A{r},{dates},{nxt}),2)'), MONEY)
        fx(ws, f"G{r}", g.format(f'IF(F{r}<>"",F{r},E{r})'), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(G{r}+H{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"ROUND({S}{R_DESK}/12,2)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(N{r}-O{r},2)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(J{r}/I{r}*100,2),0)"), PCT)
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(M{r}/J{r},2),0)"), RATIO)
        fx(ws, f"S{r}", g.format(f"IFERROR(ROUND(K{r}/J{r},2),0)"), RATIO)
        fx(ws, f"T{r}", g.format(f"IFERROR(ROUND(J{r}/L{r}*100,2),0)"), PCT)
        fx(ws, f"U{r}", g.format(f"IFERROR(ROUND(N{r}/{S}{R_COUNT},2),0)"), MONEY)
    year = dt.date.today().year
    sample = [(14000, 12, 5, 30, 14, 160), (15500, 10, 4, 26, 11, 140), (16200, 14, 6, 38, 17, 198),
              (15800, 11, 6, 33, 15, 174), (17100, 13, 5, 35, 13, 165), (16400, 12, 5, 29, 12, 150)]
    for i, (temp, recv, filled, subs, ivs, days) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        inp(ws, f"H{r}", temp, MONEY, key=(i == 0)); inp(ws, f"I{r}", recv, INT0, key=(i == 0)); inp(ws, f"J{r}", filled, INT0, key=(i == 0))
        inp(ws, f"K{r}", subs, INT0); inp(ws, f"L{r}", ivs, INT0); inp(ws, f"M{r}", days, INT0)
    note(ws, f"B{lo}", "Counts Perm Placements rows whose date falls in this month. Type your own in column C if you track placements elsewhere.")
    note(ws, f"H{lo}", "Gross margin on temp hours billed this month (bill minus pay and on-costs), from your timesheet or invoicing system.")
    note(ws, f"J{lo}", "Job orders filled this month, perm and temp.")
    note(ws, f"M{lo}", "Sum, across the job orders filled this month, of the days from job order received to offer accepted.")
    note(ws, f"O{lo}", "Assumption: desk cost is spread evenly, total desk cost (Settings) ÷ 12 each month.")
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("D", INT0), ("G", MONEY), ("H", MONEY), ("I", INT0), ("J", INT0), ("K", INT0), ("L", INT0),
                     ("M", INT0), ("N", MONEY), ("O", MONEY), ("P", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"Q{tot}", f"=IFERROR(ROUND(J{tot}/I{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"R{tot}", f"=IFERROR(ROUND(M{tot}/J{tot},2),0)", RATIO, bold=True)
    fx(ws, f"S{tot}", f"=IFERROR(ROUND(K{tot}/J{tot},2),0)", RATIO, bold=True)
    fx(ws, f"T{tot}", f"=IFERROR(ROUND(J{tot}/L{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"U{tot}", f"=IFERROR(ROUND(N{tot}/{S}{R_COUNT},2),0)", MONEY, bold=True)
    note(ws, f"Q{tot}", "Year-to-date fill rate from the column totals (filled ÷ received), not an average of the monthly rates.")
    ws.freeze_panes = "B5"


def build_reserve(ws):
    lo, hi = RES_ROWS; tot = hi + 1
    label(ws, "A1", "Rebate Reserve", bold=True).font = TITLE
    label(ws, "A2", "One row per placement under guarantee. Guarantee days, fall-off % and rebate % default from Settings; type your own to override them. Set the outcome when the guarantee ends or the candidate leaves.", muted=True)
    head(ws, 4, ["Placement", "Start date", "Net fee", "Default guarantee days (Settings)", "or type your own days",
                 "Guarantee days used", "Default fall-off % (Settings)", "or type your own fall-off %", "Fall-off % used",
                 "Default rebate % (Settings)", "or type your own rebate %", "Rebate % used", "Actual outcome",
                 "Guarantee end date", "Expected reserve", "Reserve released", "Rebate paid", "Reserve still held"],
         [32, 12, 12, 11, 11, 10, 11, 11, 10, 11, 11, 10, 14, 12, 12, 12, 12, 12])
    dv = DataValidation(type="list", formula1=OUTCOMES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"B": DATE, "C": MONEY, "E": INT0, "H": PCT, "K": PCT}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEHKM", fm); dv.add(f"M{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_GUARANTEE}"), INT0)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), INT0)
        fx(ws, f"G{r}", g.format(f"{S}{R_FALLOFF}"), PCT)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), PCT)
        fx(ws, f"J{r}", g.format(f"{S}{R_REBATE}"), PCT)
        fx(ws, f"L{r}", g.format(f'IF(K{r}<>"",K{r},J{r})'), PCT)
        fx(ws, f"N{r}", g.format(f"B{r}+F{r}"), DATE)
        fx(ws, f"O{r}", g.format(f"ROUND(C{r}*L{r}*I{r}/10000,2)"), MONEY)
        fx(ws, f"P{r}", g.format(f'IF(M{r}="retained",O{r},0)'), MONEY)
        fx(ws, f"Q{r}", g.format(f'IF(M{r}="fell off",ROUND(C{r}*L{r}/100,2),0)'), MONEY)
        fx(ws, f"R{r}", g.format(f'IF(M{r}="in guarantee",O{r},0)'), MONEY)
    year = dt.date.today().year
    sample = [("Northwind Logistics: Warehouse manager", (1, 12), 9750, None, None, None, "retained"),
              ("Brightline Dental: Practice manager", (1, 26), 9600, None, None, None, "retained"),
              ("Acme Engineering: Senior mechanical engineer", (2, 9), 18700, None, None, 50, "fell off"),
              ("Delta Software: Product designer", (4, 7), 7200, None, 12, None, "retained"),
              ("Northwind Logistics: Transport planner", (5, 14), 10400, 180, None, None, "in guarantee"),
              ("Garcia Retail: Store manager", (6, 2), 9000, 180, None, None, "in guarantee")]
    for i, (name, (m, d), fee, days, fall, rebate, outcome) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", dt.date(year, m, d), DATE); inp(ws, f"C{r}", fee, MONEY, key=True)
        inp(ws, f"E{r}", days, INT0); inp(ws, f"H{r}", fall, PCT); inp(ws, f"K{r}", rebate, PCT); inp(ws, f"M{r}", outcome)
    note(ws, f"C{lo}", "Net fee after any split, from the Perm Placements sheet.")
    note(ws, f"K{lo + 2}", "Example override: this client's terms refund half the fee on a fall-off.")
    note(ws, f"H{lo + 3}", "Example override: a role you expect to be riskier carries a higher fall-off %.")
    note(ws, f"E{lo + 4}", "Example override: a senior hire with a 180-day guarantee.")
    note(ws, f"M{lo}", "in guarantee = still inside the guarantee period; retained = guarantee passed, reserve released; fell off = rebate owed. Rows with no outcome are left out of all three columns.")
    note(ws, f"O{lo}", "Assumption: expected reserve = net fee × rebate % × fall-off % ÷ 10,000, the average rebate you expect to owe on this placement.")
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "COPQR":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    plo, phi = PERM_ROWS; tlo, thi = TEMP_ROWS; dtot = DESK_ROWS[1] + 1; rtot = RES_ROWS[1] + 1
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 80
    label(ws, "A1", "Summary", bold=True).font = TITLE
    margins = f"{TA}$P${tlo}:$P${thi}"
    rows = [(2, "Recruiters", f"={S}{R_COUNT}", INT0, "from Settings"),
            (3, "Desk cost per recruiter (annual)", f"={S}{R_DESK_EACH}", MONEY, "from Settings: loaded salary + equal share of tools and overhead"),
            (4, "Recruiter cost per hour", f"={S}{R_COST_HOUR}", MONEY, "from Settings: desk cost per recruiter ÷ working hours"),
            (5, "Year-to-date perm net fees", f"={DP}G{dtot}", MONEY, "Desk Performance totals"),
            (6, "Year-to-date temp margin", f"={DP}H{dtot}", MONEY, "Desk Performance totals"),
            (7, "Year-to-date total gross margin", f"={DP}N{dtot}", MONEY, "perm net fees + temp margin"),
            (8, "Year-to-date desk cost", f"={DP}O{dtot}", MONEY, "total desk cost ÷ 12 for each month entered"),
            (9, "Year-to-date desk profit", f"={DP}P{dtot}", MONEY, "gross margin − desk cost"),
            (10, "Average net fee per perm placement", f"=IFERROR(ROUND(AVERAGE({PP}$O${plo}:$O${phi}),2),0)", MONEY, "simple average of net fees on Perm Placements"),
            (11, "Break-even placements per recruiter (a year)", f"=IFERROR(ROUND(B3/B10,2),0)", RATIO, "desk cost per recruiter ÷ average net fee; perm only, before any temp margin"),
            (13, "Perm placements below the margin % threshold", f'=COUNTIF({PP}$R${plo}:$R${phi},"<"&B12)', INT0, "margin % on Perm Placements below the threshold above"),
            (14, "Thinnest temp assignment", f"=IFERROR(INDEX({TA}$B${tlo}:$B${thi},MATCH(MIN({margins}),{margins},0)),\"\")", None, "worker with the lowest gross margin % on Temp Assignments"),
            (15, "Its gross margin %", f"=IFERROR(MIN({margins}),0)", PCT, "lowest gross margin % on Temp Assignments"),
            (16, "Year-to-date fill rate %", f"={DP}Q{dtot}", PCT, "job orders filled ÷ received, from the Desk Performance totals"),
            (17, "Year-to-date average days to fill", f"={DP}R{dtot}", RATIO, "total days to fill ÷ job orders filled"),
            (18, "Reserve still held", f"={RR}R{rtot}", MONEY, "Rebate Reserve: placements still in guarantee"),
            (19, "Rebates paid", f"={RR}Q{rtot}", MONEY, "Rebate Reserve: fall-offs"),
            (20, "Reserve released", f"={RR}P{rtot}", MONEY, "Rebate Reserve: placements retained past the guarantee")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))
    label(ws, "A12", "Margin % threshold (your own)")
    inp(ws, "B12", 70, PCT, key=True)
    label(ws, "C12", "Default 70. Your own placeholder threshold for flagging thin perm placements, not a recommendation or an industry benchmark.", muted=True)
    note(ws, "B12", "Your own placeholder, not a recommendation and not an industry benchmark. Change it to the margin % below which you want to look at a placement again.")
    note(ws, "B11", "Placements each recruiter must make in a year for their net fees to cover their desk cost, with no temp margin.")


README = """# Staffing Agency Margin Workbook

Thank you for buying the workbook. Open `staffing-agency-margin-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** recruiters, average recruiter salary, benefits and payroll taxes %, tools and job boards, office and other overhead, working hours, and your default perm fee %, temp on-costs %, temp markup %, guarantee days, fall-off % and rebate %; desk cost per recruiter, recruiter cost per hour and the reserve % of fee follow.
2. **Perm Placements:** each placement with date, client, role, salary, fee % (or your own), split %, recruiter hours, advertising, invoiced and collected; gross fee, split fee, net fee, delivery cost, margin, margin %, net fee per hour, rebate reserve and net fee after reserve follow.
3. **Temp Assignments:** each assignment with pay rate, on-costs % and markup % (or your own), an optional agreed bill rate, hours per week and weeks; bill rate, loaded cost, gross margin per hour and %, weekly and assignment revenue and margin follow.
4. **Desk Performance:** each month's job orders received and filled, submittals, interviews, days to fill and temp margin, with perm placements and fees counted from Perm Placements; gross margin, desk cost, desk profit, fill rate, days to fill, submittals per fill, interview to fill % and margin per recruiter follow.
5. **Rebate Reserve:** placements under guarantee with start date, net fee, guarantee days, fall-off % and rebate % (or your own) and the outcome; guarantee end date, expected reserve, reserve released, rebate paid and reserve still held follow.
6. **Summary:** desk cost and cost per hour, year-to-date fees, temp margin, desk cost and desk profit, break-even placements per recruiter, placements below your margin % threshold, the thinnest temp assignment, fill rate, days to fill and the reserve position.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_perm(wb.create_sheet("Perm Placements"))
    build_temp(wb.create_sheet("Temp Assignments")); build_desk(wb.create_sheet("Desk Performance"))
    build_reserve(wb.create_sheet("Rebate Reserve")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
