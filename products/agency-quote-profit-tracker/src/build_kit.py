"""Build the Agency Quote & Project Profit Tracker workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
RATE_ROWS = (5, 14); QUOTE_ROWS = (6, 25); PROJ_ROWS = (5, 40)
S = "Settings!$B$"; RC = "'Rate Card'!"


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Agency Quote & Project Profit Tracker", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: productive hours per day and your default contingency and discount for quotes.",
        "2. Rate Card: one row per role with the bill rate and the internal cost per hour (salary, contributions and overhead divided by productive hours). Margin and markup per role are calculated.",
        "3. Quote: name the client and project, then list tasks with a role (pick from the list) and hours. Bill rate and cost rate are looked up from the Rate Card; labour, expenses, contingency and discount roll up into the quote total, blended rate, expected profit and margin. Copy the sheet for each new quote.",
        "4. Projects: one row per project with price, estimated and actual hours, cost per hour, expenses, invoiced and paid. Profit, margin, effective hourly rate, hours overrun and outstanding balance are calculated.",
        "5. Summary: totals and averages across every project row.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "Blue cells that already contain a formula (defaults pulled from Settings) may be overwritten with your own figure.",
        "", "DEFINITIONS",
        "Blended rate = total labour at bill rates ÷ total hours. Margin = (price − cost) ÷ price. Markup = bill rate ÷ cost rate.",
        "Contingency = subtotal × contingency %. Quote total = (subtotal + contingency) × (1 − discount %).",
        "Effective hourly rate = price ÷ actual hours. Hours overrun = (actual − estimated) ÷ estimated. Outstanding = invoiced − paid.",
        "Arithmetic only; not financial or legal advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 60
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [("Productive hours per working day", 8, "0.0", "Used to convert quoted hours into working days."),
            ("Default contingency on quotes (%)", 15, "0.00", "Allowance for scope the estimate missed; 10 to 20 is common."),
            ("Default discount on quotes (%)", 0, "0.00", "Applied after contingency; leave at 0 unless agreed.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True); label(ws, f"C{i}", note, muted=True)


def build_rate_card(ws):
    lo, hi = RATE_ROWS
    label(ws, "A1", "Rate Card", bold=True).font = TITLE
    label(ws, "A2", "One row per role. Cost per hour = salary + contributions + overhead ÷ productive hours (the Employee Cost Calculator on knackdesk.com gives this figure).", muted=True)
    head(ws, 4, ["Role", "Bill rate / hour", "Cost / hour", "Margin %", "Markup (bill ÷ cost)"], [24, 16, 14, 12, 18])
    for r in range(lo, hi + 1):
        blank(ws, r, "ABC", {"B": MONEY, "C": MONEY})
        fx(ws, f"D{r}", f'=IF(A{r}="","",IFERROR(ROUND((B{r}-C{r})/B{r}*100,2),""))', PCT)
        fx(ws, f"E{r}", f'=IF(A{r}="","",IFERROR(ROUND(B{r}/C{r},2),""))', "0.00")
    for i, (role, bill, cost) in enumerate([("Lead / strategist", 150, 70), ("Designer", 80, 40), ("Developer", 95, 50), ("Project manager", 85, 45)]):
        r = lo + i; inp(ws, f"A{r}", role, key=(i == 0)); inp(ws, f"B{r}", bill, MONEY, key=(i == 0)); inp(ws, f"C{r}", cost, MONEY, key=(i == 0))


def build_quote(ws):
    lo, hi = QUOTE_ROWS; tot = hi + 1; rlo, rhi = RATE_ROWS
    label(ws, "A1", "Quote", bold=True).font = TITLE
    label(ws, "A2", "Client"); inp(ws, "B2", "Example Ltd", key=True)
    label(ws, "A3", "Project"); inp(ws, "B3", "Website redesign", key=True)
    ws.column_dimensions["A"].width = 34
    head(ws, 5, ["Task", "Role", "Hours", "Bill rate", "Amount", "Cost rate", "Internal cost"], [34, 20, 10, 12, 14, 12, 14])
    dv = DataValidation(type="list", formula1=f"={RC}$A${rlo}:$A${rhi}", allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        blank(ws, r, "ABC", {"C": "0.0"}); dv.add(f"B{r}")
        fx(ws, f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX({RC}$B${rlo}:$B${rhi},MATCH(B{r},{RC}$A${rlo}:$A${rhi},0)),""))', MONEY)
        fx(ws, f"E{r}", f'=IF(B{r}="","",IFERROR(ROUND(C{r}*D{r},2),""))', MONEY)
        fx(ws, f"F{r}", f'=IF(B{r}="","",IFERROR(INDEX({RC}$C${rlo}:$C${rhi},MATCH(B{r},{RC}$A${rlo}:$A${rhi},0)),""))', MONEY)
        fx(ws, f"G{r}", f'=IF(B{r}="","",IFERROR(ROUND(C{r}*F{r},2),""))', MONEY)
    for i, (task, role, hours) in enumerate([("Discovery and strategy", "Lead / strategist", 10), ("Design", "Designer", 30), ("Build", "Developer", 40), ("Project management", "Project manager", 12)]):
        r = lo + i; inp(ws, f"A{r}", task, key=(i == 0)); inp(ws, f"B{r}", role, key=(i == 0)); inp(ws, f"C{r}", hours, "0.0", key=(i == 0))
    label(ws, f"A{tot}", "Totals", bold=True); fx(ws, f"C{tot}", f"=SUM(C{lo}:C{hi})", "0.0", bold=True); fx(ws, f"E{tot}", f"=SUM(E{lo}:E{hi})", MONEY, bold=True); fx(ws, f"G{tot}", f"=SUM(G{lo}:G{hi})", MONEY, bold=True)
    label(ws, "A28", "Expenses passed through"); inp(ws, "B28", 500, MONEY, key=True)
    label(ws, "A29", "Contingency (%)"); inp(ws, "B29", f"={S}3", "0.00")
    label(ws, "A30", "Discount (%)"); inp(ws, "B30", f"={S}4", "0.00")
    calc = [(32, "Labour at bill rates", f"=E{tot}", MONEY, False), (33, "Expenses", "=B28", MONEY, False), (34, "Subtotal", "=B32+B33", MONEY, False),
            (35, "Contingency", "=ROUND(B34*B29/100,2)", MONEY, False), (36, "Before discount", "=B34+B35", MONEY, False), (37, "Discount", "=ROUND(B36*B30/100,2)", MONEY, False),
            (38, "QUOTE TOTAL", "=B36-B37", MONEY, True), (39, "Blended hourly rate", f'=IFERROR(ROUND(E{tot}/C{tot},2),"")', MONEY, True),
            (40, "Expected internal cost", f"=G{tot}+B33", MONEY, False), (41, "Expected profit", "=B38-B40", MONEY, True), (42, "Expected margin %", '=IFERROR(ROUND(B41/B38*100,2),"")', PCT, True),
            (43, "Working days of effort", f'=IFERROR(ROUND(C{tot}/{S}2,1),"")', "0.0", False)]
    for r, name, f, fmt, bold in calc:
        label(ws, f"A{r}", name, bold=bold); fx(ws, f"B{r}", f, fmt, bold=bold)
    ws.freeze_panes = "A6"


def build_projects(ws):
    lo, hi = PROJ_ROWS; tot = hi + 1
    label(ws, "A1", "Projects", bold=True).font = TITLE
    label(ws, "A2", "One row per project. Enter actual hours as they are logged; profit and overrun update. Cost per hour = blended internal cost of the people on the job.", muted=True)
    head(ws, 4, ["Client", "Project", "Status", "Price", "Estimated hours", "Actual hours", "Cost / hour", "Expenses", "Labour cost", "Total cost", "Profit", "Margin %", "Effective rate / hour", "Hours overrun %", "Invoiced", "Paid", "Outstanding"],
         [16, 22, 10, 13, 11, 11, 11, 12, 13, 13, 13, 10, 13, 12, 13, 13, 13])
    dv = DataValidation(type="list", formula1='"quoted,active,done,lost"', allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFGHOP", {"D": MONEY, "E": "0.0", "F": "0.0", "G": MONEY, "H": MONEY, "O": MONEY, "P": MONEY}); dv.add(f"C{r}")
        fx(ws, f"I{r}", f'=IF(A{r}="","",ROUND(F{r}*G{r},2))', MONEY)
        fx(ws, f"J{r}", f'=IF(A{r}="","",I{r}+H{r})', MONEY)
        fx(ws, f"K{r}", f'=IF(A{r}="","",D{r}-J{r})', MONEY)
        fx(ws, f"L{r}", f'=IF(A{r}="","",IFERROR(ROUND(K{r}/D{r}*100,2),""))', PCT)
        fx(ws, f"M{r}", f'=IF(A{r}="","",IFERROR(ROUND(D{r}/F{r},2),""))', MONEY)
        fx(ws, f"N{r}", f'=IF(A{r}="","",IFERROR(ROUND((F{r}-E{r})/E{r}*100,2),""))', PCT)
        fx(ws, f"Q{r}", f'=IF(A{r}="","",O{r}-P{r})', MONEY)
    sample = [("Example Ltd", "Website redesign", "done", 10000, 92, 80, 55, 500, 10000, 10000), ("Sample Co", "Brand identity", "done", 6500, 60, 74, 48, 200, 6500, 3250), ("Demo Inc", "Monthly retainer (Sep)", "active", 4000, 40, 22, 50, 0, 4000, 0)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, v in zip("ABCDEFGHOP", row):
            inp(ws, f"{col}{r}", v, {"D": MONEY, "G": MONEY, "H": MONEY, "O": MONEY, "P": MONEY, "E": "0.0", "F": "0.0"}.get(col), key=(i == 0 and col in "ADF"))
    label(ws, f"A{tot}", "Totals", bold=True)
    for col in "DFHIJKOPQ": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY if col not in "F" else "0.0", bold=True)
    fx(ws, f"L{tot}", f'=IFERROR(ROUND(K{tot}/D{tot}*100,2),"")', PCT, bold=True)
    ws.freeze_panes = "C5"


def build_summary(ws):
    lo, hi = PROJ_ROWS; tot = hi + 1
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "Summary", bold=True).font = TITLE
    rows = [(2, "Projects tracked", f"=COUNTA(Projects!A{lo}:A{hi})", "0", None),
            (3, "Revenue (sum of prices)", f"=Projects!D{tot}", MONEY, None),
            (4, "Total cost", f"=Projects!J{tot}", MONEY, "labour at cost rates plus expenses"),
            (5, "Profit", f"=Projects!K{tot}", MONEY, None),
            (6, "Overall margin %", f'=IFERROR(ROUND(B5/B3*100,2),"")', PCT, "profit ÷ revenue"),
            (7, "Hours worked", f"=Projects!F{tot}", "0.0", None),
            (8, "Average effective rate per hour", '=IFERROR(ROUND(B3/B7,2),"")', MONEY, "revenue ÷ hours"),
            (9, "Invoiced", f"=Projects!O{tot}", MONEY, None),
            (10, "Paid", f"=Projects!P{tot}", MONEY, None),
            (11, "Outstanding", f"=Projects!Q{tot}", MONEY, None),
            (12, "Projects that overran their estimate", f'=COUNTIF(Projects!N{lo}:N{hi},">0")', "0", None),
            (13, "Average hours overrun %", f'=IFERROR(ROUND(AVERAGE(Projects!N{lo}:N{hi}),2),"")', PCT, "across projects with hours entered")]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True)
        if note: label(ws, f"C{r}", note, muted=True)


README = """# Agency Quote & Project Profit Tracker

Thank you for buying the tracker. Open `agency-quote-profit-tracker.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** productive hours per day, default contingency and discount.
2. **Rate Card:** bill rate and internal cost per role; margin and markup are calculated.
3. **Quote:** list tasks with a role and hours. Rates are looked up; total, blended rate, expected profit and margin follow. Duplicate the sheet per quote.
4. **Projects:** one row per project with price, hours, cost, expenses, invoiced and paid; profit, margin, effective rate, overrun and outstanding are calculated.
5. **Summary:** totals and averages across all projects.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_rate_card(wb.create_sheet("Rate Card")); build_quote(wb.create_sheet("Quote")); build_projects(wb.create_sheet("Projects")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "agency-quote-profit-tracker.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "agency-quote-profit-tracker.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "agency-quote-profit-tracker.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
