"""Build the SaaS Metrics Dashboard workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
PLAN_ROWS = (5, 15); MONTH_ROWS = (5, 40)


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "SaaS Metrics Dashboard", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Plans: one row per plan with price, billing period and current paying customers. MRR, ARR and ARPA are calculated.",
        "2. Monthly Metrics: one row per month. Enter the first month's starting MRR and customers; later months carry forward automatically. Enter new, expansion, contraction and churned MRR and customer counts; growth, churn and NRR are calculated.",
        "3. Unit Economics: gross margin, monthly churn and acquisition spend give CAC, lifetime value, LTV to CAC and payback. ARPA defaults to the Plans sheet.",
        "4. Runway: cash, monthly costs and MRR give burn, months of runway, the run-out month, and the cost level needed for a target runway.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "Blue cells that already contain a formula (ARPA, MRR links) are defaults you may overwrite.",
        "", "DEFINITIONS",
        "MRR normalises every plan to a monthly amount (annual price ÷ 12). ARR = MRR × 12. ARPA = MRR ÷ paying customers.",
        "Gross revenue churn = (contraction + churned) ÷ starting MRR. Net revenue churn also subtracts expansion. NRR = (start + expansion − contraction − churned) ÷ start.",
        "LTV = ARPA × gross margin ÷ monthly churn. CAC = spend ÷ new customers. Payback = CAC ÷ (ARPA × gross margin).",
        "Arithmetic only; not financial advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_plans(ws):
    lo, hi = PLAN_ROWS
    label(ws, "A1", "Plans", bold=True).font = TITLE
    head(ws, 4, ["Plan", "Price", "Billed every", "Paying customers", "MRR"], [22, 12, 14, 16, 14])
    dv = DataValidation(type="list", formula1='"month,quarter,year"', allow_blank=True); ws.add_data_validation(dv)
    for r in range(lo, hi + 1):
        for col in "ABCD": ws[f"{col}{r}"].font = BLUE
        ws[f"B{r}"].number_format = MONEY; dv.add(f"C{r}")
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(B{r}*D{r}/IF(C{r}="year",12,IF(C{r}="quarter",3,1)),2))', MONEY)
    inp(ws, f"A{lo}", "Basic", key=True); inp(ws, f"B{lo}", 10, MONEY); inp(ws, f"C{lo}", "month"); inp(ws, f"D{lo}", 100, "0")
    inp(ws, f"A{lo+1}", "Pro"); inp(ws, f"B{lo+1}", 240, MONEY); inp(ws, f"C{lo+1}", "year"); inp(ws, f"D{lo+1}", 50, "0")
    label(ws, "A16", "Total", bold=True); fx(ws, "D16", f"=SUM(D{lo}:D{hi})", "0", bold=True); fx(ws, "E16", f"=SUM(E{lo}:E{hi})", MONEY, bold=True)
    label(ws, "A17", "ARR", bold=True); fx(ws, "E17", "=E16*12", MONEY, bold=True)
    label(ws, "A18", "ARPA (MRR per paying customer)", bold=True); fx(ws, "E18", "=IFERROR(ROUND(E16/D16,2),0)", MONEY, bold=True)


def build_monthly(ws):
    lo, hi = MONTH_ROWS
    label(ws, "A1", "Monthly Metrics", bold=True).font = TITLE
    label(ws, "A2", "Enter the first row's starting MRR and customers; later rows carry forward from the previous month's ending values.", muted=True)
    head(ws, 4, ["Month", "Starting MRR", "New MRR", "Expansion MRR", "Contraction MRR", "Churned MRR", "Ending MRR", "Net new MRR", "Growth %", "Gross rev. churn %", "Net rev. churn %", "NRR %", "Customers start", "New customers", "Lost customers", "Customer churn %"],
         [11, 13, 12, 13, 14, 12, 13, 13, 10, 13, 13, 10, 13, 12, 12, 13])
    for r in range(lo, hi + 1):
        for col in "ACDEFNO": ws[f"{col}{r}"].font = BLUE
        ws[f"A{r}"].number_format = "mmm yyyy"
        for col in "BCDEFGH": ws[f"{col}{r}"].number_format = MONEY
        if r == lo:
            ws[f"B{r}"].font = BLUE; ws[f"M{r}"].font = BLUE
        else:
            fx(ws, f"B{r}", f'=IF(A{r}="","",G{r-1})', MONEY)
            fx(ws, f"M{r}", f'=IF(A{r}="","",M{r-1}+N{r-1}-O{r-1})', "0")
        fx(ws, f"G{r}", f'=IF(A{r}="","",B{r}+C{r}+D{r}-E{r}-F{r})', MONEY)
        fx(ws, f"H{r}", f'=IF(A{r}="","",G{r}-B{r})', MONEY)
        fx(ws, f"I{r}", f'=IF(OR(A{r}="",B{r}=0),"",ROUND(H{r}/B{r}*100,2))', PCT)
        fx(ws, f"J{r}", f'=IF(OR(A{r}="",B{r}=0),"",ROUND((E{r}+F{r})/B{r}*100,2))', PCT)
        fx(ws, f"K{r}", f'=IF(OR(A{r}="",B{r}=0),"",ROUND((E{r}+F{r}-D{r})/B{r}*100,2))', PCT)
        fx(ws, f"L{r}", f'=IF(OR(A{r}="",B{r}=0),"",ROUND((B{r}+D{r}-E{r}-F{r})/B{r}*100,2))', PCT)
        fx(ws, f"P{r}", f'=IF(OR(A{r}="",M{r}=0),"",ROUND(O{r}/M{r}*100,2))', PCT)
    first = dt.date.today().replace(day=1)
    prev = (first - dt.timedelta(days=1)).replace(day=1)
    inp(ws, f"A{lo}", prev, "mmm yyyy", key=True); inp(ws, f"B{lo}", 10000, MONEY, key=True); inp(ws, f"C{lo}", 1200, MONEY); inp(ws, f"D{lo}", 300, MONEY); inp(ws, f"E{lo}", 100, MONEY); inp(ws, f"F{lo}", 400, MONEY)
    inp(ws, f"M{lo}", 500, "0", key=True); inp(ws, f"N{lo}", 40, "0"); inp(ws, f"O{lo}", 15, "0")
    inp(ws, f"A{lo+1}", first, "mmm yyyy"); inp(ws, f"C{lo+1}", 900, MONEY); inp(ws, f"D{lo+1}", 250, MONEY); inp(ws, f"E{lo+1}", 50, MONEY); inp(ws, f"F{lo+1}", 350, MONEY); inp(ws, f"N{lo+1}", 30, "0"); inp(ws, f"O{lo+1}", 12, "0")
    ws.freeze_panes = "B5"


def build_unit(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 60
    label(ws, "A1", "Unit Economics", bold=True).font = TITLE
    rows = [("Average monthly revenue per account (ARPA)", "=Plans!E18", MONEY, "Defaults to the Plans sheet; overwrite if you prefer another figure."),
            ("Gross margin (%)", 80, "0.0", "Revenue minus hosting, payment fees and support, as a share of revenue."),
            ("Monthly customer churn (%)", 3, "0.00", "Use the latest Customer churn % from Monthly Metrics."),
            ("Sales and marketing spend in the period", 30000, MONEY, None),
            ("New customers won in the period", 100, "0", None)]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True, note=note)
        if note: label(ws, f"C{i}", note, muted=True)
    calc = [("Customer acquisition cost (CAC)", "=IFERROR(ROUND(B5/B6,2),0)", MONEY, "spend ÷ new customers"),
            ("Expected customer lifetime (months)", '=IFERROR(ROUND(1/(B4/100),2),"")', "0.00", "1 ÷ monthly churn"),
            ("Lifetime value (LTV)", '=IFERROR(ROUND(B2*B3/100/(B4/100),2),"")', MONEY, "ARPA × margin ÷ churn"),
            ("LTV to CAC ratio", '=IFERROR(ROUND(B10/B8,2),"")', "0.00", "LTV ÷ CAC"),
            ("CAC payback (months)", '=IFERROR(ROUND(B8/(B2*B3/100),2),"")', "0.00", "CAC ÷ (ARPA × margin)")]
    for i, (name, f, fmt, note) in enumerate(calc, start=8):
        label(ws, f"A{i}", name, bold=True); fx(ws, f"B{i}", f, fmt, bold=True); label(ws, f"C{i}", note, muted=True)


def build_runway(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 55
    label(ws, "A1", "Runway", bold=True).font = TITLE
    label(ws, "A2", "Cash in the bank"); inp(ws, "B2", 120000, MONEY, key=True)
    label(ws, "A3", "Monthly costs (all cash out)"); inp(ws, "B3", 25000, MONEY, key=True)
    label(ws, "A4", "Monthly recurring revenue"); inp(ws, "B4", "=Plans!E16", MONEY, note="Defaults to the Plans sheet total MRR.")
    label(ws, "A6", "Net monthly burn", bold=True); fx(ws, "B6", "=B3-B4", MONEY, bold=True); label(ws, "C6", "costs minus MRR; negative means you are cash-positive", muted=True)
    label(ws, "A7", "Months of runway", bold=True); fx(ws, "B7", '=IF(B6<=0,"Cash-positive",ROUND(B2/B6,1))', "0.0", bold=True)
    label(ws, "A8", "Cash runs out (month)", bold=True); fx(ws, "B8", '=IF(B6<=0,"",EDATE(TODAY(),INT(B2/B6)))', "mmm yyyy", bold=True); label(ws, "C8", "counting from the current month, part months rounded down", muted=True)
    label(ws, "A9", "Target runway (months)"); inp(ws, "B9", 18, "0", key=True)
    label(ws, "A10", "Maximum monthly costs for the target", bold=True); fx(ws, "B10", "=ROUND(B4+B2/B9,2)", MONEY, bold=True)
    label(ws, "A11", "Monthly cut needed", bold=True); fx(ws, "B11", "=MAX(0,B3-B10)", MONEY, bold=True)


README = """# SaaS Metrics Dashboard

Thank you for buying the dashboard. Open `saas-metrics-dashboard.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Plans:** one row per plan with price, billing period and paying customers. MRR, ARR and ARPA are calculated.
2. **Monthly Metrics:** one row per month; enter the first starting MRR and customers, then each month's movements. Growth, churn and NRR are calculated and later months carry forward.
3. **Unit Economics** and **Runway** read from Plans by default; the yellow cells are yours to set.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_plans(wb.create_sheet("Plans")); build_monthly(wb.create_sheet("Monthly Metrics")); build_unit(wb.create_sheet("Unit Economics")); build_runway(wb.create_sheet("Runway"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "saas-metrics-dashboard.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "saas-metrics-dashboard.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "saas-metrics-dashboard.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
