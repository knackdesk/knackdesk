"""Build the Freelance Pricing Kit workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'
DATE = "yyyy-mm-dd"


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK)
    return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Freelance Pricing Kit", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Rates: fill the yellow cells once. Your hourly, day and week rates are calculated from them.",
        "2. Rate Card: adjust the multipliers for rush and weekend work; the card updates from your rates.",
        "3. Quote Builder: enter the client, date and line items. Discount, VAT, deposit and milestones are calculated.",
        "4. Retainer: enter an hour block and a commitment discount to price a monthly retainer.",
        "5. Revenue Plan: enter an annual goal and your averages to see how many proposals to send.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs to set first.",
        "Blue cells that already contain a formula (for example the unit rate on a quote line) are defaults you may overwrite.",
        "", "NOTES",
        "The workbook uses formulas only, no macros. It works in Excel 2010 or later, Google Sheets and Apple Numbers.",
        "This kit does arithmetic; it is not legal, tax or financial advice. Check VAT and contract rules that apply to you.",
        "Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3):
        label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_rates(ws):
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 58
    label(ws, "A1", "Rates", bold=True).font = TITLE
    rows = [("Target income per year", 60000, MONEY, "What you want to keep after business expenses."),
            ("Business expenses per year", 6000, MONEY, "Software, insurance, equipment, accountant, training."),
            ("Tax rate to gross up for (%)", 0, "0.0", "Leave 0 if the target income is pre-tax."),
            ("Weeks off per year", 6, "0", "Holiday, sick days, public holidays, gaps between projects."),
            ("Hours worked per week", 40, "0", None),
            ("Billable share of hours (%)", 60, "0", "Most solo freelancers bill 50-70% of their hours."),
            ("Hours in a billable day", 8, "0.0", None),
            ("VAT rate on quotes (%)", 20, "0.0", "Set 0 if you do not charge VAT."),
            ("Default deposit on quotes (%)", 30, "0", None),
            ("Currency label", "USD", None, "Text only; the kit does not convert currencies.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True, note=note)
        if note: label(ws, f"C{i}", note, muted=True)
    calc = [("Working weeks", "=52-B5", "0"), ("Billable hours per year", "=ROUND(B13*B6*B7/100,2)", "0.00"),
            ("Revenue needed", "=B3+B2/(1-B4/100)", MONEY), ("Hourly rate", '=IFERROR(ROUND(B15/B14,2),0)', MONEY),
            ("Day rate", "=ROUND(B16*B8,2)", MONEY), ("Week rate (5 days)", "=ROUND(B17*5,2)", MONEY)]
    for i, (name, f, fmt) in enumerate(calc, start=13):
        label(ws, f"A{i}", name, bold=i >= 16); fx(ws, f"B{i}", f, fmt, bold=i >= 16)
    label(ws, "C16", "revenue needed ÷ billable hours; this drives every other sheet", muted=True)


def build_card(ws):
    label(ws, "A1", "Rate Card", bold=True).font = TITLE
    label(ws, "A2", "Multipliers are yours to change; rates come from the Rates sheet.", muted=True)
    head(ws, 4, ["Service tier", "Multiplier", "Hourly", "Day", "Week"], [26, 12, 14, 14, 14])
    tiers = [("Standard", 1.0), ("Rush (48-hour turnaround)", 1.5), ("Weekend or evening", 2.0), ("Existing client (loyalty)", 0.9)]
    for i, (name, mult) in enumerate(tiers, start=5):
        inp(ws, f"A{i}", name); inp(ws, f"B{i}", mult, "0.00")
        fx(ws, f"C{i}", f"=ROUND(Rates!$B$16*B{i},2)", MONEY); fx(ws, f"D{i}", f"=ROUND(Rates!$B$17*B{i},2)", MONEY); fx(ws, f"E{i}", f"=ROUND(Rates!$B$18*B{i},2)", MONEY)
    label(ws, "A10", "Minimum charge (hours)"); inp(ws, "B10", 1, "0.0"); fx(ws, "C10", "=ROUND(Rates!$B$16*B10,2)", MONEY)
    label(ws, "D10", "smallest job you will invoice", muted=True)


def build_quote(ws):
    ws.column_dimensions["A"].width = 40
    for col, w in zip("BCDEF", [12, 12, 14, 16, 16]): ws.column_dimensions[col].width = w
    label(ws, "A1", "Quote Builder", bold=True).font = TITLE
    label(ws, "A2", "Client"); inp(ws, "B2", "Example Client Ltd", key=True)
    label(ws, "A3", "Quote date"); inp(ws, "B3", dt.date.today(), DATE, key=True)
    label(ws, "A4", "Valid for (days)"); inp(ws, "B4", 30, "0")
    label(ws, "A5", "Valid until"); fx(ws, "B5", "=B3+B4", DATE)
    head(ws, 7, ["Line item", "Qty or hours", "Unit", "Unit rate", "Amount"])
    for r in range(8, 20):
        inp(ws, f"A{r}", None); inp(ws, f"B{r}", None, "0.00"); inp(ws, f"C{r}", None)
        inp(ws, f"D{r}", "=Rates!$B$16", MONEY, note="Defaults to your hourly rate; overwrite for fixed items.")
        fx(ws, f"E{r}", f'=IF(OR(A{r}="",B{r}=""),"",ROUND(B{r}*D{r},2))', MONEY)
    inp(ws, "A8", "Discovery workshop"); inp(ws, "B8", 4, "0.00"); inp(ws, "C8", "hours")
    inp(ws, "A9", "Design and build"); inp(ws, "B9", 30, "0.00"); inp(ws, "C9", "hours")
    label(ws, "A21", "Subtotal", bold=True); fx(ws, "F21", "=SUM(E8:E19)", MONEY, bold=True)
    label(ws, "A22", "Discount (%)"); inp(ws, "B22", 0, "0.0"); fx(ws, "F22", "=-ROUND(F21*B22/100,2)", MONEY)
    label(ws, "A23", "VAT (%)"); inp(ws, "B23", "=Rates!$B$9", "0.0", note="Defaults to the Rates sheet; overwrite per quote."); fx(ws, "F23", "=ROUND((F21+F22)*B23/100,2)", MONEY)
    label(ws, "A24", "Total", bold=True); fx(ws, "F24", "=F21+F22+F23", MONEY, bold=True)
    label(ws, "A26", "Deposit (%)"); inp(ws, "B26", "=Rates!$B$10", "0"); fx(ws, "F26", "=ROUND(F24*B26/100,2)", MONEY)
    label(ws, "A27", "Balance after deposit"); fx(ws, "F27", "=F24-F26", MONEY)
    label(ws, "A28", "Number of milestones for the balance"); inp(ws, "B28", 2, "0"); fx(ws, "F28", '=IF(B28>0,ROUND(F27/B28,2),"")', MONEY)
    label(ws, "A29", "Per milestone (last one absorbs rounding)", muted=True); fx(ws, "F29", '=IF(B28>0,F27-F28*(B28-1),"")', MONEY)
    ws.freeze_panes = "A8"


def build_retainer(ws):
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "Retainer", bold=True).font = TITLE
    label(ws, "A2", "Hours reserved per month"); inp(ws, "B2", 20, "0.0", key=True)
    label(ws, "A3", "Hourly rate"); inp(ws, "B3", "=Rates!$B$16", MONEY, note="Defaults to your hourly rate; overwrite if needed.")
    label(ws, "A4", "Commitment discount (%)"); inp(ws, "B4", 10, "0.0", key=True)
    label(ws, "A5", "Term (months)"); inp(ws, "B5", 12, "0", key=True)
    rows = [("List price per month", "=ROUND(B2*B3,2)"), ("Monthly retainer fee", "=ROUND(B7*(1-B4/100),2)"),
            ("Effective hourly rate", "=ROUND(B8/B2,2)"), ("Total over the term", "=B8*B5"), ("Discount given over the term", "=(B7-B8)*B5")]
    for i, (name, f) in enumerate(rows, start=7):
        label(ws, f"A{i}", name, bold=i in (8, 9)); fx(ws, f"B{i}", f, MONEY, bold=i in (8, 9))
    label(ws, "C9", "compare with the Rates sheet hourly rate; do not go below your floor", muted=True)


def build_plan(ws):
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 14
    label(ws, "A1", "Revenue Plan", bold=True).font = TITLE
    label(ws, "A2", "Annual revenue goal"); inp(ws, "B2", 120000, MONEY, key=True)
    label(ws, "A3", "Average project value"); inp(ws, "B3", 5000, MONEY, key=True)
    label(ws, "A4", "Projects a typical client buys per year"); inp(ws, "B4", 2, "0.0")
    label(ws, "A5", "Proposal win rate (%)"); inp(ws, "B5", 25, "0.0", key=True)
    head(ws, 7, ["", "Per year", "Per month"])
    rows = [("Projects to deliver", "=ROUNDUP(B2/B3,0)"), ("Clients to win", "=ROUNDUP(B8/B4,0)"), ("Proposals to send", "=ROUNDUP(B9/(B5/100),0)")]
    for i, (name, f) in enumerate(rows, start=8):
        label(ws, f"A{i}", name, bold=True); fx(ws, f"B{i}", f, "0", bold=True); fx(ws, f"C{i}", f"=ROUND(B{i}/12,1)", "0.0")


README = """# Freelance Pricing Kit

Thank you for buying the kit. Open `freelance-pricing-kit.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. Start on the **Rates** sheet and fill the yellow cells. Your hourly, day and week rates are calculated there.
2. **Rate Card**, **Quote Builder**, **Retainer** and **Revenue Plan** all read from Rates. The **Start Here** sheet explains each one.

Blue text = type here. Black text = formula, leave it alone. Blue cells with a formula are defaults you may overwrite.

No macros, no tracking. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_rates(wb.create_sheet("Rates")); build_card(wb.create_sheet("Rate Card")); build_quote(wb.create_sheet("Quote Builder"))
    build_retainer(wb.create_sheet("Retainer")); build_plan(wb.create_sheet("Revenue Plan"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "freelance-pricing-kit.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "freelance-pricing-kit.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "freelance-pricing-kit.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
