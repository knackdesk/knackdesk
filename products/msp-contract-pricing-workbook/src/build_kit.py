"""Build the MSP Contract Pricing Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "msp-contract-pricing-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"
TOOL_ROWS = (5, 24); CONTRACT_ROWS = (5, 34); MONTH_ROWS = (5, 16); BLOCK_ROWS = (5, 14); ONBOARD_ROWS = (20, 29)
S = "Settings!$B$"  # settings cell prefix
BASES = '"per user,per device,per site,fixed"'
PLANS = '"per user,per device,hybrid"'
TS = "'Tool Stack'!"; CT = "Contracts!"; TK = "Tickets!"; BO = "'Blocks & Onboarding'!"
# Settings rows (column B): referenced by every other sheet.
R_TECHS = 2; R_SALARY = 3; R_TAXES = 4; R_PROD_HRS = 5; R_OVERHEAD = 12; R_TARGET = 13; R_BLOCK_DISC = 14
R_UNUSED = 15; R_RATE = 16; R_TECH_COST = 17; R_TOTAL_HRS = 18; R_HOURLY = 19; R_MINUTE = 20
# Tool Stack summary cells.
TS_TOTAL = f"{TS}$E${TOOL_ROWS[1] + 1}"; TS_FIXED = f"{TS}$B${TOOL_ROWS[1] + 4}"; TS_UNIT = f"{TS}$B${TOOL_ROWS[1] + 5}"


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "MSP Contract Pricing Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: technician count, salary, payroll taxes and benefits, productive hours, overhead line by line, target margin, block hours discount, expected unused block hours and your default hourly rate. Your loaded hourly cost follows and every other sheet reads it.",
        "2. Tool Stack: one row per tool with its billing basis, unit cost and units in use. Monthly cost, the stack total, the fixed tool cost and the per-unit tool cost shared across contracts are calculated.",
        "3. Contracts: one row per client with plan, users, devices, monthly fee, support hours, dedicated tool costs, onboarding fee and contract months. Shared tool cost, labour cost, total cost, gross profit, margin %, fee and cost per user, the price per user at your target margin and break-even hours are calculated.",
        "4. Tickets: one row per month with tickets closed and average minutes per ticket. Technician cost and tool and overhead cost default from Settings and the Tool Stack (type your own to override). Cost per ticket, labour cost per ticket, tickets per technician and the share of productive hours spent on tickets are calculated, with totals.",
        "5. Blocks & Onboarding: block hours quotes (hours, rate, discount, expiry, expected unused hours) with block price, effective rate and effective rate on used hours; and onboarding quotes (hours, rate, setup costs, margin) with the onboarding fee, amortised per month and its share of the contract value.",
        "6. Summary: loaded hourly cost, technicians and productive hours, tool stack total, monthly recurring revenue, cost, gross profit and overall margin, contracts below target margin, most and least profitable contract, cost per ticket, block and onboarding revenue.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Loaded hourly cost: what one productive technician hour costs you = (technicians × salary × (1 + payroll taxes and benefits %) + total monthly overhead) ÷ (technicians × productive hours per technician).",
        "Cost per user: a contract's total monthly cost (shared tools + dedicated tools + labour) ÷ users. Per-unit tools are shared across contracts in proportion to each contract's users + devices; fixed tools (PSA, documentation, remote support) are not allocated to contracts and show in the Tickets cost instead.",
        "Margin on price vs markup on cost: margin % = profit ÷ price; markup % = profit ÷ cost. A 30% margin is a 42.86% markup. Price at target margin = cost ÷ (1 − target margin %), never cost × (1 + target %).",
        "Contract gross margin = (monthly fee − total contract cost) ÷ monthly fee. Break-even hours = (monthly fee − tool costs) ÷ loaded hourly cost: the support hours a contract can absorb before it loses money.",
        "Cost per ticket = (technician cost + tool and overhead cost for the month) ÷ tickets closed. Labour cost per ticket = average minutes per ticket × loaded cost per minute.",
        "Block hours breakage: prepaid hours that expire unused. Effective rate on used hours = block price ÷ hours expected to be used; it is higher than the block's nominal rate when some hours go unused.",
        "Onboarding fee amortisation: onboarding fee ÷ contract months, showing what the fee is worth per month if you spread it over the contract instead of charging it up front.",
        "The targets in Settings are your own; this workbook does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 84
    label(ws, "A1", "Settings", bold=True).font = TITLE
    staff = [(R_TECHS, "Technician count", 3, INT0, True, "Example default 3. Technicians who work client tickets and projects (full-time equivalents)."),
             (R_SALARY, "Average technician salary per month", 4500, MONEY, True, "Example default 4,500 a month. Gross salary per technician before payroll taxes and benefits."),
             (R_TAXES, "Payroll taxes and benefits %", 20, PCT, True, "Default 20. Employer payroll taxes, pension, health cover and other benefits on top of salary."),
             (R_PROD_HRS, "Productive hours per technician per month", 130, HRS, True, "Default 130 of about 173 paid hours: the rest is holidays, sickness, training, meetings and admin.")]
    for r, name, val, fmt, key, note in staff:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", note, muted=True)
    overhead = [("Office (monthly)", 1200, "Example default 1,200 a month. Rent, utilities and cleaning."),
                ("Insurance (monthly)", 300, "Example default 300 a month. Professional indemnity, cyber and liability cover."),
                ("Vehicles and travel (monthly)", 400, "Example default 400 a month. Fuel, leases, parking and travel to client sites."),
                ("Admin and sales staff (monthly)", 3500, "Example default 3,500 a month. Non-technical staff: service desk coordinator, sales, accounts."),
                ("Software not billed to clients (monthly)", 250, "Example default 250 a month. Accounting, office suite, phones; client tools belong on the Tool Stack."),
                ("Other overhead (monthly)", 300, "Example default 300 a month. Marketing, training, bank charges, anything else.")]
    for i, (name, val, note) in enumerate(overhead, start=R_PROD_HRS + 1):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, MONEY, key=True); label(ws, f"C{i}", note, muted=True)
    label(ws, f"A{R_OVERHEAD}", "Total monthly overhead", bold=True)
    fx(ws, f"B{R_OVERHEAD}", f"=SUM(B{R_PROD_HRS + 1}:B{R_OVERHEAD - 1})", MONEY, bold=True)
    label(ws, f"C{R_OVERHEAD}", "formula: sum of the six overhead rows above", muted=True)
    rows = [(R_TARGET, "Target margin % on price", 30, PCT, True, "Default 30. Your own target gross margin on each contract's monthly fee (profit ÷ price, not markup on cost)."),
            (R_BLOCK_DISC, "Block hours discount %", 10, PCT, True, "Default 10. Discount off the hourly rate for prepaid block hours; each block can override it."),
            (R_UNUSED, "Expected unused block hours %", 15, PCT, False, "Default 15. Share of prepaid block hours you expect to expire unused; each block can override it."),
            (R_RATE, "Default hourly rate for billed work", 120, MONEY, True, "Default 120. Your rate for out-of-contract work, blocks and onboarding; keep it well above the loaded hourly cost.")]
    for r, name, val, fmt, key, note in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", note, muted=True)
    ws[f"B{R_PROD_HRS}"].comment = Comment("Assumption: about 173 paid hours a month, less holidays, sickness, training, meetings and internal admin, leaves about 130 productive hours.", "Knackdesk")
    ws[f"B{R_TAXES}"].comment = Comment("Assumption: 20% is a placeholder for employer taxes and benefits; use your payroll figures.", "Knackdesk")
    ws[f"B{R_UNUSED}"].comment = Comment("Assumption: some prepaid hours expire unused (breakage). 15% is a placeholder; set it from your own expired blocks.", "Knackdesk")
    ws[f"B{R_TARGET}"].comment = Comment("Your own target, not an industry benchmark. Margin is profit ÷ price.", "Knackdesk")
    derived = [(R_TECH_COST, "Loaded technician cost per month", f"=ROUND(B{R_TECHS}*B{R_SALARY}*(1+B{R_TAXES}/100),2)", MONEY,
                "formula: technician count × salary × (1 + payroll taxes and benefits %)"),
               (R_TOTAL_HRS, "Total productive hours per month", f"=ROUND(B{R_TECHS}*B{R_PROD_HRS},2)", "0.00",
                "formula: technician count × productive hours per technician"),
               (R_HOURLY, "Loaded hourly cost", f"=IFERROR(ROUND((B{R_TECH_COST}+B{R_OVERHEAD})/B{R_TOTAL_HRS},2),0)", MONEY,
                "formula: (loaded technician cost + total overhead) ÷ productive hours; what each productive hour costs you"),
               (R_MINUTE, "Rate per minute", f"=ROUND(B{R_HOURLY}/60,4)", "0.0000", "formula: loaded hourly cost ÷ 60")]
    rows_block(ws, derived, bold_rows=(R_TECH_COST, R_TOTAL_HRS, R_HOURLY))


def build_tools(ws):
    lo, hi = TOOL_ROWS; tot = hi + 1
    label(ws, "A1", "Tool Stack", bold=True).font = TITLE
    label(ws, "A2", "One row per tool you pay for to serve clients. Per-user, per-device and per-site tools are shared across contracts by users + devices; fixed tools stay in overhead.", muted=True)
    head(ws, 4, ["Tool", "Billing basis", "Unit cost per month", "Units in use", "Monthly cost"], [30, 14, 14, 12, 14])
    dv = DataValidation(type="list", formula1=BASES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": MONEY, "D": INT0}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCD", fm); dv.add(f"B{r}")
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(C{r}*D{r},2))', MONEY)
    sample = [("RMM agent", "per device", 3, 180), ("EDR / antivirus", "per device", 4.5, 180),
              ("Email security", "per user", 2, 150), ("Backup", "per device", 6, 60),
              ("PSA (per technician seat)", "fixed", 45, 3), ("Documentation platform", "fixed", 120, 1),
              ("Password manager", "per user", 3, 150), ("Remote support", "fixed", 60, 1)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCD", row): inp(ws, f"{col}{r}", val, fm.get(col), key=col in "CD")
    ws[f"C{lo}"].comment = Comment("Example prices only. Use your distributor or vendor invoices.", "Knackdesk")
    ws[f"B{lo + 4}"].comment = Comment("PSA is billed per technician, not per client, so it is marked fixed and stays out of the contract allocation.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"D{tot}", f"=SUM(D{lo}:D{hi})", INT0, bold=True)
    fx(ws, f"E{tot}", f"=SUM(E{lo}:E{hi})", MONEY, bold=True)
    rows_block(ws, [(tot + 2, "Tool stack total per month", f"=E{tot}", MONEY, None),
                    (tot + 3, "Fixed tool cost per month", f'=SUMIF(B{lo}:B{hi},"fixed",E{lo}:E{hi})', MONEY, None),
                    (tot + 4, "Per-unit tool cost per month", f"=B{tot + 2}-B{tot + 3}", MONEY, None)], bold_rows=(tot + 2, tot + 3, tot + 4))
    label(ws, f"C{tot + 3}", "basis = fixed; counted in Tickets cost, not allocated to contracts", muted=True)
    label(ws, f"C{tot + 4}", "total − fixed; shared across contracts by users + devices", muted=True)
    ws.freeze_panes = "B5"


def build_contracts(ws):
    lo, hi = CONTRACT_ROWS; tot = hi + 1
    label(ws, "A1", "Contracts", bold=True).font = TITLE
    label(ws, "A2", "One row per client contract. Labour uses the loaded hourly cost from Settings; shared tools come from the Tool Stack per-unit total, split by users + devices.", muted=True)
    head(ws, 4, ["Client", "Plan", "Users", "Devices", "Monthly fee", "Support hours per month", "Dedicated tool costs per month",
                 "Onboarding fee charged", "Contract months", "Users + devices", "Shared tool cost allocated", "Labour cost",
                 "Total cost", "Gross profit", "Margin %", "Fee per user", "Cost per user", "Price per user at target margin",
                 "Break-even hours", "Hours over (+) / under (−) break-even", "Contract value"],
         [26, 11, 8, 8, 12, 11, 12, 12, 10, 10, 12, 12, 12, 12, 10, 10, 10, 13, 11, 13, 13])
    dv = DataValidation(type="list", formula1=PLANS, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": INT0, "D": INT0, "E": MONEY, "F": HRS, "G": MONEY, "H": MONEY, "I": INT0}
    rate = f"{S}{R_HOURLY}"
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFGHI", fm); dv.add(f"B{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"J{r}", g.format(f"C{r}+D{r}"), INT0)
        fx(ws, f"K{r}", g.format(f"IFERROR(ROUND({TS_UNIT}/$J${tot}*J{r},2),0)"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(F{r}*{rate},2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"K{r}+G{r}+L{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(E{r}-M{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/E{r}*100,2),0)"), PCT)
        fx(ws, f"P{r}", g.format(f"IFERROR(ROUND(E{r}/C{r},2),0)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(M{r}/C{r},2),0)"), MONEY)
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(Q{r}/(1-{S}{R_TARGET}/100),2),0)"), MONEY)
        fx(ws, f"S{r}", g.format(f"IFERROR(ROUND((E{r}-K{r}-G{r})/{rate},2),0)"), "0.00")
        fx(ws, f"T{r}", g.format(f"ROUND(F{r}-S{r},2)"), '0.00;-0.00;0.00')
        fx(ws, f"U{r}", g.format(f"ROUND(E{r}*I{r}+H{r},2)"), MONEY)
    sample = [("Harbour Dental Group", "per user", 25, 30, 2750, 14, 0, 1500, 24),
              ("Northgate Logistics", "hybrid", 60, 75, 6200, 38, 150, 3000, 36),
              ("Bluewater Law LLP", "per user", 18, 22, 1980, 12, 80, 1200, 24),
              ("Ridge Architects", "per device", 12, 20, 950, 16, 0, 800, 12),
              ("Summit Accounting", "per user", 30, 35, 3150, 18, 60, 1800, 36),
              ("Oakfield Veterinary", "per device", 15, 28, 1250, 10, 0, 900, 12),
              ("Kestrel Manufacturing", "hybrid", 40, 90, 3200, 45, 200, 2500, 36),
              ("Pinecrest School Trust", "per user", 50, 70, 4800, 30, 120, 2000, 24)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCDEFGHI", row): inp(ws, f"{col}{r}", val, fm.get(col), key=col in "EF")
    ws[f"F{lo}"].comment = Comment("Actual hours from your PSA time entries, or your estimate for a new contract. Include tickets, maintenance and reviews.", "Knackdesk")
    ws[f"G{lo}"].comment = Comment("Tools billed only for this client (a line-of-business app licence, a dedicated firewall subscription). Shared tools come from the Tool Stack.", "Knackdesk")
    ws[f"K{lo}"].comment = Comment("Assumption: per-unit tools are shared in proportion to users + devices across all contracts listed here.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" contracts"', bold=True)
    for col, fmt in (("C", INT0), ("D", INT0), ("E", MONEY), ("F", HRS), ("G", MONEY), ("H", MONEY), ("J", INT0),
                     ("K", MONEY), ("L", MONEY), ("M", MONEY), ("N", MONEY), ("U", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(N{tot}/E{tot}*100,2),0)", PCT, bold=True)
    ws[f"O{tot}"].comment = Comment("Overall margin is recomputed from the totals, not an average of the contract margins.", "Knackdesk")
    label(ws, f"A{tot+1}", "Total row: monthly recurring revenue (column E), monthly cost (column M), gross profit (column N) and overall margin % (column O).", muted=True)
    ws.freeze_panes = "B5"


def build_tickets(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1
    label(ws, "A1", "Tickets", bold=True).font = TITLE
    label(ws, "A2", "One row per month, filled top down. Costs default from Settings and the Tool Stack; type your own in columns E and H to override them for that month.", muted=True)
    head(ws, 4, ["Month", "Tickets closed", "Average minutes per ticket", "Technician cost (Settings)", "or type your own",
                 "Technician cost used", "Tool and overhead cost (Settings + Tool Stack)", "or type your own", "Tool and overhead cost used",
                 "Total cost", "Cost per ticket", "Labour cost per ticket", "Tickets per technician", "Labour hours on tickets",
                 "Share of productive hours on tickets %"],
         [12, 10, 11, 13, 12, 13, 15, 12, 13, 13, 11, 11, 11, 11, 13])
    fm = {"A": MONTH, "B": INT0, "C": HRS, "E": MONEY, "H": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEH", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_TECH_COST}"), MONEY)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), MONEY)
        fx(ws, f"G{r}", g.format(f"{TS_TOTAL}+{S}{R_OVERHEAD}"), MONEY)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), MONEY)
        fx(ws, f"J{r}", g.format(f"F{r}+I{r}"), MONEY)
        fx(ws, f"K{r}", g.format(f"IFERROR(ROUND(J{r}/B{r},2),0)"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(C{r}*{S}{R_MINUTE},2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"IFERROR(ROUND(B{r}/{S}{R_TECHS},2),0)"), "0.00")
        fx(ws, f"N{r}", g.format(f"ROUND(B{r}*C{r}/60,2)"), "0.00")
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/{S}{R_TOTAL_HRS}*100,2),0)"), PCT)
    year = dt.date.today().year
    sample = [(310, 32, None), (295, 35, None), (340, 30, 17500), (280, 34, None), (325, 31, None), (300, 33, None)]
    for i, (tickets, minutes, tech_override) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        inp(ws, f"B{r}", tickets, INT0, key=(i == 0)); inp(ws, f"C{r}", minutes, HRS, key=(i == 0))
        if tech_override: inp(ws, f"E{r}", tech_override, MONEY)
    ws[f"C{lo}"].comment = Comment("Average logged time per closed ticket from your PSA, including follow-up.", "Knackdesk")
    ws[f"E{lo + 2}"].comment = Comment("Example override: a month with overtime or a contractor. Leave empty to use the Settings figure.", "Knackdesk")
    ws[f"G{lo}"].comment = Comment("Default = whole Tool Stack (fixed and per-unit tools) + total overhead from Settings.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("B", INT0), ("F", MONEY), ("I", MONEY), ("J", MONEY), ("N", "0.00")):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"C{tot}", f"=IFERROR(ROUND(N{tot}*60/B{tot},1),0)", HRS, bold=True)
    fx(ws, f"K{tot}", f"=IFERROR(ROUND(J{tot}/B{tot},2),0)", MONEY, bold=True)
    fx(ws, f"L{tot}", f"=ROUND(C{tot}*{S}{R_MINUTE},2)", MONEY, bold=True)
    fx(ws, f"M{tot}", f"=IFERROR(ROUND(B{tot}/{S}{R_TECHS},2),0)", "0.00", bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(N{tot}/({S}{R_TOTAL_HRS}*COUNT(A{lo}:A{hi}))*100,2),0)", PCT, bold=True)
    ws[f"K{tot}"].comment = Comment("Totals recompute from the column totals, not an average of the monthly figures.", "Knackdesk")
    ws.freeze_panes = "B5"


def build_blocks(ws):
    lo, hi = BLOCK_ROWS; tot = hi + 1; olo, ohi = ONBOARD_ROWS; otot = ohi + 1
    label(ws, "A1", "Blocks & Onboarding", bold=True).font = TITLE
    label(ws, "A2", "Block hours quotes: rate, discount and expected unused % default from Settings; type your own in the 'or type your own' columns to override them.", muted=True)
    head(ws, 4, ["Client", "Hours in block", "Default rate (Settings)", "or type your own rate", "Rate used",
                 "Default discount % (Settings)", "or type your own discount %", "Discount % used", "Expiry months",
                 "Default unused % (Settings)", "or type your own unused %", "Unused % used", "List price", "Block price",
                 "Effective rate", "Expected used hours", "Effective rate on used hours", "Per month"],
         [26, 10, 12, 12, 11, 12, 12, 11, 10, 12, 12, 11, 12, 12, 11, 11, 13, 12])
    fm = {"B": HRS, "D": MONEY, "G": PCT, "I": INT0, "K": PCT}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABDGIK", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"C{r}", g.format(f"{S}{R_RATE}"), MONEY)
        fx(ws, f"E{r}", g.format(f'IF(D{r}<>"",D{r},C{r})'), MONEY)
        fx(ws, f"F{r}", g.format(f"{S}{R_BLOCK_DISC}"), PCT)
        fx(ws, f"H{r}", g.format(f'IF(G{r}<>"",G{r},F{r})'), PCT)
        fx(ws, f"J{r}", g.format(f"{S}{R_UNUSED}"), PCT)
        fx(ws, f"L{r}", g.format(f'IF(K{r}<>"",K{r},J{r})'), PCT)
        fx(ws, f"M{r}", g.format(f"ROUND(B{r}*E{r},2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(M{r}*(1-H{r}/100),2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/B{r},2),0)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(B{r}*(1-L{r}/100),2)"), "0.00")
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(N{r}/P{r},2),0)"), MONEY)
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(N{r}/I{r},2),0)"), MONEY)
    blocks = [("Coastal Realty", 20, None, None, 6, None), ("Meridian Studio", 40, 110, 15, 12, None),
              ("Fenwick & Sons", 10, None, None, 3, 5)]
    for i, (name, hrs, rate, disc, months, unused) in enumerate(blocks):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", hrs, HRS, key=True); inp(ws, f"D{r}", rate, MONEY)
        inp(ws, f"G{r}", disc, PCT); inp(ws, f"I{r}", months, INT0, key=True); inp(ws, f"K{r}", unused, PCT)
    ws[f"I{lo}"].comment = Comment("Months before unused hours expire. Per month = block price ÷ expiry months.", "Knackdesk")
    ws[f"K{lo}"].comment = Comment("Leave empty to use the Settings figure; any number here takes precedence for this block.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("B", HRS), ("M", MONEY), ("N", MONEY), ("P", "0.00")):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)

    label(ws, f"A{olo - 3}", "Onboarding quotes", bold=True)
    label(ws, f"A{olo - 2}", "The default rate comes from Settings; type your own rate in column D to override it for that quote.", muted=True)
    head(ws, olo - 1, ["Client", "Onboarding hours", "Default rate (Settings)", "or type your own rate", "Rate used",
                       "Setup costs", "Margin on costs %", "Contract months", "Monthly fee", "Labour", "Costs with margin",
                       "Onboarding fee", "Amortised per month", "Contract value", "Fee share of contract %"])
    fm = {"B": HRS, "D": MONEY, "F": MONEY, "G": PCT, "H": INT0, "I": MONEY}
    for r in range(olo, ohi + 1):
        blank(ws, r, "ABDFGHI", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"C{r}", g.format(f"{S}{R_RATE}"), MONEY)
        fx(ws, f"E{r}", g.format(f'IF(D{r}<>"",D{r},C{r})'), MONEY)
        fx(ws, f"J{r}", g.format(f"ROUND(B{r}*E{r},2)"), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND(F{r}*(1+G{r}/100),2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"J{r}+K{r}"), MONEY)
        fx(ws, f"M{r}", g.format(f"IFERROR(ROUND(L{r}/H{r},2),0)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(I{r}*H{r}+L{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(L{r}/N{r}*100,2),0)"), PCT)
    quotes = [("Harbour Dental Group", 16, None, 600, 15, 24, 2750), ("Brightwell Clinic", 24, None, 1200, 20, 36, 3900),
              ("Vega Creative", 8, 100, 250, 10, 12, 900)]
    for i, (name, hrs, rate, setup, margin, months, fee) in enumerate(quotes):
        r = olo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", hrs, HRS, key=True); inp(ws, f"D{r}", rate, MONEY)
        inp(ws, f"F{r}", setup, MONEY); inp(ws, f"G{r}", margin, PCT); inp(ws, f"H{r}", months, INT0)
        inp(ws, f"I{r}", fee, MONEY)
    ws[f"B{olo}"].comment = Comment("Hours to document the environment, deploy agents, migrate accounts and train staff. Estimate generously.", "Knackdesk")
    ws[f"F{olo}"].comment = Comment("Hardware, licences and third-party costs you pay to onboard the client.", "Knackdesk")
    label(ws, f"A{otot}", "Total", bold=True)
    for col, fmt in (("B", HRS), ("J", MONEY), ("K", MONEY), ("L", MONEY), ("N", MONEY)):
        fx(ws, f"{col}{otot}", f"=SUM({col}{olo}:{col}{ohi})", fmt, bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    lo, hi = CONTRACT_ROWS; tot = hi + 1; t_lo, t_hi = MONTH_ROWS; t_tot = t_hi + 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 64
    label(ws, "A1", "Summary", bold=True).font = TITLE
    m = lambda col: f"{CT}${col}${lo}:${col}${hi}"
    rows = [(2, "Loaded hourly cost", f"={S}{R_HOURLY}", MONEY, "from Settings"),
            (3, "Technicians", f"={S}{R_TECHS}", INT0, "from Settings"),
            (4, "Productive hours per month", f"={S}{R_TOTAL_HRS}", "0.00", "from Settings"),
            (5, "Tool stack monthly total", f"={TS_TOTAL}", MONEY, "from the Tool Stack"),
            (6, "Fixed tool cost (not allocated to contracts)", f"={TS_FIXED}", MONEY, "from the Tool Stack; counted in the Tickets cost"),
            (7, "Contracts", f'=COUNTA({m("A")})', INT0, None),
            (8, "Monthly recurring revenue", f"={CT}E{tot}", MONEY, "sum of monthly fees"),
            (9, "Total monthly contract cost", f"={CT}M{tot}", MONEY, "shared tools + dedicated tools + labour"),
            (10, "Monthly gross profit", f"={CT}N{tot}", MONEY, None),
            (11, "Overall margin %", f"={CT}O{tot}", PCT, "gross profit ÷ monthly recurring revenue"),
            (12, "Contracts below target margin", f'=SUMPRODUCT(({m("A")}<>"")*({m("O")}<{S}{R_TARGET}))', INT0, "margin % below your target in Settings"),
            (13, "Most profitable contract", f'=IFERROR(INDEX({m("A")},MATCH(MAX({m("N")}),{m("N")},0)),"")', None, "highest monthly gross profit"),
            (14, "Least profitable contract", f'=IFERROR(INDEX({m("A")},MATCH(MIN({m("N")}),{m("N")},0)),"")', None, "lowest monthly gross profit; re-price first"),
            (15, "Latest month cost per ticket", f"=IFERROR(INDEX({TK}$K${t_lo}:$K${t_hi},MATCH(MAX({TK}$A${t_lo}:$A${t_hi}),{TK}$A${t_lo}:$A${t_hi},0)),0)", MONEY, "the most recent month on Tickets"),
            (16, "Year-to-date tickets", f"={TK}B{t_tot}", INT0, "Tickets totals"),
            (17, "Year-to-date cost per ticket", f"={TK}K{t_tot}", MONEY, "Tickets totals"),
            (18, "Block hours revenue", f"={BO}N{BLOCK_ROWS[1] + 1}", MONEY, "sum of block prices"),
            (19, "Onboarding revenue", f"={BO}L{ONBOARD_ROWS[1] + 1}", MONEY, "sum of onboarding fees")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))


README = """# MSP Contract Pricing Workbook

Thank you for buying the workbook. Open `msp-contract-pricing-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** technician count, salary, payroll taxes and benefits, productive hours, overhead line by line, target margin, block hours discount, expected unused block hours and default hourly rate; your loaded hourly cost follows.
2. **Tool Stack:** each tool with billing basis, unit cost and units in use; monthly cost, stack total, fixed tool cost and per-unit tool cost follow.
3. **Contracts:** clients with plan, users, devices, monthly fee, support hours, dedicated tools, onboarding fee and contract months; shared tool cost, labour, total cost, gross profit, margin %, cost per user, price per user at target margin and break-even hours follow.
4. **Tickets:** tickets closed and average minutes per month, with technician and tool and overhead cost (or your own); cost per ticket, labour cost per ticket, tickets per technician and share of productive hours on tickets follow.
5. **Blocks & Onboarding:** block hours quotes with block price, effective rate and effective rate on used hours; onboarding quotes with the onboarding fee, amortised per month and share of contract value.
6. **Summary:** loaded hourly cost, monthly recurring revenue, cost, gross profit and overall margin, contracts below target, most and least profitable contract, cost per ticket, block and onboarding revenue.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_tools(wb.create_sheet("Tool Stack"))
    build_contracts(wb.create_sheet("Contracts")); build_tickets(wb.create_sheet("Tickets"))
    build_blocks(wb.create_sheet("Blocks & Onboarding")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
