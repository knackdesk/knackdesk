"""Build the Loan & Lease Comparison Workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
SCHED_FIRST = 14; SCHED_ROWS = 360; SCHED_LAST = SCHED_FIRST + SCHED_ROWS - 1
OFFERS = "BCDE"


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Loan & Lease Comparison Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Loan Compare: one column per offer (up to four) with lender, amount, annual rate, term in months and upfront fees. Payment, totals, true APR and the cost ranking are calculated.",
        "2. Schedule: enter one loan (amount, rate, term, start date) and an optional extra monthly payment. The table shows every month's interest, principal and balance until payoff, with the payoff month and the interest saved by the extra payment.",
        "3. Lease vs Buy: the lease terms on one side and the purchase, deposit, loan and resale value on the other; both totals over the term and the cheaper option.",
        "4. DSCR: net operating income and current annual debt service give today's coverage ratio; add a proposed loan payment to see the ratio after it, and the largest payment a lender's target allows.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Payment = PMT(rate ÷ 12, term, amount). Total interest = payment × term − amount. Total cost = total interest + fees.",
        "APR = RATE(term, payment, amount − fees) × 12: the rate that would give the same payments on the money actually received.",
        "Schedule: interest = opening balance × rate ÷ 12; the last scheduled month pays off whatever remains, so cent rounding never leaves a residue.",
        "Lease total = payments + buyout − resale (resale only when you buy the asset out). Buy total = deposit + loan payments − resale. DSCR = net operating income ÷ annual debt service.",
        "Fixed rates and equal monthly payments are assumed. Arithmetic only; not financial, credit or tax advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_compare(ws):
    ws.column_dimensions["A"].width = 34
    for c in OFFERS: ws.column_dimensions[c].width = 16
    label(ws, "A1", "Loan Compare", bold=True).font = TITLE
    label(ws, "A2", "One column per offer. Leave the amount blank for unused columns.", muted=True)
    head(ws, 4, ["", "Offer 1", "Offer 2", "Offer 3", "Offer 4"])
    labels = {5: "Lender", 6: "Loan amount", 7: "Annual rate (%)", 8: "Term (months)", 9: "Upfront fees", 10: "Monthly payment", 11: "Total repaid", 12: "Total interest", 13: "Total cost (interest + fees)", 14: "APR (%)", 15: "Amount received after fees", 16: "Rank by total cost (1 = cheapest)"}
    for r, t in labels.items(): label(ws, f"A{r}", t, bold=r >= 10)
    for c in OFFERS:
        for r, fmt in ((6, MONEY), (7, "0.00"), (8, "0"), (9, MONEY)): ws[f"{c}{r}"].font = BLUE; ws[f"{c}{r}"].number_format = fmt
        ws[f"{c}5"].font = BLUE
        fx(ws, f"{c}10", f'=IF({c}6="","",ROUND(-PMT({c}7/100/12,{c}8,{c}6),2))', MONEY, bold=True)
        fx(ws, f"{c}11", f'=IF({c}6="","",{c}10*{c}8)', MONEY)
        fx(ws, f"{c}12", f'=IF({c}6="","",{c}11-{c}6)', MONEY)
        fx(ws, f"{c}13", f'=IF({c}6="","",{c}12+{c}9)', MONEY, bold=True)
        fx(ws, f"{c}14", f'=IF({c}6="","",IF({c}9=0,{c}7,ROUND(RATE({c}8,-{c}10,{c}6-{c}9)*12*100,2)))', "0.00", bold=True)
        fx(ws, f"{c}15", f'=IF({c}6="","",{c}6-{c}9)', MONEY)
        fx(ws, f"{c}16", f'=IF({c}6="","",COUNTIF($B$13:$E$13,"<"&{c}13)+1)', "0", bold=True)
    samples = [("Bank A", 50000, 6, 60, 500), ("Bank B", 50000, 5.5, 60, 1250), ("Online lender", 50000, 7.2, 48, 0)]
    for i, (name, amt, rate, term, fees) in enumerate(samples):
        c = OFFERS[i]; inp(ws, f"{c}5", name, key=(i == 0)); inp(ws, f"{c}6", amt, MONEY, key=(i == 0)); inp(ws, f"{c}7", rate, "0.00", key=(i == 0)); inp(ws, f"{c}8", term, "0", key=(i == 0)); inp(ws, f"{c}9", fees, MONEY, key=(i == 0))


def build_schedule(ws):
    ws.column_dimensions["A"].width = 30
    for c, w in zip("BCDEFG", (12, 14, 12, 12, 12, 14)): ws.column_dimensions[c].width = w
    label(ws, "A1", "Schedule", bold=True).font = TITLE
    label(ws, "A2", "Loan amount"); inp(ws, "B2", 20000, MONEY, key=True)
    label(ws, "A3", "Annual rate (%)"); inp(ws, "B3", 7, "0.00", key=True)
    label(ws, "A4", "Term (months, up to 360)"); inp(ws, "B4", 48, "0", key=True)
    label(ws, "A5", "Extra payment each month"); inp(ws, "B5", 150, MONEY, key=True)
    label(ws, "A6", "First payment date"); inp(ws, "B6", dt.date.today().replace(day=1), DATE, key=True)
    label(ws, "A8", "Scheduled monthly payment", bold=True); fx(ws, "B8", "=ROUND(-PMT(B3/100/12,B4,B2),2)", MONEY, bold=True)
    label(ws, "A9", "Interest without extra payments", bold=True); fx(ws, "B9", "=ROUND(B8*B4-B2,2)", MONEY, bold=True)
    label(ws, "A10", "Months to payoff", bold=True); fx(ws, "B10", f"=MAX(A{SCHED_FIRST}:A{SCHED_LAST})", "0", bold=True)
    label(ws, "A11", "Interest paid", bold=True); fx(ws, "B11", f"=SUM(E{SCHED_FIRST}:E{SCHED_LAST})", MONEY, bold=True)
    label(ws, "A12", "Interest saved by extra payments", bold=True); fx(ws, "B12", "=ROUND(B9-B11,2)", MONEY, bold=True)
    label(ws, "C10", "Payoff month", bold=True); fx(ws, "D10", '=IF(B10=0,"",EDATE(B6,B10-1))', "mmm yyyy", bold=True)
    label(ws, "C11", "Months saved", bold=True); fx(ws, "D11", "=MAX(0,B4-B10)", "0", bold=True)
    head(ws, 13, ["#", "Date", "Opening balance", "Payment", "Interest", "Principal", "Closing balance"])
    for r in range(SCHED_FIRST, SCHED_LAST + 1):
        p = r - 1
        fx(ws, f"A{r}", '=IF($B$2="","",1)' if r == SCHED_FIRST else f'=IF(A{p}="","",IF(G{p}<=0.005,"",A{p}+1))', "0")
        fx(ws, f"B{r}", f'=IF(A{r}="","",EDATE($B$6,A{r}-1))', DATE)
        fx(ws, f"C{r}", '=IF(A{0}="","",$B$2)'.format(r) if r == SCHED_FIRST else f'=IF(A{r}="","",G{p})', MONEY)
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(C{r}*$B$3/100/12,2))', MONEY)
        fx(ws, f"D{r}", f'=IF(A{r}="","",IF(A{r}>=$B$4,C{r}+E{r},MIN($B$8+$B$5,C{r}+E{r})))', MONEY)
        fx(ws, f"F{r}", f'=IF(A{r}="","",D{r}-E{r})', MONEY)
        fx(ws, f"G{r}", f'=IF(A{r}="","",ROUND(C{r}-F{r},2))', MONEY)
    ws.freeze_panes = f"A{SCHED_FIRST}"


def build_lease(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "Lease vs Buy", bold=True).font = TITLE
    rows = [(3, "Purchase price", 30000, MONEY), (4, "Deposit paid upfront", 0, MONEY), (5, "Loan annual rate (%)", 6, "0.00"), (6, "Loan term (months)", 60, "0"), (7, "Expected resale value at the end", 9000, MONEY),
            (8, "Lease payment per month", 550, MONEY), (9, "Lease term (months)", 60, "0"), (10, "Lease buyout at the end (0 if handed back)", 3000, MONEY)]
    for r, name, val, fmt in rows: label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=True)
    label(ws, "A11", "Loan monthly payment", bold=True); fx(ws, "B11", "=ROUND(-PMT(B5/100/12,B6,B3-B4),2)", MONEY, bold=True)
    label(ws, "A12", "Buy: total cost over the term", bold=True); fx(ws, "B12", "=ROUND(B4+B11*B6-B7,2)", MONEY, bold=True); label(ws, "C12", "deposit + loan payments − resale", muted=True)
    label(ws, "A13", "Lease: total cost over the term", bold=True); fx(ws, "B13", "=ROUND(B8*B9+B10-IF(B10>0,B7,0),2)", MONEY, bold=True); label(ws, "C13", "payments + buyout − resale (only if bought out)", muted=True)
    label(ws, "A14", "Cheaper option", bold=True); fx(ws, "B14", '=IF(B12=B13,"Same",IF(B12<B13,"Buy","Lease"))', None, bold=True)
    label(ws, "A15", "Difference", bold=True); fx(ws, "B15", "=ABS(B12-B13)", MONEY, bold=True)
    label(ws, "A16", "Buy: average cost per month", bold=True); fx(ws, "B16", '=IFERROR(ROUND(B12/B6,2),"")', MONEY, bold=True)
    label(ws, "A17", "Lease: average cost per month", bold=True); fx(ws, "B17", '=IFERROR(ROUND(B13/B9,2),"")', MONEY, bold=True)
    label(ws, "A19", "Tax treatment, maintenance and usage limits are not included; weigh them after the totals.", muted=True)


def build_dscr(ws):
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "DSCR", bold=True).font = TITLE
    label(ws, "A2", "Net operating income for the year"); inp(ws, "B2", 120000, MONEY, key=True); label(ws, "C2", "revenue − operating expenses, before interest, principal and tax", muted=True)
    label(ws, "A3", "Current annual debt service"); inp(ws, "B3", 80000, MONEY, key=True); label(ws, "C3", "all loan principal and interest paid in a year", muted=True)
    label(ws, "A4", "Lender's target ratio"); inp(ws, "B4", 1.25, "0.00", key=True)
    label(ws, "A5", "DSCR today", bold=True); fx(ws, "B5", '=IFERROR(ROUND(B2/B3,2),"")', "0.00", bold=True)
    label(ws, "A6", "Meets target", bold=True); fx(ws, "B6", '=IF(B5="","",IF(B5>=B4,"Yes","No"))', None, bold=True)
    label(ws, "A7", "Largest annual debt service at the target", bold=True); fx(ws, "B7", "=ROUND(MAX(0,B2)/B4,2)", MONEY, bold=True)
    label(ws, "A8", "Largest monthly payment at the target", bold=True); fx(ws, "B8", "=ROUND(B7/12,2)", MONEY, bold=True)
    label(ws, "A9", "Headroom in annual debt service", bold=True); fx(ws, "B9", "=ROUND(B7-B3,2)", MONEY, bold=True); label(ws, "C9", "negative = reduction needed to reach the target", muted=True)
    label(ws, "A11", "Proposed new loan: monthly payment"); inp(ws, "B11", "='Loan Compare'!B10", MONEY, note="Defaults to Offer 1 on the Loan Compare sheet; overwrite if needed.")
    label(ws, "A12", "Annual debt service after the new loan", bold=True); fx(ws, "B12", "=B3+N(B11)*12", MONEY, bold=True)
    label(ws, "A13", "DSCR after the new loan", bold=True); fx(ws, "B13", '=IFERROR(ROUND(B2/B12,2),"")', "0.00", bold=True)
    label(ws, "A14", "Meets target after the new loan", bold=True); fx(ws, "B14", '=IF(B13="","",IF(B13>=B4,"Yes","No"))', None, bold=True)


README = """# Loan & Lease Comparison Workbook

Thank you for buying the workbook. Open `loan-lease-comparison-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Loan Compare:** up to four offers side by side with payment, total cost, true APR and a cost ranking.
2. **Schedule:** one loan's month-by-month amortisation with optional extra payments, payoff month and interest saved.
3. **Lease vs Buy:** both totals over the same period and the cheaper option.
4. **DSCR:** coverage ratio today and after a proposed loan, with the largest payment a target allows.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Fixed rates and equal payments assumed. Arithmetic only, not financial, credit or tax advice. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_compare(wb.create_sheet("Loan Compare")); build_schedule(wb.create_sheet("Schedule")); build_lease(wb.create_sheet("Lease vs Buy")); build_dscr(wb.create_sheet("DSCR"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "loan-lease-comparison-workbook.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "loan-lease-comparison-workbook.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "loan-lease-comparison-workbook.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
