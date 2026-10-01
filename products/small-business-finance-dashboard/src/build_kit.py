"""Build the Small Business Finance Dashboard workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PL_ROWS = (5, 16); EXP_ROWS = (5, 20); EXP_TOTAL = 21
S = "Settings!$B$"; PL = "'Monthly P&L'!"; EX = "Expenses!"


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Small Business Finance Dashboard", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: the year and the share of pre-tax profit you want to set aside for tax. The percentage is yours to choose; the workbook does not know your tax rate.",
        "2. Expenses: one row per operating expense category, one column per month. Monthly totals feed the P&L automatically; yearly totals and each category's share are calculated.",
        "3. Monthly P&L: enter revenue, cost of sales, interest and owner draws for each month as it closes. Gross, operating and pre-tax profit, margins, the tax set-aside, retained profit and the running total are calculated. Operating expenses default to the Expenses sheet total.",
        "4. Summary: year-to-date totals and margins, average monthly revenue, best and worst months, tax set aside, draws and the largest expense category.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "Blue cells that already contain a formula (operating expenses pulled from the Expenses sheet) may be overwritten with your own figure.",
        "", "DEFINITIONS",
        "Gross profit = revenue − cost of sales. Operating profit = gross profit − operating expenses. Pre-tax profit = operating profit − interest.",
        "Tax set-aside = pre-tax profit × your percentage (zero in a loss month). Retained = pre-tax profit − tax set-aside − owner draws. Margins are each profit ÷ revenue.",
        "Enter figures excluding VAT or sales tax you collect. Arithmetic only; not tax, accounting or financial advice. Support: hello@knackdesk.com"]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Settings", bold=True).font = TITLE
    label(ws, "A2", "Year"); inp(ws, "B2", dt.date.today().year, "0", key=True); label(ws, "C2", "Shown on the P&L and Summary.", muted=True)
    label(ws, "A3", "Tax set-aside (% of pre-tax profit)"); inp(ws, "B3", 25, "0.00", key=True); label(ws, "C3", "Last year's tax and contributions ÷ last year's profit is a reasonable starting point.", muted=True)
    label(ws, "A4", "Business name"); inp(ws, "B4", "Example Studio", key=True); label(ws, "C4", "Optional label.", muted=True)


def build_expenses(ws):
    lo, hi = EXP_ROWS
    label(ws, "A1", "Expenses", bold=True).font = TITLE
    label(ws, "A2", "Operating expenses by category and month, excluding cost of sales (which goes on the P&L). Blank months count as zero.", muted=True)
    head(ws, 4, ["Category"] + MONTHS + ["Year total", "Share %"], [26] + [11] * 12 + [13, 10])
    for r in range(lo, hi + 1):
        ws[f"A{r}"].font = BLUE
        for c in range(2, 14):
            cell = ws.cell(row=r, column=c); cell.font = BLUE; cell.number_format = MONEY
        fx(ws, f"N{r}", f'=IF(A{r}="","",SUM(B{r}:M{r}))', MONEY)
        fx(ws, f"O{r}", f'=IF(A{r}="","",IFERROR(ROUND(N{r}/$N${EXP_TOTAL}*100,2),""))', PCT)
    sample = [("Rent and utilities", 900), ("Salaries and contractors", 4200), ("Software and subscriptions", 260), ("Marketing", 500), ("Insurance", 120), ("Accounting and legal", 150), ("Travel", 80), ("Other", 100)]
    for i, (name, amt) in enumerate(sample):
        r = lo + i; inp(ws, f"A{r}", name, key=(i == 0))
        for c, v in zip("BCD", (amt, amt, round(amt * 1.1, 2))): inp(ws, f"{c}{r}", v, MONEY, key=(i == 0 and c == "B"))
    label(ws, f"A{EXP_TOTAL}", "Total per month", bold=True)
    for c in range(2, 15):
        col = get_column_letter(c); fx(ws, f"{col}{EXP_TOTAL}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    ws.freeze_panes = "B5"


def build_pl(ws):
    lo, hi = PL_ROWS; tot = hi + 1
    label(ws, "A1", "Monthly P&L", bold=True).font = TITLE
    fx(ws, "B1", f'={S}4&" · "&{S}2', None, bold=True)
    label(ws, "A2", "Enter revenue, cost of sales, interest and owner draws each month. Operating expenses come from the Expenses sheet. Leave future months' revenue blank.", muted=True)
    head(ws, 4, ["Month", "Revenue", "Cost of sales", "Gross profit", "Gross margin %", "Operating expenses", "Operating profit", "Operating margin %", "Interest", "Pre-tax profit", "Tax set-aside", "Profit after set-aside", "Margin after set-aside %", "Owner draws", "Retained in business", "Cumulative retained"],
         [8, 12, 12, 12, 11, 13, 13, 11, 10, 12, 12, 13, 12, 12, 13, 13])
    for i, r in enumerate(range(lo, hi + 1)):
        label(ws, f"A{r}", MONTHS[i], bold=True)
        for col in "BCIN": ws[f"{col}{r}"].font = BLUE; ws[f"{col}{r}"].number_format = MONEY
        exp_col = get_column_letter(2 + i)
        fx(ws, f"F{r}", f'=IF(B{r}="","",{EX}{exp_col}{EXP_TOTAL})', MONEY).font = BLUE
        fx(ws, f"D{r}", f'=IF(B{r}="","",B{r}-C{r})', MONEY)
        fx(ws, f"E{r}", f'=IF(B{r}="","",IFERROR(ROUND(D{r}/B{r}*100,2),""))', PCT)
        fx(ws, f"G{r}", f'=IF(B{r}="","",D{r}-F{r})', MONEY)
        fx(ws, f"H{r}", f'=IF(B{r}="","",IFERROR(ROUND(G{r}/B{r}*100,2),""))', PCT)
        fx(ws, f"J{r}", f'=IF(B{r}="","",G{r}-I{r})', MONEY)
        fx(ws, f"K{r}", f'=IF(B{r}="","",ROUND(MAX(0,J{r})*{S}3/100,2))', MONEY)
        fx(ws, f"L{r}", f'=IF(B{r}="","",J{r}-K{r})', MONEY)
        fx(ws, f"M{r}", f'=IF(B{r}="","",IFERROR(ROUND(L{r}/B{r}*100,2),""))', PCT)
        fx(ws, f"O{r}", f'=IF(B{r}="","",L{r}-N{r})', MONEY)
        fx(ws, f"P{r}", f'=IF(B{r}="","",O{r}+N(P{r-1}))' if r > lo else f'=IF(B{r}="","",O{r})', MONEY)
    for i, (rev, cos, intr, draw) in enumerate([(14000, 3500, 50, 3000), (12500, 3100, 50, 3000), (16800, 4200, 50, 3500)]):
        r = lo + i; inp(ws, f"B{r}", rev, MONEY, key=(i == 0)); inp(ws, f"C{r}", cos, MONEY, key=(i == 0)); inp(ws, f"I{r}", intr, MONEY); inp(ws, f"N{r}", draw, MONEY, key=(i == 0))
    label(ws, f"A{tot}", "Year", bold=True)
    for col in "BCDFGIJKLNO": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"E{tot}", f'=IFERROR(ROUND(D{tot}/B{tot}*100,2),"")', PCT, bold=True)
    fx(ws, f"H{tot}", f'=IFERROR(ROUND(G{tot}/B{tot}*100,2),"")', PCT, bold=True)
    fx(ws, f"M{tot}", f'=IFERROR(ROUND(L{tot}/B{tot}*100,2),"")', PCT, bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    lo, hi = PL_ROWS; tot = hi + 1; elo, ehi = EXP_ROWS
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "Summary (year to date)", bold=True).font = TITLE
    rows = [(2, "Revenue", f"={PL}B{tot}", MONEY, None),
            (3, "Cost of sales", f"={PL}C{tot}", MONEY, None),
            (4, "Gross profit", f"={PL}D{tot}", MONEY, None),
            (5, "Gross margin %", f"={PL}E{tot}", PCT, None),
            (6, "Operating expenses", f"={PL}F{tot}", MONEY, None),
            (7, "Operating profit", f"={PL}G{tot}", MONEY, None),
            (8, "Operating margin %", f"={PL}H{tot}", PCT, None),
            (9, "Pre-tax profit", f"={PL}J{tot}", MONEY, "after interest"),
            (10, "Set aside for tax", f"={PL}K{tot}", MONEY, "at the Settings percentage"),
            (11, "Owner draws", f"={PL}N{tot}", MONEY, None),
            (12, "Retained in the business", f"={PL}O{tot}", MONEY, "pre-tax profit − tax set-aside − draws"),
            (13, "Months entered", f'=COUNT({PL}B{lo}:B{hi})', "0", None),
            (14, "Average monthly revenue", f'=IFERROR(ROUND(B2/B13,2),"")', MONEY, None),
            (15, "Best month (revenue)", f'=IFERROR(INDEX({PL}A{lo}:A{hi},MATCH(MAX({PL}B{lo}:B{hi}),{PL}B{lo}:B{hi},0)),"")', None, None),
            (16, "Worst month (revenue)", f'=IFERROR(INDEX({PL}A{lo}:A{hi},MATCH(MIN({PL}B{lo}:B{hi}),{PL}B{lo}:B{hi},0)),"")', None, None),
            (17, "Largest expense category", f'=IFERROR(INDEX({EX}A{elo}:A{ehi},MATCH(MAX({EX}N{elo}:N{ehi}),{EX}N{elo}:N{ehi},0)),"")', None, None),
            (18, "Largest category share of expenses %", f'=IFERROR(ROUND(MAX({EX}N{elo}:N{ehi})/{EX}N{EXP_TOTAL}*100,2),"")', PCT, None),
            (19, "Operating expenses as % of revenue", f'=IFERROR(ROUND(B6/B2*100,2),"")', PCT, None)]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True)
        if note: label(ws, f"C{r}", note, muted=True)


README = """# Small Business Finance Dashboard

Thank you for buying the dashboard. Open `small-business-finance-dashboard.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** the year and the share of pre-tax profit to set aside for tax (your choice; not tax advice).
2. **Expenses:** operating expenses by category and month; totals feed the P&L.
3. **Monthly P&L:** revenue, cost of sales, interest and owner draws per month; profits, margins, tax set-aside, retained profit and the running total are calculated.
4. **Summary:** year-to-date totals, margins, best and worst months and the largest expense category.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not tax, accounting or financial advice. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_expenses(wb.create_sheet("Expenses")); build_pl(wb.create_sheet("Monthly P&L")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "small-business-finance-dashboard.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "small-business-finance-dashboard.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "small-business-finance-dashboard.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
