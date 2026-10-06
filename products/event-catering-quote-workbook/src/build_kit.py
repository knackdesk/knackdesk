"""Build the Event & Catering Quote Workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "event-catering-quote-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"
MENU_ROWS = (9, 20); RENT_ROWS = (24, 33); MIX_ROWS = (6, 8); EVENT_ROWS = (5, 40)
S = "Settings!$B$"  # settings cell prefix
G = "Quote!$B$5"; H = "Quote!$B$6"  # guests, event hours


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Event & Catering Quote Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: your labour rates, staffing ratios, setup and minimum hours, drinks-per-guest estimates, target margin and deposit. Every other sheet reads them.",
        "2. Quote: client, event, guests and hours, then the menu (cost per guest), rentals and other costs. Staffing, bar, total cost, price at your target margin, price per guest, deposit and balance are calculated.",
        "3. Bar: the drinks estimate for the quote, split by beer, wine and spirits with your cost per serving. Shares must total 100.",
        "4. Events: one row per event you quoted or ran, with the quoted price and the actual costs. Profit, margin and outstanding balance are calculated, with totals in row 41.",
        "5. Summary: events tracked, quoted total, cost, profit, margin, average price per guest, events below target margin and money outstanding, read from the Events sheet.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Price at target margin = total cost ÷ (1 − target margin %). A 30% margin means 30% of the price you charge is profit, which is not the same as a 30% markup on cost.",
        "Staff hours = the larger of (event hours + setup and teardown hours) and the minimum hours per staff member. Staff needed = guests ÷ guests per staff, rounded up.",
        "Staffing ratios (one server per 25 guests, one bartender per 50) are common starting points, not rules; plated service, buffets and venues differ.",
        "Drinks per guest defaults (2 in the first hour, 1 each hour after) are estimates; adjust them for your crowd and event type.",
        "Events: profit = quoted price − actual costs; margin = profit ÷ quoted price. Leave lost events' actual costs blank and remove the quoted price if you do not want them in the totals.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [("Server labour rate per hour", 22, MONEY, "What a server costs you per hour, including any employer costs."),
            ("Bartender rate per hour", 25, MONEY, "What a bartender costs you per hour."),
            ("Guests per server", 25, "0", "Common starting point; plated service often needs fewer guests per server."),
            ("Guests per bartender", 50, "0", "Common starting point; a full bar or cocktails may need more bartenders."),
            ("Setup and teardown hours", 2, "0.0", "Added to the event hours for every staff member."),
            ("Minimum hours per staff", 4, "0.0", "Minimum paid shift length."),
            ("Drinks per guest, first hour", 2, "0.0", "Estimate; adjust for your crowd."),
            ("Drinks per guest, each later hour", 1, "0.0", "Estimate; adjust for your crowd."),
            ("Target margin (% of price)", 30, "0.0", "Profit as a share of the price you charge, not a markup on cost."),
            ("Deposit (% of price)", 40, "0.0", "Share of the price collected up front.")]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, fmt, key=True); label(ws, f"C{i}", note, muted=True)


def quote_inputs(ws):
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 16; ws.column_dimensions["D"].width = 16
    label(ws, "A1", "Quote", bold=True).font = TITLE
    y = dt.date.today().year
    for r, name, val, fmt in [(2, "Client", "Sample client", None), (3, "Event name", "Company dinner", None), (4, "Event date", dt.date(y + 1, 6, 20), DATE),
                              (5, "Guests", 100, "0"), (6, "Event hours", 5, "0.0")]:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=True)


def quote_menu(ws):
    lo, hi = MENU_ROWS; tot = hi + 1
    head(ws, 8, ["Menu item", "Cost per guest", "Per-guest total", "Event total"])
    for r in range(lo, hi + 1):
        blank(ws, r, "AB", {"B": MONEY})
        fx(ws, f"C{r}", f'=IF(A{r}="","",B{r})', MONEY)
        fx(ws, f"D{r}", f'=IF(A{r}="","",ROUND(B{r}*$B$5,2))', MONEY)
    for i, (item, cost) in enumerate([("Starter", 6), ("Main", 14), ("Dessert", 5)]):
        inp(ws, f"A{lo+i}", item); inp(ws, f"B{lo+i}", cost, MONEY, key=(i == 1))
    label(ws, f"A{tot}", "Food total", bold=True)
    for col in "CD": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)


def quote_rentals(ws):
    lo, hi = RENT_ROWS; tot = hi + 1
    head(ws, 23, ["Rentals and other costs", "Cost"])
    for r in range(lo, hi + 1): blank(ws, r, "AB", {"B": MONEY})
    inp(ws, f"A{lo}", "Tables and linen"); inp(ws, f"B{lo}", 450, MONEY)
    inp(ws, f"A{lo+1}", "Sound"); inp(ws, f"B{lo+1}", 300, MONEY)
    label(ws, f"A{tot}", "Rentals total", bold=True); fx(ws, f"B{tot}", f"=SUM(B{lo}:B{hi})", MONEY, bold=True)


def quote_staffing(ws):
    label(ws, "A36", "Staffing", bold=True).font = TITLE
    rows = [(37, "Hours per staff member", f"=MAX($B$6+{S}6,{S}7)", "0.0", "event hours + setup and teardown, at least the minimum shift"),
            (38, "Servers needed", f"=IFERROR(ROUNDUP($B$5/{S}4,0),0)", "0", "guests ÷ guests per server, rounded up"),
            (39, "Server cost", f"=ROUND(B38*B37*{S}2,2)", MONEY, "servers × hours × server rate"),
            (40, "Bartenders needed", f"=IFERROR(ROUNDUP($B$5/{S}5,0),0)", "0", "guests ÷ guests per bartender, rounded up"),
            (41, "Bartender cost", f"=ROUND(B40*B37*{S}3,2)", MONEY, "bartenders × hours × bartender rate"),
            (42, "Staffing total", "=B39+B41", MONEY, None)]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=(r == 42)); fx(ws, f"B{r}", f, fmt, bold=(r == 42))
        if note: label(ws, f"C{r}", note, muted=True)


def quote_summary(ws):
    m_tot = MENU_ROWS[1] + 1; r_tot = RENT_ROWS[1] + 1; b_tot = MIX_ROWS[1] + 1
    label(ws, "A44", "Quote summary", bold=True).font = TITLE
    rows = [(45, "Food total", f"=D{m_tot}", None), (46, "Staffing total", "=B42", None),
            (47, "Bar total", f"=Bar!E{b_tot}", "from the Bar sheet"), (48, "Rentals and other", f"=B{r_tot}", None),
            (49, "Total cost", "=B45+B46+B47+B48", None),
            (50, "Price at target margin", f"=IFERROR(ROUND(B49/(1-{S}10/100),2),0)", "total cost ÷ (1 − target margin %)"),
            (51, "Profit", "=B50-B49", None),
            (52, "Price per guest", "=IFERROR(ROUND(B50/$B$5,2),0)", None),
            (53, "Deposit", f"=ROUND(B50*{S}11/100,2)", "price × deposit %"),
            (54, "Balance due", "=B50-B53", None)]
    for r, name, f, note in rows:
        key = r in (49, 50, 52)
        label(ws, f"A{r}", name, bold=key); fx(ws, f"B{r}", f, MONEY, bold=key)
        if note: label(ws, f"C{r}", note, muted=True)


def build_quote(ws):
    quote_inputs(ws); quote_menu(ws); quote_rentals(ws); quote_staffing(ws); quote_summary(ws)
    label(ws, "A7", "Menu: cost per guest for each course or item. Rentals: fixed costs for this event.", muted=True)


def build_bar(ws):
    lo, hi = MIX_ROWS; tot = hi + 1
    label(ws, "A1", "Bar", bold=True).font = TITLE
    head(ws, 5, ["Drink type", "Share %", "Drinks", "Cost per serving", "Cost"], [30, 12, 12, 16, 14]); ws.column_dimensions["F"].width = 28
    for r, name, f, fmt in [(2, "Guests (from Quote)", f"={G}", "0"), (3, "Event hours (from Quote)", f"={H}", "0.0"),
                            (4, "Drinks total", f"=ROUND({G}*({S}8+{S}9*MAX({H}-1,0)),0)", "0")]:
        label(ws, f"A{r}", name, bold=(r == 4)); fx(ws, f"B{r}", f, fmt, bold=(r == 4))
    label(ws, "C4", "guests × (first-hour drinks + later-hour drinks × (hours − 1))", muted=True)
    for i, (kind, share, cost) in enumerate([("Beer", 40, 2.5), ("Wine", 40, 3), ("Spirits", 20, 4)]):
        r = lo + i
        inp(ws, f"A{r}", kind); inp(ws, f"B{r}", share, "0.0", key=True); inp(ws, f"D{r}", cost, MONEY, key=True)
        fx(ws, f"C{r}", f'=IF(A{r}="","",ROUND($B$4*B{r}/100,0))', "0")
        fx(ws, f"E{r}", f'=IF(A{r}="","",ROUND(C{r}*D{r},2))', MONEY)
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=SUM(B{lo}:B{hi})", "0.0", bold=True); fx(ws, f"C{tot}", f"=SUM(C{lo}:C{hi})", "0", bold=True)
    fx(ws, f"E{tot}", f"=SUM(E{lo}:E{hi})", MONEY, bold=True)
    fx(ws, f"F{tot}", f'=IF(ROUND(B{tot},2)=100,"","Shares must total 100")', bold=True).font = Font(name=FONT, bold=True, color="C00000")
    label(ws, f"A{tot+1}", "Bar cost per guest", bold=True); fx(ws, f"B{tot+1}", f"=IFERROR(ROUND(E{tot}/B2,2),0)", MONEY, bold=True)


def build_events(ws):
    lo, hi = EVENT_ROWS; tot = hi + 1
    label(ws, "A1", "Events", bold=True).font = TITLE
    label(ws, "A2", "One row per event. Enter the quoted price and the actual costs once known; profit, margin and outstanding are calculated.", muted=True)
    head(ws, 4, ["Event", "Date", "Client", "Guests", "Status", "Quoted price", "Actual food", "Actual staffing", "Actual bar", "Actual rentals",
                 "Total cost", "Profit", "Margin %", "Invoiced", "Paid", "Outstanding"],
         [22, 12, 18, 9, 11, 13, 12, 12, 12, 12, 13, 13, 10, 13, 13, 13])
    dv = DataValidation(type="list", formula1='"quoted,confirmed,done,lost"', allow_blank=True); ws.add_data_validation(dv)
    fm = {c: MONEY for c in "FGHIJNO"}; fm.update({"B": DATE, "D": "0"})
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFGHIJNO", fm); dv.add(f"E{r}")
        fx(ws, f"K{r}", f'=IF(A{r}="","",G{r}+H{r}+I{r}+J{r})', MONEY)
        fx(ws, f"L{r}", f'=IF(A{r}="","",F{r}-K{r})', MONEY)
        fx(ws, f"M{r}", f'=IF(A{r}="","",IFERROR(ROUND(L{r}/F{r}*100,2),""))', PCT)
        fx(ws, f"P{r}", f'=IF(A{r}="","",N{r}-O{r})', MONEY)
    y = dt.date.today().year
    sample = [("Harbour wedding", dt.date(y, 5, 16), "Smith family", 120, "done", 10600, 3600, 1450, 1300, 900, 10600, 10600),
              ("Summer staff party", dt.date(y, 7, 11), "Acme Ltd", 80, "done", 6300, 2100, 980, 820, 450, 6300, 3000),
              ("50th birthday", dt.date(y, 9, 26), "J. Lee", 45, "confirmed", 2900, 950, 600, 520, 300, 1160, 1160)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCDEFGHIJNO", row): inp(ws, f"{col}{r}", val, fm.get(col), key=(i == 0 and col in "AF"))
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "DFGHIJKLNOP": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", "0" if col == "D" else MONEY, bold=True)
    fx(ws, f"M{tot}", f"=IFERROR(ROUND(L{tot}/F{tot}*100,2),0)", PCT, bold=True)
    ws.freeze_panes = "B5"


def build_summary(ws):
    lo, hi = EVENT_ROWS; tot = hi + 1
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 50
    label(ws, "A1", "Summary", bold=True).font = TITLE
    rows = [(2, "Events tracked", f"=COUNTA(Events!A{lo}:A{hi})", "0", None),
            (3, "Quoted total", f"=Events!F{tot}", MONEY, None),
            (4, "Total cost", f"=Events!K{tot}", MONEY, None),
            (5, "Profit", f"=Events!L{tot}", MONEY, None),
            (6, "Margin %", f"=Events!M{tot}", PCT, "profit ÷ quoted total"),
            (7, "Total guests", f"=Events!D{tot}", "0", None),
            (8, "Average price per guest", f"=IFERROR(ROUND(B3/B7,2),0)", MONEY, "quoted total ÷ total guests"),
            (9, "Events below target margin", f'=COUNTIF(Events!M{lo}:M{hi},"<"&{S}10)', "0", "margin % under Settings target"),
            (10, "Outstanding", f"=Events!P{tot}", MONEY, "invoiced − paid")]
    for r, name, f, fmt, note in rows:
        label(ws, f"A{r}", name, bold=True); fx(ws, f"B{r}", f, fmt, bold=True)
        if note: label(ws, f"C{r}", note, muted=True)


README = """# Event & Catering Quote Workbook

Thank you for buying the workbook. Open `event-catering-quote-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** labour rates, guests per server and per bartender, setup and minimum hours, drinks-per-guest estimates, target margin and deposit.
2. **Quote:** guests, hours, menu cost per guest and rentals. Staffing, bar, total cost, price at your target margin, price per guest, deposit and balance follow.
3. **Bar:** the drinks estimate split by beer, wine and spirits with your cost per serving.
4. **Events:** quoted price against actual costs per event, with profit, margin and outstanding balance.
5. **Summary:** totals, average price per guest and how many events fell below your target margin.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; staffing ratios and drinks estimates are starting points, not rules. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_quote(wb.create_sheet("Quote")); build_bar(wb.create_sheet("Bar"))
    build_events(wb.create_sheet("Events")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx"); wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
