"""Build the Bookkeeping Practice Pricing Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "bookkeeping-pricing-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"
CLIENT_ROWS = (5, 34); QUOTE_ROWS = (5, 24); MONTH_ROWS = (5, 16)
S = "Settings!$B$"  # settings cell prefix
ENTITIES = '"sole trader,partnership,company,nonprofit,other"'
PB = "'Package Builder'!"; CQ = "'Catch-Up Quotes'!"; RL = "'Realization Log'!"
# Settings rows (column B): referenced by every other sheet.
R_OVERHEAD = 8; R_INCOME = 9; R_STAFF = 10; R_WEEKS = 11; R_HOURS = 12; R_BILLABLE = 13; R_STD_RATE = 14
R_BUFFER = 15; R_DISCOUNT = 16; R_TX = 17; R_ACCOUNT = 18; R_BILL_HRS = 19; R_COST_RATE = 20; R_RATE = 21; R_MINUTE = 22


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Bookkeeping Practice Pricing Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: your monthly overhead line by line, owner income goal, staff cost, working weeks, hours, billable share, standard hourly rate, scope buffer, commitment discount, per-transaction rate and per-account fee. Every other sheet reads them.",
        "2. Package Builder: one row per client with entity type, transactions, bank and card accounts, add-ons, hours, base fee and what you charge today. The suggested package fee, hourly-equivalent fee, recommended fee, the gap, effective hourly rate and annual fee are calculated.",
        "3. Catch-Up Quotes: one row per catch-up or clean-up job with months behind, hours per month, rate, setup fee and discount. Total hours, labour cost, discount, quote total and the per-month figure are calculated.",
        "4. Capacity: how many clients your billable hours can hold after admin time, hours used, spare hours, open client slots and revenue at capacity, read from the Package Builder unless you type your own figures.",
        "5. Realization Log: one row per month with hours worked, standard rate, amount billed and amount collected. Standard fees, billing, collection and overall realization, write-offs, uncollected and effective rate are calculated, with totals.",
        "6. Summary: cost-recovery rate against the rate you price with, clients priced below recommended, revenue now and at recommended fees, highest and lowest effective hourly rate clients, open slots and year-to-date realization.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Target hourly rate: the rate you price work at (rate used for pricing on Settings). Set it at or above the cost-recovery rate = (total monthly overhead + owner income goal + staff cost) ÷ billable hours per month. Billable hours per month = hours per week × billable % × working weeks per year ÷ 12.",
        "Package fee: a fixed monthly fee for ongoing bookkeeping = base fee + transactions × per-transaction rate + bank and card accounts × per-account fee + monthly add-ons (payroll runs, sales tax filings, reporting).",
        "Fixed fee buffer: the scope buffer % added on top of estimated hours × rate when pricing a fixed fee, so scope creep and messy months do not come out of your pocket. Hourly-equivalent fee = hours × rate × (1 + buffer %) × (1 − commitment discount %). Recommended fee = the higher of the package fee and the hourly-equivalent fee.",
        "Capacity: the most clients your usable hours can hold. Usable hours = billable hours per month × (1 − admin share %). Max clients = usable hours ÷ hours per client, rounded down.",
        "Billing realization = amount billed ÷ standard fees (hours worked × standard rate): how much of your time made it onto an invoice.",
        "Collection realization = amount collected ÷ amount billed: how much of what you invoiced was paid.",
        "Overall realization = amount collected ÷ standard fees: billing realization × collection realization.",
        "The targets in Settings are your own; this workbook does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 80
    label(ws, "A1", "Settings", bold=True).font = TITLE
    overhead = [("Software and subscriptions (monthly)", 250, "Example default 250 a month. Ledger software seats, receipt capture, practice management, e-signature."),
                ("Insurance (monthly)", 90, "Example default 90 a month. Professional indemnity and cyber cover."),
                ("Office or coworking (monthly)", 300, "Example default 300 a month. Rent, desk or a share of home office costs."),
                ("Marketing (monthly)", 100, "Example default 100 a month. Website, listings, directory fees, ads."),
                ("Professional fees and CPE (monthly)", 60, "Example default 60 a month. Memberships, licences, continuing education spread over 12 months."),
                ("Other overhead (monthly)", 100, "Example default 100 a month. Phone, internet, bank charges, anything else.")]
    for i, (name, val, note) in enumerate(overhead, start=2):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, MONEY, key=True); label(ws, f"C{i}", note, muted=True)
    label(ws, f"A{R_OVERHEAD}", "Total monthly overhead", bold=True)
    fx(ws, f"B{R_OVERHEAD}", "=SUM(B2:B7)", MONEY, bold=True)
    label(ws, f"C{R_OVERHEAD}", "formula: sum of the six overhead rows above", muted=True)
    rows = [(R_INCOME, "Monthly owner income goal", 4500, MONEY, True, "Example default 4,500 a month. What you want to pay yourself before personal tax."),
            (R_STAFF, "Staff cost per month", 0, MONEY, False, "Default 0. Wages, contractor and payroll costs for anyone else working on client files."),
            (R_WEEKS, "Working weeks per year", 48, INT0, True, "Default 48: 52 weeks less 4 weeks of holiday, sickness and training."),
            (R_HOURS, "Hours worked per week", 40, HRS, True, "Default 40. All hours at work, including admin, sales and learning."),
            (R_BILLABLE, "Share of hours that are billable %", 70, PCT, True, "Default 70. The share of your hours spent on client work; the rest is admin, sales and learning."),
            (R_STD_RATE, "Standard hourly rate", 90, MONEY, True, "Default 90. The hourly rate you price with; set it at or above the cost-recovery rate in row 20."),
            (R_BUFFER, "Scope buffer % for fixed fees", 15, PCT, True, "Default 15. Added to hours × rate when turning hours into a fixed fee, to cover scope creep and messy months."),
            (R_DISCOUNT, "Commitment discount %", 5, PCT, False, "Default 5. Taken off the hourly-equivalent fee for clients who commit to a monthly package."),
            (R_TX, "Per-transaction rate", 1.5, MONEY, True, "Default 1.50 per transaction a month. Used in the package fee on the Package Builder."),
            (R_ACCOUNT, "Per-bank-account fee", 25, MONEY, True, "Default 25 per bank or card account a month. Covers reconciling each account.")]
    for r, name, val, fmt, key, note in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", note, muted=True)
    ws[f"B{R_BILLABLE}"].comment = Comment("Assumption: nobody bills every hour at work. 70% leaves 30% for admin, sales, learning and follow-up.", "Knackdesk")
    ws[f"B{R_WEEKS}"].comment = Comment("Assumption: 4 weeks a year off for holiday, sickness and training.", "Knackdesk")
    ws[f"B{R_BUFFER}"].comment = Comment("Assumption: fixed-fee work runs over the estimate; 15% is a placeholder, set it from your own overruns.", "Knackdesk")
    derived = [(R_BILL_HRS, "Billable hours per month", f"=ROUND(B{R_HOURS}*B{R_BILLABLE}/100*B{R_WEEKS}/12,2)", "0.00",
                "formula: hours per week × billable % × working weeks ÷ 12"),
               (R_COST_RATE, "Cost-recovery hourly rate", f"=IFERROR(ROUND((B{R_OVERHEAD}+B{R_INCOME}+B{R_STAFF})/B{R_BILL_HRS},2),0)", MONEY,
                "formula: (total overhead + income goal + staff cost) ÷ billable hours per month; the least each billable hour must earn"),
               (R_RATE, "Rate used for pricing", f"=B{R_STD_RATE}", MONEY,
                "formula: the standard hourly rate in row 14; set that at or above the cost-recovery rate above"),
               (R_MINUTE, "Rate per minute", f"=ROUND(B{R_RATE}/60,4)", "0.0000", "formula: rate used for pricing ÷ 60")]
    rows_block(ws, derived, bold_rows=(R_BILL_HRS, R_COST_RATE, R_RATE))


def build_builder(ws):
    lo, hi = CLIENT_ROWS; tot = hi + 1
    label(ws, "A1", "Package Builder", bold=True).font = TITLE
    label(ws, "A2", "One row per client. Charges use the per-transaction rate, per-account fee, rate, scope buffer and commitment discount from Settings; the recommended fee is the higher of the package fee and the hourly-equivalent fee.", muted=True)
    head(ws, 4, ["Client", "Entity type", "Transactions per month", "Bank and card accounts", "Add-ons monthly",
                 "Estimated hours per month", "Base fee", "Current fee charged", "Transaction charge", "Account charge",
                 "Suggested monthly fee", "Hourly-equivalent fee", "Recommended fee", "Gap (current − recommended)",
                 "Effective hourly rate at current fee", "Annual at current fee"],
         [26, 13, 13, 12, 11, 12, 10, 12, 12, 11, 12, 13, 13, 14, 14, 13])
    dv = DataValidation(type="list", formula1=ENTITIES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": INT0, "D": INT0, "E": MONEY, "F": HRS, "G": MONEY, "H": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFGH", fm); dv.add(f"B{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"I{r}", g.format(f"ROUND(C{r}*{S}{R_TX},2)"), MONEY)
        fx(ws, f"J{r}", g.format(f"ROUND(D{r}*{S}{R_ACCOUNT},2)"), MONEY)
        fx(ws, f"K{r}", g.format(f"G{r}+I{r}+J{r}+E{r}"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(F{r}*{S}{R_RATE}*(1+{S}{R_BUFFER}/100)*(1-{S}{R_DISCOUNT}/100),2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"MAX(K{r},L{r})"), MONEY)
        fx(ws, f"N{r}", g.format(f"H{r}-M{r}"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(H{r}/F{r},2),0)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(H{r}*12,2)"), MONEY)
    sample = [("Riverside Cafe", "company", 180, 3, 50, 7, 250, 550),
              ("Okafor Consulting", "sole trader", 40, 1, 0, 2.5, 150, 220),
              ("Lindqvist Dental Partners", "partnership", 260, 4, 120, 9, 300, 950),
              ("Hope Street Food Bank", "nonprofit", 90, 2, 40, 4, 200, 380),
              ("Nguyen Builders Ltd", "company", 320, 5, 150, 12, 350, 600),
              ("Patel Design Studio", "sole trader", 60, 2, 0, 3, 150, 340),
              ("Green & Co Landscaping", "partnership", 120, 2, 30, 5, 250, 520),
              ("Mason Family Trust", "other", 25, 2, 0, 1.5, 150, 260)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCDEFGH", row): inp(ws, f"{col}{r}", val, fm.get(col), key=col in "CFH")
    ws[f"E{lo}"].comment = Comment("Add-ons are monthly extras priced separately: payroll runs, sales tax filings, management reports.", "Knackdesk")
    ws[f"G{lo}"].comment = Comment("Example base fee 250: the minimum monthly fee for this client before volume charges. Your own figure.", "Knackdesk")
    ws[f"F{lo}"].comment = Comment("Your estimate of monthly hours on this file; check it against your time records.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" clients"', bold=True)
    fx(ws, f"F{tot}", f"=SUM(F{lo}:F{hi})", HRS, bold=True)
    for col in "HKLMP": fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"N{tot}", f"=H{tot}-M{tot}", MONEY, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(H{tot}/F{tot},2),0)", MONEY, bold=True)
    label(ws, f"A{tot+1}", "Total row: monthly revenue at current fees (column H), annual (column P), total hours (column F) and client count (column B).", muted=True)
    ws.freeze_panes = "B5"


def build_quotes(ws):
    lo, hi = QUOTE_ROWS; tot = hi + 1
    label(ws, "A1", "Catch-Up Quotes", bold=True).font = TITLE
    label(ws, "A2", "One row per catch-up or clean-up job. The default rate comes from Settings; type your own rate in column E to override it for that quote.", muted=True)
    head(ws, 4, ["Client", "Months behind", "Hours per month", "Default rate (Settings)", "or type your own rate", "Rate used",
                 "Setup fee", "Discount %", "Total hours", "Labour cost", "Subtotal", "Discount", "Quote total", "Per month caught up"],
         [26, 10, 11, 12, 12, 11, 10, 10, 10, 12, 12, 11, 13, 13])
    fm = {"B": INT0, "C": HRS, "E": MONEY, "G": MONEY, "H": PCT}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEGH", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_RATE}"), MONEY)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), MONEY)
        fx(ws, f"I{r}", g.format(f"ROUND(B{r}*C{r},2)"), HRS)
        fx(ws, f"J{r}", g.format(f"ROUND(I{r}*F{r},2)"), MONEY)
        fx(ws, f"K{r}", g.format(f"J{r}+G{r}"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(K{r}*H{r}/100,2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"K{r}-L{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(M{r}/B{r},2),0)"), MONEY)
    sample = [("Harbour Bakery", 6, 5, None, 150, 10), ("Kowalski Plumbing", 12, 3.5, 80, 200, 0),
              ("Sunrise Yoga Studio", 3, 4, None, 100, 5), ("Ellis Property Group", 9, 8, 85, 250, 10)]
    for i, (name, months, hrs, rate, setup, disc) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", months, INT0, key=True); inp(ws, f"C{r}", hrs, HRS, key=True)
        inp(ws, f"E{r}", rate, MONEY); inp(ws, f"G{r}", setup, MONEY); inp(ws, f"H{r}", disc, PCT)
    ws[f"C{lo}"].comment = Comment("Catch-up months usually take longer than a normal month: missing statements, queries, fixing coding. Estimate generously.", "Knackdesk")
    ws[f"E{lo}"].comment = Comment("Leave empty to use the Settings rate; any number here takes precedence for this quote.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=SUM(B{lo}:B{hi})", INT0, bold=True)
    for col, fmt in (("G", MONEY), ("I", HRS), ("J", MONEY), ("K", MONEY), ("L", MONEY), ("M", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    ws.freeze_panes = "B5"


def override(ws, r, name, note):
    label(ws, f"A{r}", name)
    inp(ws, f"B{r}", None, ws[f"B{r-1}"].number_format, key=True, note=note)
    label(ws, f"C{r}", "optional override; leave empty to use the line above", muted=True)


def build_capacity(ws):
    tot = CLIENT_ROWS[1] + 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Capacity", bold=True).font = TITLE
    label(ws, "A2", "How many clients your billable hours can hold, and what they could earn.", muted=True)
    rows_block(ws, [(4, "Billable hours per month", f"={S}{R_BILL_HRS}", "0.00", "from Settings")])
    label(ws, "A5", "Admin share % of billable time")
    inp(ws, "B5", 10, PCT, key=True, note="Assumption: 10% of billable time goes on client admin that you cannot invoice (emails, chasing documents).")
    label(ws, "C5", "default 10: client emails, chasing documents, onboarding", muted=True)
    keep = "Leave empty to use the Package Builder figure above; any number here takes precedence."
    rows_block(ws, [(6, "Usable hours per month", "=ROUND(B4*(1-B5/100),2)", "0.00", "billable hours × (1 − admin share %)"),
                    (7, "Current clients from Package Builder", f"={PB}B{tot}", INT0, "clients listed on the Package Builder")])
    override(ws, 8, "or type your own number of clients", keep)
    rows_block(ws, [(9, "Clients used", '=IF(B8="",B7,B8)', INT0, "your override if typed, otherwise the Package Builder figure"),
                    (10, "Hours per client from Package Builder", f"=IFERROR(ROUND({PB}F{tot}/{PB}B{tot},2),0)", "0.00", "total hours ÷ clients")])
    override(ws, 11, "or type your own hours per client", keep)
    rows_block(ws, [(12, "Hours per client used", '=IF(B11="",B10,B11)', "0.00", "your override if typed, otherwise the Package Builder figure"),
                    (13, "Average fee from Package Builder", f"=IFERROR(ROUND({PB}H{tot}/{PB}B{tot},2),0)", MONEY, "monthly revenue at current fees ÷ clients")])
    override(ws, 14, "or type your own average monthly fee", keep)
    rows = [(15, "Average fee used", '=IF(B14="",B13,B14)', MONEY, "your override if typed, otherwise the Package Builder figure"),
            (16, "Max clients", "=IFERROR(INT(B6/B12),0)", INT0, "usable hours ÷ hours per client, rounded down"),
            (17, "Hours used", "=ROUND(B9*B12,2)", "0.00", "clients × hours per client"),
            (18, "Spare hours per month", "=ROUND(B6-B17,2)", '0.00;-0.00;0.00', "usable hours − hours used; negative means overbooked"),
            (19, "Open client slots", "=MAX(0,B16-B9)", INT0, "max clients − clients used, never below zero"),
            (20, "Monthly revenue at capacity", "=ROUND(B16*B15,2)", MONEY, "max clients × average fee"),
            (21, "Current monthly revenue", "=ROUND(B9*B15,2)", MONEY, "clients × average fee"),
            (22, "Status", '=IF(B9>B16,"Overbooked by "&(B9-B16)&" clients",IF(B19>0,B19&" open client slots","At capacity"))', None, None)]
    rows_block(ws, rows, bold_rows=(16, 19, 20, 22))


def build_log(ws):
    lo, hi = MONTH_ROWS; tot = hi + 1
    label(ws, "A1", "Realization Log", bold=True).font = TITLE
    label(ws, "A2", "One row per month, filled top down. The default rate comes from Settings; type your own standard rate in column D to override it for that month.", muted=True)
    head(ws, 4, ["Month", "Hours worked", "Default rate (Settings)", "or type your own rate", "Standard rate used", "Amount billed",
                 "Amount collected", "Standard fees", "Billing realization %", "Collection realization %", "Overall realization %",
                 "Write-offs", "Uncollected", "Effective rate"],
         [12, 10, 12, 12, 12, 12, 12, 12, 12, 12, 12, 11, 11, 11])
    fm = {"A": MONTH, "B": HRS, "D": MONEY, "F": MONEY, "G": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABDFG", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"C{r}", g.format(f"{S}{R_RATE}"), MONEY)
        fx(ws, f"E{r}", g.format(f'IF(D{r}<>"",D{r},C{r})'), MONEY)
        fx(ws, f"H{r}", g.format(f"ROUND(B{r}*E{r},2)"), MONEY)
        fx(ws, f"I{r}", g.format(f"IFERROR(ROUND(F{r}/H{r}*100,2),0)"), PCT)
        fx(ws, f"J{r}", g.format(f"IFERROR(ROUND(G{r}/F{r}*100,2),0)"), PCT)
        fx(ws, f"K{r}", g.format(f"IFERROR(ROUND(G{r}/H{r}*100,2),0)"), PCT)
        fx(ws, f"L{r}", g.format(f"H{r}-F{r}"), MONEY)
        fx(ws, f"M{r}", g.format(f"F{r}-G{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(G{r}/B{r},2),0)"), MONEY)
    year = dt.date.today().year
    sample = [(110, 9200, 8800), (105, 8900, 8700), (120, 10300, 9500), (100, 8600, 8400), (115, 9800, 9600), (108, 9300, 8900)]
    for i, (hrs, billed, collected) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        inp(ws, f"B{r}", hrs, HRS, key=(i == 0)); inp(ws, f"F{r}", billed, MONEY, key=(i == 0)); inp(ws, f"G{r}", collected, MONEY, key=(i == 0))
    ws[f"B{lo}"].comment = Comment("Hours worked on client files this month, billed or not, from your own time records.", "Knackdesk")
    ws[f"G{lo}"].comment = Comment("Cash received this month against this month's invoices; late payments can be added back when they arrive.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("B", HRS), ("F", MONEY), ("G", MONEY), ("H", MONEY), ("L", MONEY), ("M", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"I{tot}", f"=IFERROR(ROUND(F{tot}/H{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"J{tot}", f"=IFERROR(ROUND(G{tot}/F{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"K{tot}", f"=IFERROR(ROUND(G{tot}/H{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(G{tot}/B{tot},2),0)", MONEY, bold=True)
    ws[f"I{tot}"].comment = Comment("Totals recompute the percentages from the column totals, not an average of the monthly percentages.", "Knackdesk")
    ws.freeze_panes = "B5"


def build_summary(ws):
    lo, hi = CLIENT_ROWS; tot = hi + 1; l_tot = MONTH_ROWS[1] + 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 60
    label(ws, "A1", "Summary", bold=True).font = TITLE
    m = lambda col: f"{PB}{col}{lo}:{col}{hi}"
    rows = [(2, "Cost-recovery hourly rate", f"={S}{R_COST_RATE}", MONEY, "from Settings"),
            (3, "Rate used for pricing", f"={S}{R_RATE}", MONEY, "from Settings; keep it at or above the cost-recovery rate"),
            (4, "Clients on the Package Builder", f"={PB}B{tot}", INT0, None),
            (5, "Clients priced below recommended", f'=COUNTIF({m("N")},"<0")', INT0, "gap column below zero"),
            (6, "Monthly revenue at current fees", f"={PB}H{tot}", MONEY, None),
            (7, "Annual revenue at current fees", f"={PB}P{tot}", MONEY, None),
            (8, "Monthly revenue at recommended fees", f"={PB}M{tot}", MONEY, "if every client moved to the recommended fee"),
            (9, "Monthly uplift", "=B8-B6", MONEY, "recommended − current; negative means you charge above recommended overall"),
            (10, "Annual uplift", "=ROUND(B9*12,2)", MONEY, "monthly uplift × 12"),
            (11, "Highest effective hourly rate client", f'=IFERROR(INDEX({m("A")},MATCH(MAX({m("O")}),{m("O")},0)),"")', None, "current fee ÷ hours"),
            (12, "Lowest effective hourly rate client", f'=IFERROR(INDEX({m("A")},MATCH(MIN({m("O")}),{m("O")},0)),"")', None, "re-price first"),
            (13, "Open client slots", "=Capacity!B19", INT0, "from Capacity"),
            (14, "Spare hours per month", "=Capacity!B18", "0.00", "from Capacity"),
            (15, "Year-to-date billing realization", f"={RL}I{l_tot}", PCT, "Realization Log totals"),
            (16, "Year-to-date collection realization", f"={RL}J{l_tot}", PCT, "Realization Log totals"),
            (17, "Year-to-date overall realization", f"={RL}K{l_tot}", PCT, "Realization Log totals"),
            (18, "Year-to-date write-offs", f"={RL}L{l_tot}", MONEY, "standard fees not billed"),
            (19, "Year-to-date uncollected", f"={RL}M{l_tot}", MONEY, "billed but not collected")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))


README = """# Bookkeeping Practice Pricing Workbook

Thank you for buying the workbook. Open `bookkeeping-pricing-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** your monthly overhead line by line, owner income goal, staff cost, working weeks, hours, billable share, standard hourly rate, scope buffer, commitment discount, per-transaction rate and per-account fee; your cost-recovery rate follows.
2. **Package Builder:** clients with entity type, transactions, accounts, add-ons, hours, base fee and current fee; the suggested package fee, hourly-equivalent fee, recommended fee, the gap, effective hourly rate and annual fee follow.
3. **Catch-Up Quotes:** catch-up jobs with months behind, hours per month, rate (or your own), setup fee and discount; total hours, labour cost, quote total and per-month figure follow.
4. **Capacity:** usable hours after admin time, max clients, hours used, spare hours, open client slots and revenue at capacity.
5. **Realization Log:** hours, standard rate, billed and collected per month with billing, collection and overall realization, write-offs, uncollected, effective rate and totals.
6. **Summary:** cost-recovery rate vs the rate you use, clients below recommended, revenue now and at recommended fees, highest and lowest effective rate clients, open slots and year-to-date realization.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_builder(wb.create_sheet("Package Builder"))
    build_quotes(wb.create_sheet("Catch-Up Quotes")); build_capacity(wb.create_sheet("Capacity"))
    build_log(wb.create_sheet("Realization Log")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
