"""Build the Home Care Agency Pricing Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "home-care-pricing-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"
RATIO = "0.00"
RATE_ROWS = (5, 14); CG_ROWS = (5, 34); CL_ROWS = (5, 24); TO_ROWS = (16, 27)
S = "Settings!$B$"  # settings cell prefix
CG_STATUS = '"active,on leave,left"'
CL_STATUS = '"active,pending,ended"'
SR = "'Service Rates'!"; CG = "Caregivers!"; CQ = "'Clients & Quotes'!"; CT = "'Capacity & Turnover'!"
WEEKS_PER_MONTH = "52/12"
# Settings rows (column B): referenced by every other sheet.
R_PAY = 2; R_ONCOST = 3; R_NONBILL = 4; R_OH_FIRST = 5; R_OH_LAST = 10; R_OVERHEAD = 11; R_MARGIN = 12
R_WEEKEND = 13; R_OVERNIGHT = 14; R_MILEAGE = 15; R_HOURS_WK = 16; R_UTIL = 17; R_WEEKS = 18
R_LOADED = 19; R_LABOUR = 20; R_BILL_HRS = 21; R_OH_HOUR = 22; R_COST = 23; R_BILL = 24
# Capacity & Turnover top block rows (column B).
C_ACTIVE = 5; C_PAID = 6; C_BILLABLE = 7; C_COMMITTED = 8; C_SPARE = 9; C_STATUS = 10; C_REVENUE = 11
C_CEILING = 12; C_TRAINER = 13


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def note(ws, ref, text):
    ws[ref].comment = Comment(text, "Knackdesk")


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Home Care Agency Pricing Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: default caregiver pay rate, on-costs %, non-billable paid hours %, six monthly overhead lines, your target margin % on the bill rate, weekend uplift %, overnight rate, mileage rate, hours per caregiver per week, your target utilization % and weeks per quote period. Loaded pay rate, labour cost and overhead per billable hour, cost per billable hour and the default bill rate at your target margin follow and every other sheet reads them.",
        "2. Service Rates: one row per service level (companion, personal care, dementia care, respite, live-in and your own), with pay rate, on-costs % and margin % (defaults from Settings, type your own to override) and the bill rate you charge today. Cost per billable hour, the bill rate at your margin, margin per hour and the gap between what you charge and the computed rate are calculated.",
        "3. Caregivers: one row per caregiver with status (active, on leave, left), paid hours, billable hours, travel hours and training hours per week, and pay rate (default from Settings, type your own to override). Non-billable hours, utilization %, other non-billable hours, the weekly cost of non-billable time and the gap to your target utilization are calculated, with totals for active caregivers.",
        "4. Clients & Quotes: one row per client or quote with service level, hours per week, weekend hours, overnight hours, miles per week, weeks (default from Settings, type your own to override) and status (active, pending, ended). The bill rate is looked up from Service Rates; standard, weekend, overnight and mileage charges, the weekly and period total and the effective hourly rate are calculated, with totals for active clients.",
        "5. Capacity & Turnover: the top block compares the hours your active caregivers can bill at your target utilization with the hours committed to active clients, and shows spare hours, weekly revenue committed and the revenue ceiling. The lower block takes one row per month with caregivers at the start, leavers, hires, recruiting cost, onboarding hours, trainer cost per hour, unfilled hours and margin per billable hour (defaults you can override), and calculates turnover %, turnover cost, cost per leaver and caregivers at the end.",
        "6. Summary: cost per billable hour, default bill rate, service levels priced below the computed rate, active caregivers and utilization against your target, active clients, hours committed and spare hours, weekly and period revenue committed, year-to-date leavers, turnover % and turnover cost, and the client with the highest weekly total.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Pay rate vs bill rate: the pay rate is what you pay the caregiver per hour; the bill rate is what you charge the client per hour. The difference has to cover on-costs, non-billable time, overhead and margin.",
        "On-costs: what a caregiver costs you on top of their pay rate: employer payroll taxes, workers' compensation, holiday pay, pension and payroll fees. Loaded pay rate = pay rate × (1 + on-costs %).",
        "Non-billable paid hours: hours you pay for but cannot bill, such as travel between visits, training, supervision and waiting time, entered as a % of billable hours. Labour cost per billable hour = loaded pay rate × (1 + non-billable %).",
        "Cost per billable hour: labour cost per billable hour plus overhead per billable hour (monthly overhead ÷ planned billable hours per month).",
        "Margin on bill rate: profit as a share of the bill rate, not a markup on cost. Bill rate at margin = cost per billable hour ÷ (1 − margin %). A 25% margin on a 30.00 cost is a 40.00 bill rate, not 37.50.",
        "Utilization: billable hours ÷ paid hours, as a %. The total for the team is computed from the totals, not averaged across caregivers.",
        "Capacity: the hours your active caregivers can bill in a week at your target utilization = active caregivers × hours per caregiver per week × target utilization %.",
        "Care plan quote: the weekly and period price for one client's care plan: standard hours at the service level's bill rate, weekend hours with the uplift, overnight hours at the overnight rate, and mileage.",
        "Turnover cost: recruiting cost + onboarding hours × trainer cost per hour + unfilled hours × margin per billable hour lost. Turnover % for the year = leavers ÷ average headcount.",
        "The targets in this workbook are your own; it does not supply industry benchmarks. Set them from your own history and plan.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial or care advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [(R_PAY, "Default caregiver pay rate (per hour)", 16, MONEY, True, "Example default 16.00 an hour. What you pay a caregiver per hour before on-costs; each service level and caregiver can override it."),
            (R_ONCOST, "On-costs %", 22, PCT, True, "Default 22. Employer payroll taxes, workers' compensation, holiday pay and payroll fees on top of pay; each service level can override it."),
            (R_NONBILL, "Non-billable paid % of billable hours", 10, PCT, False, "Default 10. Paid travel, training, supervision and waiting time, as a % of billable hours."),
            (5, "Overhead: office (monthly)", 1200, MONEY, False, "Example default 1,200 a month. Rent, utilities, phones and office supplies."),
            (6, "Overhead: scheduling software (monthly)", 300, MONEY, False, "Example default 300 a month. Scheduling, visit verification and care management software."),
            (7, "Overhead: insurance (monthly)", 900, MONEY, False, "Example default 900 a month. Liability, bonding and other business insurance."),
            (8, "Overhead: recruiting and training (monthly)", 1000, MONEY, False, "Example default 1,000 a month. Job ads, background checks and ongoing training not costed on Capacity & Turnover."),
            (9, "Overhead: management salaries (monthly)", 3500, MONEY, True, "Example default 3,500 a month. Owner, care manager, scheduler and office staff pay, loaded."),
            (10, "Overhead: other (monthly)", 700, MONEY, False, "Example default 700 a month. Accounting, marketing, licences and everything else."),
            (R_MARGIN, "Target margin % on bill rate (your own)", 25, PCT, True, "Default 25. Your own placeholder, not a recommendation or an industry benchmark; each service level can override it."),
            (R_WEEKEND, "Weekend uplift %", 15, PCT, False, "Default 15. Added to the bill rate for weekend hours on Clients & Quotes."),
            (R_OVERNIGHT, "Overnight rate (per overnight hour)", 20, MONEY, False, "Example default 20.00 an hour. Flat rate billed per overnight hour on Clients & Quotes; type 0 if you do not offer overnight care."),
            (R_MILEAGE, "Mileage rate (per mile)", 0.5, MONEY, False, "Default 0.50 a mile. Charged to the client for miles driven on their errands or visits."),
            (R_HOURS_WK, "Hours per caregiver per week", 32, HRS, False, "Default 32. Planned paid hours per active caregiver, used for capacity on Capacity & Turnover."),
            (R_UTIL, "Target utilization % (your own)", 85, PCT, True, "Default 85. Your own placeholder, not a recommendation or an industry benchmark; the share of paid hours you plan to bill."),
            (R_WEEKS, "Weeks per period", 4, HRS, False, "Default 4. Weeks in a quote or billing period on Clients & Quotes; each client can override it.")]
    for r, name, val, fmt, key, txt in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", txt, muted=True)
    note(ws, f"B{R_PAY}", "Assumption: 16.00 is an example pay rate; use what you actually pay your caregivers.")
    note(ws, f"B{R_ONCOST}", "Assumption: 22% is a placeholder. On-costs depend on your country, state and the caregiver's contract; use your payroll provider's figures.")
    note(ws, f"B{R_NONBILL}", "Assumption: 10% is a placeholder. Compare paid hours with billed hours on your last few payrolls to set it.")
    note(ws, f"B{R_MARGIN}", "Your own placeholder, not a recommendation and not an industry benchmark. Margin is on the bill rate, not a markup on cost.")
    note(ws, f"B{R_OVERNIGHT}", "Assumption: overnight hours are billed at this flat hourly rate instead of the service level's bill rate. 0 means overnight care is not offered.")
    note(ws, f"B{R_UTIL}", "Your own placeholder, not a recommendation and not an industry benchmark. Look at Caregivers for the utilization you actually reach.")
    note(ws, f"B{R_WEEKS}", "Assumption: 4 weeks per period. Use 4.33 for a calendar month or 1 for a weekly quote.")
    derived = [(R_OVERHEAD, "Total monthly overhead", f"=SUM(B{R_OH_FIRST}:B{R_OH_LAST})", MONEY,
                "formula: sum of the six overhead lines above"),
               (R_LOADED, "Loaded pay rate", f"=ROUND(B{R_PAY}*(1+B{R_ONCOST}/100),2)", MONEY,
                "formula: pay rate × (1 + on-costs %)"),
               (R_LABOUR, "Labour cost per billable hour", f"=ROUND(B{R_LOADED}*(1+B{R_NONBILL}/100),2)", MONEY,
                "formula: loaded pay rate × (1 + non-billable paid %)"),
               (R_BILL_HRS, "Planned billable hours per month", f"=ROUND({CG}$D${CG_ROWS[1] + 1}*{WEEKS_PER_MONTH},2)", HRS,
                "formula: billable hours per week of active caregivers (Caregivers totals) × 52 ÷ 12"),
               (R_OH_HOUR, "Overhead per billable hour", f"=IFERROR(ROUND(B{R_OVERHEAD}/B{R_BILL_HRS},2),0)", MONEY,
                "formula: total monthly overhead ÷ planned billable hours per month"),
               (R_COST, "Cost per billable hour", f"=ROUND(B{R_LABOUR}+B{R_OH_HOUR},2)", MONEY,
                "formula: labour cost per billable hour + overhead per billable hour"),
               (R_BILL, "Default bill rate at target margin", f"=IFERROR(ROUND(B{R_COST}/(1-B{R_MARGIN}/100),2),0)", MONEY,
                "formula: cost per billable hour ÷ (1 − target margin %)")]
    rows_block(ws, derived, bold_rows=(R_OVERHEAD, R_COST, R_BILL))
    note(ws, f"B{R_BILL_HRS}", "Assumption: a month is 52 ÷ 12 = 4.33 weeks. Only caregivers with status active count.")


def build_rates(ws):
    lo, hi = RATE_ROWS
    label(ws, "A1", "Service Rates", bold=True).font = TITLE
    label(ws, "A2", "One row per service level. Pay rate, on-costs % and margin % default from Settings; type your own to override them. Type the bill rate you charge today to see the gap to the computed rate.", muted=True)
    head(ws, 4, ["Service level", "Default pay rate (Settings)", "or type your own pay rate", "Pay rate used",
                 "Default on-costs % (Settings)", "or type your own on-costs %", "On-costs % used",
                 "Default margin % (Settings)", "or type your own margin %", "Margin % used", "Loaded pay rate",
                 "Labour cost per billable hour", "Overhead per billable hour", "Cost per billable hour",
                 "Bill rate at margin", "Margin per hour", "Current bill rate charged", "Gap (current − computed)"],
         [20, 11, 11, 10, 11, 11, 10, 11, 11, 10, 11, 12, 12, 12, 11, 11, 12, 12])
    fm = {"C": MONEY, "F": PCT, "I": PCT, "Q": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ACFIQ", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"B{r}", g.format(f"{S}{R_PAY}"), MONEY)
        fx(ws, f"D{r}", g.format(f'IF(C{r}<>"",C{r},B{r})'), MONEY)
        fx(ws, f"E{r}", g.format(f"{S}{R_ONCOST}"), PCT)
        fx(ws, f"G{r}", g.format(f'IF(F{r}<>"",F{r},E{r})'), PCT)
        fx(ws, f"H{r}", g.format(f"{S}{R_MARGIN}"), PCT)
        fx(ws, f"J{r}", g.format(f'IF(I{r}<>"",I{r},H{r})'), PCT)
        fx(ws, f"K{r}", g.format(f"ROUND(D{r}*(1+G{r}/100),2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(K{r}*(1+{S}{R_NONBILL}/100),2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"{S}{R_OH_HOUR}"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(L{r}+M{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/(1-J{r}/100),2),0)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(O{r}-N{r},2)"), MONEY)
        fx(ws, f"R{r}", f'=IF(OR(A{r}="",Q{r}=""),"",ROUND(Q{r}-O{r},2))', MONEY)
    sample = [("Companion", 15, None, None, 36), ("Personal care", None, None, None, 38),
              ("Dementia care", 18, None, 28, 44), ("Respite", None, None, None, 34),
              ("Live-in", 14, None, None, 34)]
    for i, (name, pay, oncost, margin, current) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"C{r}", pay, MONEY); inp(ws, f"F{r}", oncost, PCT)
        inp(ws, f"I{r}", margin, PCT); inp(ws, f"Q{r}", current, MONEY, key=True)
    inp(ws, f"A{lo + 5}", "Other")
    note(ws, f"A{lo + 5}", "Add your own service levels in the empty rows; Clients & Quotes lists every name typed here.")
    note(ws, f"C{lo}", "Example override: companion care pays a little less than personal care.")
    note(ws, f"I{lo + 2}", "Example override: a higher margin on dementia care for the extra supervision and training it needs.")
    note(ws, f"Q{lo + 3}", "Example: respite is charged below the computed rate, so the gap is negative and the Summary counts it.")
    note(ws, f"M{lo}", "Assumption: overhead is spread evenly over every billable hour, whatever the service level.")
    note(ws, f"O{lo}", "Bill rate at margin = cost per billable hour ÷ (1 − margin %). Clients & Quotes prices each client at this rate.")
    note(ws, f"R{lo}", "Negative = you charge less than the rate that reaches your margin %. Blank when no current bill rate is typed.")
    ws.freeze_panes = "B5"


def build_caregivers(ws):
    lo, hi = CG_ROWS; tot = hi + 1
    label(ws, "A1", "Caregivers", bold=True).font = TITLE
    label(ws, "A2", "One row per caregiver, with typical weekly hours. Pay rate defaults from Settings; type your own to override it. Totals count active caregivers only.", muted=True)
    head(ws, 4, ["Caregiver", "Status", "Paid hours per week", "Billable hours per week", "Travel hours per week",
                 "Training hours per week", "Default pay rate (Settings)", "or type your own pay rate", "Pay rate used",
                 "Non-billable hours", "Utilization %", "Other non-billable hours", "Non-billable cost per week",
                 "vs target utilization (points)"],
         [20, 11, 10, 10, 10, 10, 11, 11, 10, 11, 11, 11, 12, 12])
    dv = DataValidation(type="list", formula1=CG_STATUS, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": HRS, "D": HRS, "E": HRS, "F": HRS, "H": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFH", fm); dv.add(f"B{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"{S}{R_PAY}"), MONEY)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), MONEY)
        fx(ws, f"J{r}", g.format(f"C{r}-D{r}"), HRS)
        fx(ws, f"K{r}", g.format(f"IFERROR(ROUND(D{r}/C{r}*100,2),0)"), PCT)
        fx(ws, f"L{r}", g.format(f"J{r}-E{r}-F{r}"), HRS)
        fx(ws, f"M{r}", g.format(f"ROUND(J{r}*I{r}*(1+{S}{R_ONCOST}/100),2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(K{r}-{S}{R_UTIL},2)"), PCT)
    sample = [("Ana R.", "active", 38, 32, 3, 1, None), ("Ben O.", "active", 36, 30, 4, 0, None),
              ("Carla M.", "active", 32, 28, 2, 1, None), ("Dev P.", "active", 40, 33, 5, 1, None),
              ("Erin K.", "active", 24, 20, 2, 0, None), ("Femi A.", "active", 35, 29, 3, 2, 18),
              ("Grace L.", "active", 30, 24, 4, 0, None), ("Hugo S.", "on leave", 0, 0, 0, 0, None),
              ("Ines T.", "active", 38, 31, 4, 1, None), ("Jamal W.", "left", 0, 0, 0, 0, None),
              ("Kim N.", "active", 28, 22, 3, 1, None), ("Luis G.", "active", 36, 27, 5, 2, None)]
    for i, (name, status, paid, billable, travel, training, pay) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", status); inp(ws, f"C{r}", paid, HRS, key=True)
        inp(ws, f"D{r}", billable, HRS, key=True); inp(ws, f"E{r}", travel, HRS); inp(ws, f"F{r}", training, HRS)
        inp(ws, f"H{r}", pay, MONEY)
    note(ws, f"C{lo}", "Hours you pay this caregiver in a typical week, including travel, training and waiting time.")
    note(ws, f"D{lo}", "Hours billed to clients in a typical week.")
    note(ws, f"H{lo + 5}", "Example override: a senior caregiver on a higher pay rate.")
    note(ws, f"K{lo}", "Utilization = billable hours ÷ paid hours. 0 when there are no paid hours.")
    note(ws, f"M{lo}", "Assumption: non-billable hours × pay rate × (1 + on-costs % from Settings).")
    label(ws, f"A{tot}", "Total (active)", bold=True)
    st = f"$B${lo}:$B${hi}"
    fx(ws, f"B{tot}", f'=COUNTIF({st},"active")', '0" active"', bold=True)
    for col, fmt in (("C", HRS), ("D", HRS), ("E", HRS), ("F", HRS), ("J", HRS), ("L", HRS), ("M", MONEY)):
        fx(ws, f"{col}{tot}", f'=SUMIF({st},"active",{col}{lo}:{col}{hi})', fmt, bold=True)
    fx(ws, f"K{tot}", f"=IFERROR(ROUND(D{tot}/C{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"N{tot}", f"=ROUND(K{tot}-{S}{R_UTIL},2)", PCT, bold=True)
    note(ws, f"K{tot}", "Team utilization from the totals of active caregivers, not an average of the rows.")
    ws.freeze_panes = "C5"


def build_clients(ws):
    lo, hi = CL_ROWS; tot = hi + 1; rlo, rhi = RATE_ROWS
    label(ws, "A1", "Clients & Quotes", bold=True).font = TITLE
    label(ws, "A2", "One row per client or care plan quote. Weekend and overnight hours are part of hours per week. The bill rate comes from Service Rates for the service level chosen. Weeks default from Settings; type your own to override. Totals count active clients only.", muted=True)
    head(ws, 4, ["Client", "Service level", "Hours per week", "Weekend hours", "Overnight hours", "Miles per week",
                 "Default weeks (Settings)", "or type your own weeks", "Weeks used", "Status", "Bill rate",
                 "Standard hours", "Standard hours charge", "Weekend charge", "Overnight charge", "Mileage charge",
                 "Weekly total", "Period total", "Effective hourly rate"],
         [20, 16, 10, 10, 10, 10, 10, 10, 9, 10, 10, 10, 12, 11, 11, 11, 12, 12, 11])
    dv_level = DataValidation(type="list", formula1=f"{SR}$A${rlo}:$A${rhi}", allow_blank=True)
    dv_status = DataValidation(type="list", formula1=CL_STATUS, allow_blank=True)
    ws.add_data_validation(dv_level); ws.add_data_validation(dv_status)
    fm = {"C": HRS, "D": HRS, "E": HRS, "F": HRS, "H": HRS}
    names = f"{SR}$A${rlo}:$A${rhi}"; rates = f"{SR}$O${rlo}:$O${rhi}"
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFHJ", fm); dv_level.add(f"B{r}"); dv_status.add(f"J{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"{S}{R_WEEKS}"), HRS)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), HRS)
        fx(ws, f"K{r}", g.format(f"IFERROR(INDEX({rates},MATCH(B{r},{names},0)),{S}{R_BILL})"), MONEY)
        fx(ws, f"L{r}", g.format(f"MAX(0,C{r}-D{r}-E{r})"), HRS)
        fx(ws, f"M{r}", g.format(f"ROUND(L{r}*K{r},2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(D{r}*K{r}*(1+{S}{R_WEEKEND}/100),2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"ROUND(E{r}*{S}{R_OVERNIGHT},2)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(F{r}*{S}{R_MILEAGE},2)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"ROUND(M{r}+N{r}+O{r}+P{r},2)"), MONEY)
        fx(ws, f"R{r}", g.format(f"ROUND(Q{r}*I{r},2)"), MONEY)
        fx(ws, f"S{r}", g.format(f"IFERROR(ROUND(Q{r}/C{r},2),0)"), MONEY)
    sample = [("Mrs. Alvarez", "Personal care", 40, 8, 0, 30, None, "active"),
              ("Mr. Brennan", "Dementia care", 56, 16, 0, 20, None, "active"),
              ("Mrs. Chen", "Companion", 20, 4, 0, 40, None, "active"),
              ("Mr. Dlamini", "Respite", 12, 6, 0, 10, 6, "active"),
              ("Mrs. Evans", "Live-in", 84, 24, 0, 0, None, "active"),
              ("Mr. Fischer", "Personal care", 30, 0, 8, 15, None, "active"),
              ("Mrs. Gupta", "Companion", 15, 0, 0, 20, None, "pending"),
              ("Mr. Harris", "Personal care", 25, 5, 0, 25, None, "ended")]
    for i, (name, level, hpw, wknd, night, miles, weeks, status) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", level); inp(ws, f"C{r}", hpw, HRS, key=True)
        inp(ws, f"D{r}", wknd, HRS); inp(ws, f"E{r}", night, HRS); inp(ws, f"F{r}", miles, HRS)
        inp(ws, f"H{r}", weeks, HRS); inp(ws, f"J{r}", status)
    note(ws, f"C{lo}", "All hours of care in a typical week, including the weekend and overnight hours in the next two columns.")
    note(ws, f"H{lo + 3}", "Example override: a six-week respite package.")
    note(ws, f"K{lo}", "Bill rate at margin for this service level on Service Rates. If the level is blank or not found, the default bill rate from Settings is used.")
    note(ws, f"L{lo}", "Hours per week − weekend hours − overnight hours, never below 0.")
    note(ws, f"N{lo}", "Assumption: weekend hours are billed at the bill rate × (1 + weekend uplift % from Settings).")
    note(ws, f"O{lo}", "Assumption: overnight hours are billed at the flat overnight rate from Settings, not the service level's bill rate.")
    note(ws, f"S{lo}", "Weekly total ÷ hours per week, including uplifts and mileage.")
    label(ws, f"A{tot}", "Total (active)", bold=True)
    st = f"$J${lo}:$J${hi}"
    fx(ws, f"B{tot}", f'=COUNTIF({st},"active")', '0" active"', bold=True)
    for col, fmt in (("C", HRS), ("D", HRS), ("E", HRS), ("F", HRS), ("L", HRS), ("M", MONEY), ("N", MONEY),
                     ("O", MONEY), ("P", MONEY), ("Q", MONEY), ("R", MONEY)):
        fx(ws, f"{col}{tot}", f'=SUMIF({st},"active",{col}{lo}:{col}{hi})', fmt, bold=True)
    fx(ws, f"S{tot}", f"=IFERROR(ROUND(Q{tot}/C{tot},2),0)", MONEY, bold=True)
    note(ws, f"S{tot}", "Weekly total ÷ hours per week for active clients, not an average of the rows.")
    ws.freeze_panes = "B5"


def build_capacity(ws):
    clo, chi = CG_ROWS; qlo, qhi = CL_ROWS; lo, hi = TO_ROWS; tot = hi + 1
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 14
    label(ws, "A1", "Capacity & Turnover", bold=True).font = TITLE
    label(ws, "A2", "Top: weekly capacity of active caregivers against hours committed to active clients. Below: one row per month of caregiver turnover, filled top down.", muted=True)
    label(ws, "A4", "CAPACITY (per week)", bold=True)
    cl_st = f"{CQ}$J${qlo}:$J${qhi}"
    rows = [(C_ACTIVE, "Active caregivers", f'=COUNTIF({CG}$B${clo}:$B${chi},"active")', INT0, None),
            (C_PAID, "Paid capacity per week (hours)", f"=ROUND(B{C_ACTIVE}*{S}{R_HOURS_WK},2)", HRS, None),
            (C_BILLABLE, "Billable capacity at target utilization (hours)", f"=ROUND(B{C_PAID}*{S}{R_UTIL}/100,2)", HRS, None),
            (C_COMMITTED, "Hours committed to active clients", f'=SUMIF({cl_st},"active",{CQ}$C${qlo}:$C${qhi})', HRS, None),
            (C_SPARE, "Spare hours", f"=ROUND(B{C_BILLABLE}-B{C_COMMITTED},2)", HRS, None),
            (C_STATUS, "Status", f'=IF(B{C_SPARE}<0,"overcommitted","spare")', None, None),
            (C_REVENUE, "Weekly revenue committed", f'=SUMIF({cl_st},"active",{CQ}$Q${qlo}:$Q${qhi})', MONEY, None),
            (C_CEILING, "Weekly revenue ceiling", f"=IFERROR(ROUND(B{C_REVENUE}/B{C_COMMITTED}*B{C_BILLABLE},2),ROUND(B{C_BILLABLE}*{S}{R_BILL},2))", MONEY, None)]
    rows_block(ws, rows, bold_rows=(C_SPARE, C_STATUS, C_REVENUE))
    label(ws, f"A{C_TRAINER}", "Default trainer cost per hour")
    inp(ws, f"B{C_TRAINER}", 30, MONEY, key=True)
    note(ws, f"B{C_PAID}", "Assumption: active caregivers × hours per caregiver per week (Settings). The Caregivers sheet shows the hours actually paid.")
    note(ws, f"B{C_SPARE}", "Billable capacity − hours committed. Negative means more hours are promised than your team can bill at your target utilization.")
    note(ws, f"B{C_CEILING}", "Billable capacity × the average hourly revenue of active clients (weekly revenue ÷ hours committed); the default bill rate from Settings if no client is active.")
    note(ws, f"B{C_TRAINER}", "Example default 30.00 an hour. Loaded cost of the person who trains and shadows new caregivers; each month can override it.")
    head(ws, lo - 1, ["Month", "Caregivers at start", "Leavers", "Hires", "Recruiting cost", "Onboarding hours",
                      "Default trainer cost per hour", "or type your own trainer cost", "Trainer cost per hour used",
                      "Unfilled hours", "Default margin per billable hour (Settings)", "or type your own margin per hour",
                      "Margin per billable hour used", "Turnover % for the month", "Cost per leaver",
                      "Turnover cost for the month", "Caregivers at end"])
    for col, w in zip("DEFGHIJKLMNOPQ", [9, 12, 11, 11, 11, 11, 10, 12, 12, 11, 11, 11, 13, 11]):
        ws.column_dimensions[col].width = w
    fm = {"A": MONTH, "B": INT0, "C": INT0, "D": INT0, "E": MONEY, "F": HRS, "H": MONEY, "J": HRS, "L": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFHJL", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"$B${C_TRAINER}"), MONEY)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND({S}{R_BILL}-{S}{R_COST},2)"), MONEY)
        fx(ws, f"M{r}", g.format(f'IF(L{r}<>"",L{r},K{r})'), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(C{r}/B{r}*100,2),0)"), PCT)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(P{r}/C{r},2),0)"), MONEY)
        fx(ws, f"P{r}", g.format(f"ROUND(E{r}+F{r}*I{r}+J{r}*M{r},2)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"B{r}-C{r}+D{r}"), INT0)
    year = dt.date.today().year
    sample = [(11, 1, 2, 600, 32, 20), (12, 2, 1, 900, 16, 40), (11, 1, 1, 450, 16, 12),
              (11, 0, 1, 500, 16, 0), (12, 2, 1, 800, 16, 30), (11, 1, 1, 400, 16, 10)]
    for i, (start, leavers, hires, recruit, onboard, unfilled) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH); inp(ws, f"B{r}", start, INT0, key=(i == 0))
        inp(ws, f"C{r}", leavers, INT0, key=(i == 0)); inp(ws, f"D{r}", hires, INT0)
        inp(ws, f"E{r}", recruit, MONEY); inp(ws, f"F{r}", onboard, HRS); inp(ws, f"J{r}", unfilled, HRS)
    note(ws, f"E{lo}", "Job ads, background checks, drug tests and agency fees spent on replacing caregivers this month.")
    note(ws, f"F{lo}", "Assumption: trainer hours spent onboarding and shadowing this month's new hires.")
    note(ws, f"J{lo}", "Client hours you could not staff this month because of a vacancy, so the margin on them was lost.")
    note(ws, f"K{lo}", "Default bill rate − cost per billable hour, both from Settings: the margin lost on each unfilled hour.")
    note(ws, f"P{lo}", "Assumption: turnover cost = recruiting cost + onboarding hours × trainer cost per hour + unfilled hours × margin per billable hour.")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=IFERROR(ROUND((AVERAGE(B{lo}:B{hi})+AVERAGE(Q{lo}:Q{hi}))/2,2),0)", RATIO, bold=True)
    for col, fmt in (("C", INT0), ("D", INT0), ("E", MONEY), ("F", HRS), ("J", HRS), ("P", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(C{tot}/B{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(P{tot}/C{tot},2),0)", MONEY, bold=True)
    note(ws, f"B{tot}", "Average headcount: the average of caregivers at the start and at the end across the months entered.")
    note(ws, f"N{tot}", "Year-to-date turnover % = leavers ÷ average headcount, not a sum of the monthly rates.")
    ws.freeze_panes = f"B{lo}"


def build_summary(ws):
    rlo, rhi = RATE_ROWS; qlo, qhi = CL_ROWS; cgtot = CG_ROWS[1] + 1; cltot = CL_ROWS[1] + 1; totr = TO_ROWS[1] + 1
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 80
    label(ws, "A1", "Summary", bold=True).font = TITLE
    weekly = f"{CQ}$Q${qlo}:$Q${qhi}"
    rows = [(2, "Cost per billable hour", f"={S}{R_COST}", MONEY, "from Settings: labour cost + overhead per billable hour"),
            (3, "Default bill rate at target margin", f"={S}{R_BILL}", MONEY, "from Settings: cost per billable hour ÷ (1 − target margin %)"),
            (4, "Service levels priced below the computed rate", f'=COUNTIF({SR}$R${rlo}:$R${rhi},"<0")', INT0, "Service Rates rows where the current bill rate is below the bill rate at margin"),
            (5, "Active caregivers", f"={CT}B{C_ACTIVE}", INT0, "Caregivers with status active"),
            (6, "Team utilization %", f"={CG}K{cgtot}", PCT, "billable ÷ paid hours of active caregivers, from the Caregivers totals"),
            (7, "Target utilization % (your own)", f"={S}{R_UTIL}", PCT, "from Settings; your own placeholder, not a benchmark"),
            (8, "Utilization vs target (points)", "=ROUND(B6-B7,2)", PCT, "negative = below your own target"),
            (9, "Active clients", f'=COUNTIF({CQ}$J${qlo}:$J${qhi},"active")', INT0, "Clients & Quotes with status active"),
            (10, "Weekly hours committed", f"={CT}B{C_COMMITTED}", HRS, "hours per week for active clients"),
            (11, "Spare hours per week", f"={CT}B{C_SPARE}", HRS, "billable capacity at target utilization − hours committed"),
            (12, "Weekly revenue committed", f"={CQ}Q{cltot}", MONEY, "weekly totals for active clients"),
            (13, "Period revenue committed", f"={CQ}R{cltot}", MONEY, "period totals for active clients"),
            (14, "Year-to-date leavers", f"={CT}C{totr}", INT0, "Capacity & Turnover totals"),
            (15, "Year-to-date turnover %", f"={CT}N{totr}", PCT, "leavers ÷ average headcount"),
            (16, "Year-to-date turnover cost", f"={CT}P{totr}", MONEY, "recruiting + onboarding + margin lost on unfilled hours"),
            (17, "Client with the highest weekly total", f'=IFERROR(INDEX({CQ}$A${qlo}:$A${qhi},MATCH(MAX({weekly}),{weekly},0)),"")', None, "largest weekly total on Clients & Quotes, any status"),
            (18, "Its weekly total", f"=IFERROR(MAX({weekly}),0)", MONEY, "highest weekly total on Clients & Quotes")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))
    note(ws, "B4", "Count of service levels with a negative gap on Service Rates. Raise those rates or accept a lower margin knowingly.")
    note(ws, "B15", "Turnover for the months entered so far, not annualised.")


README = """# Home Care Agency Pricing Workbook

Thank you for buying the workbook. Open `home-care-pricing-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** default caregiver pay rate, on-costs %, non-billable paid %, six monthly overhead lines, your target margin %, weekend uplift %, overnight rate, mileage rate, hours per caregiver, your target utilization % and weeks per period; loaded pay rate, cost per billable hour and the default bill rate at your margin follow.
2. **Service Rates:** each service level with pay rate, on-costs % and margin % (or your own) and the bill rate you charge today; cost per billable hour, bill rate at margin, margin per hour and the gap to what you charge follow.
3. **Caregivers:** each caregiver with status, paid, billable, travel and training hours per week and pay rate (or your own); non-billable hours, utilization %, non-billable cost per week and the gap to your target utilization follow.
4. **Clients & Quotes:** each client or care plan quote with service level, hours, weekend and overnight hours, miles, weeks (or your own) and status; bill rate, standard, weekend, overnight and mileage charges, weekly and period totals and the effective hourly rate follow.
5. **Capacity & Turnover:** billable capacity of active caregivers against hours committed to active clients, spare hours and revenue committed; plus each month's caregivers, leavers, hires, recruiting cost, onboarding and unfilled hours, with turnover %, turnover cost and cost per leaver.
6. **Summary:** cost per billable hour, default bill rate, service levels priced too low, utilization against your target, hours committed and spare, revenue committed, year-to-date turnover % and cost, and your largest client.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial or care advice; the targets are your own, not industry benchmarks. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_rates(wb.create_sheet("Service Rates"))
    build_caregivers(wb.create_sheet("Caregivers")); build_clients(wb.create_sheet("Clients & Quotes"))
    build_capacity(wb.create_sheet("Capacity & Turnover")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
