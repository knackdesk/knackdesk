"""Build the Law Firm Pricing Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "law-firm-pricing-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; DATE = "yyyy-mm-dd"; MONTH = "mmm yyyy"; INT0 = "0"; HRS = "0.0"
SIGNED = '0.00;-0.00;0.00'
TK_ROWS = (5, 24); FLAT_ROWS = (5, 24); MATTER_ROWS = (5, 34); LOCK_ROWS = (5, 16); CONT_ROWS = (22, 31)
S = "Settings!$B$"  # settings cell prefix
ROLES = '"partner,associate,paralegal,other"'
MATTER_TYPES = '"conveyancing,will and estate,incorporation,contract,immigration,family,litigation stage,other"'
FEE_TYPES = '"hourly,flat,contingency"'
YES_NO = '"yes,no"'
TK = "Timekeepers!"; FF = "'Flat Fee Pricer'!"; MT = "'Matter Tracker'!"; LC = "'Lock-Up & Contingency'!"
# Settings rows (column B): referenced by every other sheet.
R_OH_FIRST = 2; R_OH_LAST = 7; R_OVERHEAD = 8; R_METHOD = 9; R_BENEFITS = 10; R_HOURS = 11; R_TARGET = 12
R_CONTINGENCY = 13; R_DAYS = 14; R_REFERRAL = 15; R_WORKDAYS = 16; R_COUNT = 17; R_OH_EACH = 18; R_FIRM_HRS = 19; R_OH_HOUR = 20


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Law Firm Pricing Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: annual firm overhead line by line, benefits and payroll taxes %, default billable hours per timekeeper, your target margin on flat fees, scope contingency %, days in the year, default referral fee % and working days per month. Overhead per timekeeper follows and every other sheet reads it.",
        "2. Timekeepers: one row per lawyer or paralegal with role, annual salary or draw, benefits % and billable hours (defaults from Settings, type your own to override) and bill rate. Loaded salary, overhead share, total annual cost, cost rate, margin per hour, margin % and break-even hours are calculated.",
        "3. Flat Fee Pricer: one row per matter type or quote with partner, associate and paralegal hours and disbursements. Average cost rates per role come from the Timekeepers sheet; labour cost, cost with scope contingency, total cost, the flat fee at your target margin, effective hourly rate and the gap to the fee you charge today are calculated.",
        "4. Matter Tracker: one row per matter with fee type, budget hours and fees, actual hours and fees (time at standard rates), agreed fee, billed and collected. Hours and fees variance, realized rate, write-off, recovery % and an over / under budget status are calculated, with totals.",
        "5. Lock-Up & Contingency: one row per month with unbilled work in progress, accounts receivable and fees billed in the trailing 12 months; WIP days, debtor days, lock-up days and locked-up cash follow. Below, contingency matters: settlement, contingency %, case costs, whether costs come off before the fee and any referral fee; gross fee, net fee to the firm, net to the client and effective hourly rate follow.",
        "6. Summary: timekeepers, average cost and bill rate, timekeepers with negative margin, flat fees priced below the computed fee and the uplift, matters over budget, write-offs, recovery %, lock-up days and locked-up cash, and net contingency fees.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Timekeeper: anyone whose time is recorded and billed to matters (partner, associate, paralegal).",
        "Cost rate: what one billable hour of a timekeeper costs the firm = (salary or draw × (1 + benefits and payroll taxes %) + overhead per timekeeper) ÷ billable hours per year. Overhead is shared equally per timekeeper.",
        "Flat fee margin on price: margin % = profit ÷ fee, not markup on cost. Flat fee at target margin = total cost ÷ (1 − target margin %), never cost × (1 + target %). A 35% margin is a 53.85% markup.",
        "Scope contingency vs contingency fee: scope contingency is a cushion you add to the estimated labour cost of a flat fee for work that runs over. A contingency fee is a fee paid as a percentage of a settlement or award. The two are unrelated.",
        "Matter variance: actual minus budget, for hours and for fees (time at standard rates). Positive = over budget. Variance % = variance ÷ budget.",
        "Realization: how much of the value of time worked turns into fees. Write-off = time value above the agreed fee (flat or capped matters) or above the amount billed (hourly). Recovery % = collected ÷ actual fees. Realized rate = collected ÷ actual hours.",
        "Work in progress (WIP): time and disbursements recorded but not yet billed.",
        "Lock-up days: WIP days + debtor days = (unbilled WIP + accounts receivable) ÷ (fees billed in the trailing 12 months ÷ days in year). The number of days of fees tied up between doing the work and collecting the cash.",
        "The targets in Settings are your own; this workbook does not supply industry benchmarks. Set them from your own history and plan.",
        "The rules on fee arrangements, contingency fees and referral fees are a matter for your jurisdiction and professional body; check them before you use any figure here with a client.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not legal or financial advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    overhead = [("Rent and facilities (annual)", 60000, "Example default 60,000 a year. Rent, service charges, utilities, cleaning and repairs."),
                ("Support staff (annual)", 90000, "Example default 90,000 a year. Receptionist, secretaries, bookkeeper and practice manager, loaded with their payroll costs."),
                ("Software and research (annual)", 12000, "Example default 12,000 a year. Practice management, document management, legal research subscriptions and IT."),
                ("Insurance (annual)", 8000, "Example default 8,000 a year. Professional indemnity, office and cyber cover."),
                ("Marketing (annual)", 10000, "Example default 10,000 a year. Website, directories, advertising and business development."),
                ("Other overhead (annual)", 10000, "Example default 10,000 a year. Practising certificates, training, bank charges, accounting, anything else.")]
    for i, (name, val, note) in enumerate(overhead, start=R_OH_FIRST):
        label(ws, f"A{i}", name); inp(ws, f"B{i}", val, MONEY, key=True); label(ws, f"C{i}", note, muted=True)
    label(ws, f"A{R_OVERHEAD}", "Total annual firm overhead", bold=True)
    fx(ws, f"B{R_OVERHEAD}", f"=SUM(B{R_OH_FIRST}:B{R_OH_LAST})", MONEY, bold=True)
    label(ws, f"C{R_OVERHEAD}", "formula: sum of the six overhead rows above", muted=True)
    label(ws, f"A{R_METHOD}", "Overhead allocation method")
    label(ws, f"B{R_METHOD}", "per timekeeper")
    label(ws, f"C{R_METHOD}", "Documented, not an input: total overhead is shared equally across every timekeeper listed on the Timekeepers sheet.", muted=True)
    rows = [(R_BENEFITS, "Benefits and payroll taxes %", 20, PCT, True, "Default 20. Employer payroll taxes, pension, health cover and other benefits on top of salary or draw; each timekeeper can override it."),
            (R_HOURS, "Default billable hours per year per timekeeper", 1500, INT0, True, "Default 1,500. Hours recorded and billed to matters, not hours worked; each timekeeper can override it."),
            (R_TARGET, "Target margin % on flat fees", 35, PCT, True, "Default 35. Your own placeholder target margin on price (profit ÷ fee), not a recommendation or an industry benchmark; each matter can override it."),
            (R_CONTINGENCY, "Scope contingency %", 15, PCT, True, "Default 15. Cushion added to the estimated labour cost of a flat fee for work that runs over; each matter can override it."),
            (R_DAYS, "Days in year", 365, INT0, False, "Default 365. Used to turn trailing 12-month fees into fees per day for lock-up days."),
            (R_REFERRAL, "Default referral fee %", 0, PCT, False, "Default 0. Share of a contingency fee paid to a referring lawyer, where your rules allow it; each matter can override it."),
            (R_WORKDAYS, "Working days per month", 21, INT0, False, "Default 21. Used to show billable hours per working day on the Timekeepers sheet (12 months a year).")]
    for r, name, val, fmt, key, note in rows:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", note, muted=True)
    ws[f"B{R_BENEFITS}"].comment = Comment("Assumption: 20% is a placeholder for employer taxes and benefits; use your payroll figures. For equity partners, use the share of draw you treat as a cost.", "Knackdesk")
    ws[f"B{R_HOURS}"].comment = Comment("Assumption: 1,500 billable hours is a placeholder. Use your own recorded and billed hours from last year.", "Knackdesk")
    ws[f"B{R_TARGET}"].comment = Comment("Your own placeholder target, not a recommendation and not an industry benchmark. Margin is profit ÷ fee.", "Knackdesk")
    ws[f"B{R_CONTINGENCY}"].comment = Comment("Assumption: 15% is a placeholder for scope creep on fixed-fee work. Set it from how far your past flat-fee matters ran over.", "Knackdesk")
    ws[f"B{R_REFERRAL}"].comment = Comment("Whether referral fees are allowed, and how much, depends on your jurisdiction and professional body.", "Knackdesk")
    derived = [(R_COUNT, "Number of timekeepers", f"=COUNTA({TK}$A${TK_ROWS[0]}:$A${TK_ROWS[1]})", INT0,
                "formula: count of names on the Timekeepers sheet"),
               (R_OH_EACH, "Overhead per timekeeper (annual)", f"=IFERROR(ROUND(B{R_OVERHEAD}/B{R_COUNT},2),0)", MONEY,
                "formula: total annual overhead ÷ number of timekeepers"),
               (R_FIRM_HRS, "Total billable hours per year (all timekeepers)", f"={TK}$I${TK_ROWS[1] + 1}", INT0,
                "formula: billable hours used, summed on the Timekeepers sheet"),
               (R_OH_HOUR, "Overhead per billable hour (firm)", f"=IFERROR(ROUND(B{R_OVERHEAD}/B{R_FIRM_HRS},2),0)", MONEY,
                "formula: total annual overhead ÷ total billable hours; what overhead adds to each hour billed")]
    rows_block(ws, derived, bold_rows=(R_COUNT, R_OH_EACH, R_OH_HOUR))


def build_timekeepers(ws):
    lo, hi = TK_ROWS; tot = hi + 1
    label(ws, "A1", "Timekeepers", bold=True).font = TITLE
    label(ws, "A2", "One row per timekeeper. Benefits % and billable hours default from Settings; type your own in the 'or type your own' columns to override them.", muted=True)
    head(ws, 4, ["Name", "Role", "Annual salary or draw", "Default benefits % (Settings)", "or type your own benefits %",
                 "Benefits % used", "Default billable hours (Settings)", "or type your own hours", "Billable hours used",
                 "Bill rate", "Loaded salary", "Overhead allocation", "Total annual cost", "Cost rate", "Margin per hour",
                 "Margin %", "Break-even hours", "Billable hours per working day"],
         [22, 12, 14, 12, 12, 11, 12, 12, 11, 11, 13, 13, 13, 11, 11, 10, 11, 12])
    dv = DataValidation(type="list", formula1=ROLES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": MONEY, "E": PCT, "H": INT0, "J": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEHJ", fm); dv.add(f"B{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"D{r}", g.format(f"{S}{R_BENEFITS}"), PCT)
        fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), PCT)
        fx(ws, f"G{r}", g.format(f"{S}{R_HOURS}"), INT0)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), INT0)
        fx(ws, f"K{r}", g.format(f"ROUND(C{r}*(1+F{r}/100),2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"{S}{R_OH_EACH}"), MONEY)
        fx(ws, f"M{r}", g.format(f"K{r}+L{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(M{r}/I{r},2),0)"), MONEY)
        fx(ws, f"O{r}", g.format(f"ROUND(J{r}-N{r},2)"), MONEY)
        fx(ws, f"P{r}", g.format(f"IFERROR(ROUND(O{r}/J{r}*100,2),0)"), PCT)
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(M{r}/J{r},2),0)"), "0.00")
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(I{r}/({S}{R_WORKDAYS}*12),2),0)"), "0.00")
    sample = [("Partner A", "partner", 180000, None, None, 450), ("Partner B", "partner", 160000, None, 1300, 400),
              ("Associate A", "associate", 95000, None, None, 275), ("Associate B", "associate", 85000, 25, None, 250),
              ("Associate C", "associate", 70000, None, None, 225), ("Paralegal A", "paralegal", 50000, None, 1200, 150)]
    for i, (name, role, pay, ben, hrs, rate) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", role); inp(ws, f"C{r}", pay, MONEY, key=True)
        inp(ws, f"E{r}", ben, PCT); inp(ws, f"H{r}", hrs, INT0); inp(ws, f"J{r}", rate, MONEY, key=True)
    ws[f"C{lo}"].comment = Comment("For an equity partner, use the draw you would have to pay someone to do this work, not the full profit share.", "Knackdesk")
    ws[f"H{lo + 1}"].comment = Comment("Example override: a partner who spends more time on management and business development bills fewer hours.", "Knackdesk")
    ws[f"L{lo}"].comment = Comment("Assumption: overhead is shared equally per timekeeper (Settings). A paralegal carries the same overhead as a partner.", "Knackdesk")
    ws[f"Q{lo}"].comment = Comment("Billable hours a year at this bill rate needed to cover this timekeeper's total annual cost.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" timekeepers"', bold=True)
    for col, fmt in (("C", MONEY), ("I", INT0), ("K", MONEY), ("L", MONEY), ("M", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(M{tot}/I{tot},2),0)", MONEY, bold=True)
    ws[f"N{tot}"].comment = Comment("Blended firm cost rate: total annual cost ÷ total billable hours, not an average of the rows.", "Knackdesk")
    ws.freeze_panes = "B5"


def role_rate(role):
    lo, hi = TK_ROWS
    return f'=IFERROR(ROUND(AVERAGEIF({TK}$B${lo}:$B${hi},"{role}",{TK}$N${lo}:$N${hi}),2),0)'


def build_flat(ws):
    lo, hi = FLAT_ROWS; tot = hi + 1
    label(ws, "A1", "Flat Fee Pricer", bold=True).font = TITLE
    for col, name, role in (("A", "Partner cost rate", "partner"), ("C", "Associate cost rate", "associate"), ("E", "Paralegal cost rate", "paralegal")):
        nxt = chr(ord(col) + 1)
        label(ws, f"{col}2", name, bold=True); fx(ws, f"{nxt}2", role_rate(role), MONEY, bold=True)
    ws["B2"].comment = Comment("Average cost rate of every timekeeper with this role on the Timekeepers sheet (0 if there is none).", "Knackdesk")
    label(ws, "A3", "One row per matter type or quote. Contingency % and margin % default from Settings; type your own to override them. Current fee is what you charge today.", muted=True)
    head(ws, 4, ["Matter", "Matter type", "Partner hours", "Associate hours", "Paralegal hours", "Disbursements",
                 "Default contingency % (Settings)", "or type your own contingency %", "Contingency % used",
                 "Default margin % (Settings)", "or type your own margin %", "Margin % used", "Labour cost",
                 "Cost with contingency", "Total cost", "Flat fee", "Total hours", "Effective hourly rate",
                 "Current fee charged", "Gap: current − flat fee"],
         [28, 15, 10, 10, 10, 12, 12, 12, 11, 12, 12, 11, 12, 12, 12, 12, 10, 12, 12, 13])
    dv = DataValidation(type="list", formula1=MATTER_TYPES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"C": HRS, "D": HRS, "E": HRS, "F": MONEY, "H": PCT, "K": PCT, "S": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFHKS", fm); dv.add(f"B{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"G{r}", g.format(f"{S}{R_CONTINGENCY}"), PCT)
        fx(ws, f"I{r}", g.format(f'IF(H{r}<>"",H{r},G{r})'), PCT)
        fx(ws, f"J{r}", g.format(f"{S}{R_TARGET}"), PCT)
        fx(ws, f"L{r}", g.format(f'IF(K{r}<>"",K{r},J{r})'), PCT)
        fx(ws, f"M{r}", g.format(f"ROUND(C{r}*$B$2+D{r}*$D$2+E{r}*$F$2,2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(M{r}*(1+I{r}/100),2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"N{r}+F{r}"), MONEY)
        fx(ws, f"P{r}", g.format(f"IFERROR(ROUND(O{r}/(1-L{r}/100),2),0)"), MONEY)
        fx(ws, f"Q{r}", g.format(f"C{r}+D{r}+E{r}"), HRS)
        fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(P{r}/Q{r},2),0)"), MONEY)
        fx(ws, f"T{r}", f'=IF(OR(A{r}="",S{r}=""),"",ROUND(S{r}-P{r},2))', MONEY)
    sample = [("Residential purchase", "conveyancing", 2, 6, 4, 150, None, None, 1800),
              ("Simple will", "will and estate", 1, 2, 3, 0, None, None, 1500),
              ("Company incorporation", "incorporation", 1, 4, 2, 300, None, None, 2200),
              ("Commercial contract review", "contract", 3, 5, 0, 0, None, None, 2500),
              ("Work visa application", "immigration", 2, 8, 6, 500, None, None, 4200),
              ("Litigation: pleadings stage", "litigation stage", 10, 25, 10, 1200, 25, 30, 12000)]
    for i, (name, kind, ph, ah, lh, disb, cont, margin, fee) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", kind)
        for col, val in zip("CDE", (ph, ah, lh)): inp(ws, f"{col}{r}", val, HRS, key=True)
        inp(ws, f"F{r}", disb, MONEY); inp(ws, f"H{r}", cont, PCT); inp(ws, f"K{r}", margin, PCT); inp(ws, f"S{r}", fee, MONEY)
    ws[f"C{lo}"].comment = Comment("Estimated hours per role for a typical matter of this type. Use the recorded time on your last few similar matters.", "Knackdesk")
    ws[f"F{lo}"].comment = Comment("Assumption: disbursements are included in the flat fee and carry the margin. Leave them out if you pass them through at cost.", "Knackdesk")
    ws[f"H{lo + 5}"].comment = Comment("Example override: litigation stages run over more often, so this matter carries a larger scope contingency.", "Knackdesk")
    ws[f"T{lo}"].comment = Comment("Negative = you charge less today than the flat fee at your target margin.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("C", HRS), ("D", HRS), ("E", HRS), ("F", MONEY), ("M", MONEY), ("N", MONEY), ("O", MONEY),
                     ("P", MONEY), ("Q", HRS), ("S", MONEY), ("T", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    ws.freeze_panes = "B5"


def build_matters(ws):
    lo, hi = MATTER_ROWS; tot = hi + 1
    label(ws, "A1", "Matter Tracker", bold=True).font = TITLE
    label(ws, "A2", "One row per matter. Actual fees = recorded time at standard rates. Agreed fee is the flat or capped fee; leave it empty for hourly matters.", muted=True)
    head(ws, 4, ["Matter", "Client", "Fee type", "Budget hours", "Budget fees", "Actual hours", "Actual fees",
                 "Agreed fee", "Billed", "Collected", "Hours variance", "Hours variance %", "Fees variance",
                 "Fees variance %", "Realized rate", "Write-off", "Recovery %", "Status"],
         [28, 18, 12, 10, 12, 10, 12, 12, 12, 12, 10, 10, 12, 10, 11, 12, 10, 13])
    dv = DataValidation(type="list", formula1=FEE_TYPES, allow_blank=True); ws.add_data_validation(dv)
    fm = {"D": HRS, "E": MONEY, "F": HRS, "G": MONEY, "H": MONEY, "I": MONEY, "J": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEFGHIJ", fm); dv.add(f"C{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"K{r}", g.format(f"ROUND(F{r}-D{r},2)"), SIGNED)
        fx(ws, f"L{r}", g.format(f"IFERROR(ROUND(K{r}/D{r}*100,2),0)"), PCT)
        fx(ws, f"M{r}", g.format(f"ROUND(G{r}-E{r},2)"), MONEY)
        fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(M{r}/E{r}*100,2),0)"), PCT)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(J{r}/F{r},2),0)"), MONEY)
        fx(ws, f"P{r}", g.format(f"IF(N(H{r})>0,MAX(0,ROUND(G{r}-H{r},2)),MAX(0,ROUND(G{r}-I{r},2)))"), MONEY)
        fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(J{r}/G{r}*100,2),0)"), PCT)
        fx(ws, f"R{r}", g.format(f'IF(G{r}>E{r},"Over budget",IF(G{r}<E{r},"Under budget","On budget"))'))
    sample = [("Smith residential purchase", "J. Smith", "flat", 12, 1800, 15, 2600, 1800, 1800, 1800),
              ("Shareholder agreement", "Acme Ltd", "hourly", 20, 5000, 18, 4600, None, 4600, 4200),
              ("Patel v Northwind", "R. Patel", "hourly", 60, 15000, 72, 18500, None, 17000, 15000),
              ("Lee estate plan", "M. Lee", "flat", 6, 1500, 5, 1250, 1500, 1500, 1500),
              ("Garcia injury claim", "L. Garcia", "contingency", 80, 20000, 95, 24000, 33600, 33600, 33600),
              ("Lease review", "Brightline Cafe", "hourly", 10, 2500, 10, 2500, None, 2500, 2500),
              ("Okafor work visa", "C. Okafor", "flat", 14, 3500, 16, 3900, 3500, 3500, 2000),
              ("Harbor employment dispute", "Harbor Freight Co", "hourly", 40, 10000, 35, 8800, None, 8000, 6000)]
    for i, row in enumerate(sample):
        r = lo + i
        for col, val in zip("ABCDEFGHIJ", row): inp(ws, f"{col}{r}", val, fm.get(col), key=col in "EG")
    ws[f"G{lo}"].comment = Comment("Recorded time on the matter at each timekeeper's standard bill rate, whatever the fee arrangement.", "Knackdesk")
    ws[f"H{lo}"].comment = Comment("The flat, capped or contingency fee agreed with the client. Leave empty for hourly matters: the write-off then compares actual fees with the amount billed.", "Knackdesk")
    ws[f"P{lo}"].comment = Comment("Assumption: write-off = actual fees above the agreed fee (if any), otherwise actual fees above the amount billed. Never negative.", "Knackdesk")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" matters"', bold=True)
    for col, fmt in (("D", HRS), ("E", MONEY), ("F", HRS), ("G", MONEY), ("H", MONEY), ("I", MONEY), ("J", MONEY),
                     ("K", SIGNED), ("M", MONEY), ("P", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"L{tot}", f"=IFERROR(ROUND(K{tot}/D{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(M{tot}/E{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(J{tot}/F{tot},2),0)", MONEY, bold=True)
    fx(ws, f"Q{tot}", f"=IFERROR(ROUND(J{tot}/G{tot}*100,2),0)", PCT, bold=True)
    ws[f"Q{tot}"].comment = Comment("Totals recompute from the column totals, not an average of the matter rows.", "Knackdesk")
    ws.freeze_panes = "B5"


def build_lockup(ws):
    lo, hi = LOCK_ROWS
    label(ws, "A1", "Lock-Up & Contingency", bold=True).font = TITLE
    label(ws, "A2", "Lock-up: one row per month, filled top down, from your month-end WIP and debtors reports.", muted=True)
    head(ws, 4, ["Month", "Unbilled WIP", "Accounts receivable", "Fees billed, trailing 12 months", "Fees per day",
                 "WIP days", "Debtor days", "Lock-up days", "Locked-up cash", "Cash per day of lock-up",
                 "Change in lock-up days vs previous month"],
         [26, 13, 13, 14, 12, 10, 10, 10, 13, 13, 14])
    fm = {"A": MONTH, "B": MONEY, "C": MONEY, "D": MONEY}
    days = f"{S}{R_DAYS}"
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCD", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"E{r}", g.format(f"IFERROR(ROUND(D{r}/{days},2),0)"), MONEY)
        fx(ws, f"F{r}", g.format(f"IFERROR(ROUND(B{r}/E{r},2),0)"), "0.00")
        fx(ws, f"G{r}", g.format(f"IFERROR(ROUND(C{r}/E{r},2),0)"), "0.00")
        fx(ws, f"H{r}", g.format(f"F{r}+G{r}"), "0.00")
        fx(ws, f"I{r}", g.format(f"B{r}+C{r}"), MONEY)
        fx(ws, f"J{r}", g.format(f"IFERROR(ROUND(I{r}/H{r},2),0)"), MONEY)
        if r == lo:
            fx(ws, f"K{r}", '=""')
        else:
            fx(ws, f"K{r}", f'=IF(OR(A{r}="",A{r - 1}=""),"",ROUND(H{r}-H{r - 1},2))', SIGNED)
    year = dt.date.today().year
    sample = [(180000, 220000, 1800000), (175000, 230000, 1820000), (170000, 210000, 1830000),
              (160000, 205000, 1850000), (165000, 195000, 1860000), (150000, 190000, 1880000)]
    for i, (wip, ar, fees) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", dt.date(year, i + 1, 1), MONTH)
        inp(ws, f"B{r}", wip, MONEY, key=(i == 0)); inp(ws, f"C{r}", ar, MONEY, key=(i == 0)); inp(ws, f"D{r}", fees, MONEY, key=(i == 0))
    ws[f"B{lo}"].comment = Comment("Recorded time and disbursements not yet billed at month end, at the value you expect to bill.", "Knackdesk")
    ws[f"D{lo}"].comment = Comment("Fees billed in the 12 months to this month end. Fees per day = this ÷ days in year (Settings).", "Knackdesk")
    ws[f"J{lo}"].comment = Comment("Locked-up cash ÷ lock-up days: roughly the cash released for each day you take off lock-up.", "Knackdesk")

    clo, chi = CONT_ROWS; ctot = chi + 1
    label(ws, f"A{clo - 3}", "Contingency matters", bold=True)
    label(ws, f"A{clo - 2}", "Referral fee % defaults from Settings; type your own in column G to override it. Whether costs come off before the fee, and any referral fee, depend on your agreement and your rules.", muted=True)
    head(ws, clo - 1, ["Matter", "Settlement amount", "Contingency %", "Case costs", "Costs deducted before fee",
                       "Default referral fee % (Settings)", "or type your own referral %", "Referral fee % used",
                       "Hours worked", "Fee base", "Gross fee", "Referral fee", "Net fee to firm", "Net to client",
                       "Effective hourly rate"])
    dv = DataValidation(type="list", formula1=YES_NO, allow_blank=True); ws.add_data_validation(dv)
    fm = {"B": MONEY, "C": PCT, "D": MONEY, "G": PCT, "I": HRS}
    for r in range(clo, chi + 1):
        blank(ws, r, "ABCDEGI", fm); dv.add(f"E{r}")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"F{r}", g.format(f"{S}{R_REFERRAL}"), PCT)
        fx(ws, f"H{r}", g.format(f'IF(G{r}<>"",G{r},F{r})'), PCT)
        fx(ws, f"J{r}", g.format(f'IF(E{r}="yes",B{r}-D{r},B{r})'), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND(J{r}*C{r}/100,2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(K{r}*H{r}/100,2)"), MONEY)
        fx(ws, f"M{r}", g.format(f"K{r}-L{r}"), MONEY)
        fx(ws, f"N{r}", g.format(f"ROUND(B{r}-K{r}-D{r},2)"), MONEY)
        fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(M{r}/I{r},2),0)"), MONEY)
    matters = [("Garcia injury claim", 120000, 30, 8000, "yes", None, 95),
               ("Nguyen slip and fall", 45000, 33.33, 3000, "no", 25, 60),
               ("Delta product claim", 250000, 25, 22000, "yes", None, 310)]
    for i, (name, settle, pct, costs, before, ref, hrs) in enumerate(matters):
        r = clo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", settle, MONEY, key=True); inp(ws, f"C{r}", pct, PCT, key=True)
        inp(ws, f"D{r}", costs, MONEY); inp(ws, f"E{r}", before); inp(ws, f"G{r}", ref, PCT); inp(ws, f"I{r}", hrs, HRS)
    ws[f"E{clo}"].comment = Comment("yes = the fee is a percentage of the settlement after case costs are repaid; no = of the gross settlement.", "Knackdesk")
    ws[f"G{clo + 1}"].comment = Comment("Example override: a quarter of the fee goes to the referring lawyer. Check your rules first.", "Knackdesk")
    ws[f"N{clo}"].comment = Comment("Assumption: case costs are repaid out of the settlement either way; net to client = settlement − gross fee − case costs.", "Knackdesk")
    label(ws, f"A{ctot}", "Total", bold=True)
    for col, fmt in (("B", MONEY), ("D", MONEY), ("I", HRS), ("K", MONEY), ("L", MONEY), ("M", MONEY), ("N", MONEY)):
        fx(ws, f"{col}{ctot}", f"=SUM({col}{clo}:{col}{chi})", fmt, bold=True)
    fx(ws, f"O{ctot}", f"=IFERROR(ROUND(M{ctot}/I{ctot},2),0)", MONEY, bold=True)
    ws[f"O{ctot}"].comment = Comment("Total net fee ÷ total hours, not an average of the matter rates.", "Knackdesk")
    ws.freeze_panes = "B5"


def build_summary(ws):
    tlo, thi = TK_ROWS; flo, fhi = FLAT_ROWS; mlo, mhi = MATTER_ROWS; llo, lhi = LOCK_ROWS; clo, chi = CONT_ROWS
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 20; ws.column_dimensions["C"].width = 70
    label(ws, "A1", "Summary", bold=True).font = TITLE
    months = f"{LC}$A${llo}:$A${lhi}"
    latest = lambda col: f"INDEX({LC}${col}${llo}:${col}${lhi},MATCH(MAX({months}),{months},0))"
    rows = [(2, "Timekeepers", f"={S}{R_COUNT}", INT0, "from Settings"),
            (3, "Average cost rate", f"=IFERROR(ROUND(AVERAGE({TK}$N${tlo}:$N${thi}),2),0)", MONEY, "simple average of the timekeeper cost rates"),
            (4, "Average bill rate", f"=IFERROR(ROUND(AVERAGE({TK}$J${tlo}:$J${thi}),2),0)", MONEY, "simple average of the timekeeper bill rates"),
            (5, "Timekeepers with negative margin", f'=COUNTIF({TK}$O${tlo}:$O${thi},"<0")', INT0, "bill rate below cost rate"),
            (6, "Total annual cost of timekeepers", f"={TK}M{thi + 1}", MONEY, "loaded salaries + overhead"),
            (7, "Flat fee matters priced below the computed fee", f'=COUNTIF({FF}$T${flo}:$T${fhi},"<0")', INT0, "current fee below the flat fee at your target margin"),
            (8, "Monthly uplift at the computed fee", f'=-SUMIF({FF}$T${flo}:$T${fhi},"<0",{FF}$T${flo}:$T${fhi})', MONEY, "assumes each under-priced matter is sold once a month"),
            (9, "Matters over budget", f'=COUNTIF({MT}$R${mlo}:$R${mhi},"Over budget")', INT0, "actual fees above budget fees"),
            (10, "Total write-offs", f"={MT}P{mhi + 1}", MONEY, "Matter Tracker totals"),
            (11, "Overall recovery %", f"={MT}Q{mhi + 1}", PCT, "collected ÷ actual fees, from the totals"),
            (12, "Latest lock-up days", f"=IFERROR({latest('H')},0)", "0.00", "the most recent month on Lock-Up & Contingency"),
            (13, "Latest locked-up cash", f"=IFERROR({latest('I')},0)", MONEY, "unbilled WIP + accounts receivable"),
            (14, "Change in lock-up days vs first month", f"=IFERROR(ROUND(B12-{LC}H{llo},2),0)", SIGNED, "negative = cash is coming in faster"),
            (15, "Total net contingency fees to firm", f"={LC}M{chi + 1}", MONEY, "after referral fees"),
            (16, "Average effective hourly rate on contingency matters", f"={LC}O{chi + 1}", MONEY, "total net fee ÷ total hours worked")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))


README = """# Law Firm Pricing Workbook

Thank you for buying the workbook. Open `law-firm-pricing-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** annual overhead line by line, benefits and payroll taxes %, billable hours per timekeeper, your target margin on flat fees, scope contingency %, days in year, referral fee % and working days per month; overhead per timekeeper follows.
2. **Timekeepers:** each lawyer and paralegal with role, salary or draw, benefits % and billable hours (or your own) and bill rate; loaded salary, overhead share, total annual cost, cost rate, margin per hour, margin % and break-even hours follow.
3. **Flat Fee Pricer:** matter types with partner, associate and paralegal hours and disbursements; labour cost at average cost rates per role, cost with scope contingency, total cost, flat fee at your target margin, effective hourly rate and the gap to your current fee follow.
4. **Matter Tracker:** matters with fee type, budget and actual hours and fees, agreed fee, billed and collected; hours and fees variance, realized rate, write-off, recovery % and budget status follow.
5. **Lock-Up & Contingency:** monthly WIP, debtors and trailing 12-month fees with WIP days, debtor days, lock-up days and locked-up cash; contingency matters with gross fee, referral fee, net fee to the firm, net to the client and effective hourly rate.
6. **Summary:** timekeepers, average cost and bill rate, negative-margin timekeepers, under-priced flat fees and the uplift, matters over budget, write-offs, recovery %, lock-up days and locked-up cash, and net contingency fees.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not legal or financial advice; the targets are your own, not industry benchmarks. Licensed for personal or single-firm use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_timekeepers(wb.create_sheet("Timekeepers"))
    build_flat(wb.create_sheet("Flat Fee Pricer")); build_matters(wb.create_sheet("Matter Tracker"))
    build_lockup(wb.create_sheet("Lock-Up & Contingency")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
