"""Build the Freelancer Finance Kit workbook, README and zip. Usage: build_kit.py --out DIR"""
import argparse, datetime as dt, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'
DATE = "yyyy-mm-dd"
PCT = '0.0"%"'
INV_ROWS = (5, 104)
EXP_ROWS = (5, 204)


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]
    c.value = value
    c.font = BLUE
    if fmt:
        c.number_format = fmt
    if key:
        c.fill = YELLOW
    if note:
        c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]
    c.value = formula
    c.font = BOLD if bold else BLACK
    if fmt:
        c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]
    c.value = text
    c.font = BOLD if bold else (MUTED if muted else BLACK)
    return c


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Freelancer Finance Kit", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = [
        "",
        "HOW TO USE",
        "1. Settings: fill in the yellow cells once (your business name, default payment terms, late fee rate, hours and income goal).",
        "2. Invoices: add one row per invoice. Due date, days overdue, late fee and status are calculated for you.",
        "3. Rate Calculator: reads Settings and shows the hourly and day rate you need. Change Settings to explore.",
        "4. Payment Schedule: enter a project total, deposit %, number of milestones and dates to get a payment plan.",
        "5. Expenses: log costs with a date and category.",
        "6. Summary: pick a start month and see invoiced, paid, expenses and net by month and quarter.",
        "",
        "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs to set first.",
        "Each table has one example row showing the expected format. Overwrite it with your own data.",
        "",
        "NOTES",
        "Dates are calendar days. Late fee = amount × monthly rate ÷ 100 × days overdue ÷ 30 (simple interest, prorated by day).",
        "The workbook uses formulas only, no macros. It works in Excel 2010 or later, Google Sheets and Apple Numbers.",
        "This kit does arithmetic; it is not legal, tax or financial advice. Check enforceable late-fee rates for your jurisdiction.",
        "Support: hello@knackdesk.com",
    ]
    for i, t in enumerate(lines, start=3):
        label(ws, f"A{i}", t, bold=t.isupper() and t != "")


def build_settings(ws):
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    label(ws, "A1", "Settings", bold=True).font = TITLE
    rows = [
        ("Business name", "Your Studio", None, "Shown nowhere else; for your reference."),
        ("Currency label", "USD", None, "Text only; the kit does not convert currencies."),
        ("Default payment terms (days)", 30, "0", "Used when an invoice row leaves Terms blank. Net 30 = 30."),
        ("Late fee (% per month)", 1.5, "0.0", "Prorated by day. Use the rate in your contract; check local caps."),
        ("Hours in a billable day", 8, "0.0", None),
        ("Weeks off per year", 6, "0", "Holiday, sick days, public holidays, gaps between projects."),
        ("Hours worked per week", 40, "0", None),
        ("Billable share of hours (%)", 60, "0", "Most solo freelancers bill 50-70% of their hours."),
        ("Target income per year", 60000, MONEY, "What you want to keep after business expenses."),
        ("Business expenses per year", 6000, MONEY, "Software, insurance, equipment, accountant, training."),
        ("Tax rate to gross up for (%)", 0, "0.0", "Leave 0 if the target income is pre-tax."),
    ]
    for i, (name, val, fmt, note) in enumerate(rows, start=2):
        label(ws, f"A{i}", name)
        inp(ws, f"B{i}", val, fmt, key=True, note=note)
        if note:
            label(ws, f"C{i}", note, muted=True)
    label(ws, "A13", "Today (auto)")
    fx(ws, "B13", "=TODAY()", DATE)
    label(ws, "C13", "Recalculates each time the file opens; drives days overdue.", muted=True)


def build_invoices(ws):
    lo, hi = INV_ROWS
    label(ws, "A1", "Invoices", bold=True).font = TITLE
    label(ws, "A2", "Outstanding", bold=True)
    fx(ws, "B2", f'=SUMIFS(D{lo}:D{hi},J{lo}:J{hi},"Open")+SUMIFS(D{lo}:D{hi},J{lo}:J{hi},"Overdue")', MONEY, bold=True)
    label(ws, "C2", "Overdue", bold=True)
    fx(ws, "D2", f'=SUMIFS(D{lo}:D{hi},J{lo}:J{hi},"Overdue")', MONEY, bold=True)
    label(ws, "E2", "Late fees accrued", bold=True)
    fx(ws, "F2", f"=SUM(I{lo}:I{hi})", MONEY, bold=True)
    label(ws, "G2", "Paid to date", bold=True)
    fx(ws, "H2", f'=SUMIFS(D{lo}:D{hi},J{lo}:J{hi},"Paid")', MONEY, bold=True)
    head(ws, 4, ["Invoice #", "Client", "Issue date", "Amount", "Terms (days)", "Due date", "Paid date", "Days overdue", "Late fee", "Status", "Total due"],
         [12, 22, 13, 14, 12, 13, 13, 13, 12, 11, 14])
    for r in range(lo, hi + 1):
        for col in "ABCDEG":
            ws[f"{col}{r}"].font = BLUE
        ws[f"C{r}"].number_format = DATE
        ws[f"G{r}"].number_format = DATE
        ws[f"D{r}"].number_format = MONEY
        fx(ws, f"F{r}", f'=IF(C{r}="","",C{r}+IF(E{r}="",Settings!$B$4,E{r}))', DATE)
        fx(ws, f"H{r}", f'=IF(OR(C{r}="",G{r}<>""),0,MAX(0,Settings!$B$13-F{r}))', "0")
        fx(ws, f"I{r}", f"=IF(H{r}>0,ROUND(D{r}*Settings!$B$5/100*H{r}/30,2),0)", MONEY)
        fx(ws, f"J{r}", f'=IF(C{r}="","",IF(G{r}<>"","Paid",IF(H{r}>0,"Overdue","Open")))')
        fx(ws, f"K{r}", f'=IF(C{r}="","",D{r}+I{r})', MONEY)
    example = dt.date.today() - dt.timedelta(days=45)
    inp(ws, f"A{lo}", "INV-001"); inp(ws, f"B{lo}", "Example Client Ltd")
    inp(ws, f"C{lo}", example, DATE); inp(ws, f"D{lo}", 1500, MONEY); inp(ws, f"E{lo}", 30, "0")
    ws.freeze_panes = "A5"


def build_rate(ws):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 55
    label(ws, "A1", "Rate Calculator", bold=True).font = TITLE
    label(ws, "A2", "All inputs live on the Settings sheet.", muted=True)
    rows = [
        ("Working weeks", "=52-Settings!B7", "0", "52 minus weeks off"),
        ("Hours available", "=B4*Settings!B8", "0", "working weeks × hours per week"),
        ("Billable hours", "=ROUND(B5*Settings!B9/100,2)", "0.00", "hours available × billable share"),
        ("Revenue needed", "=Settings!B11+Settings!B10/(1-Settings!B12/100)", MONEY, "expenses + income grossed up for tax"),
        ("Hourly rate needed", '=IFERROR(ROUND(B7/B6,2),"check inputs")', MONEY, "revenue needed ÷ billable hours"),
        ("Day rate needed", '=IFERROR(ROUND(B8*Settings!B6,2),"check inputs")', MONEY, "hourly × hours in a billable day"),
    ]
    for i, (name, f, fmt, note) in enumerate(rows, start=4):
        label(ws, f"A{i}", name, bold=i >= 8)
        fx(ws, f"B{i}", f, fmt, bold=i >= 8)
        label(ws, f"C{i}", note, muted=True)


def build_schedule(ws):
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 50
    label(ws, "A1", "Payment Schedule", bold=True).font = TITLE
    inputs = [("Project total", 5000, MONEY), ("Deposit (%)", 30, "0"), ("Number of milestones (1-6)", 2, "0"),
              ("Start date", dt.date.today(), DATE), ("Project length (weeks)", 8, "0.0")]
    for i, (name, val, fmt) in enumerate(inputs, start=2):
        label(ws, f"A{i}", name)
        inp(ws, f"B{i}", val, fmt, key=True)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="6", showErrorMessage=True, errorTitle="Milestones", error="Enter a whole number from 1 to 6.")
    ws.add_data_validation(dv)
    dv.add("B4")
    head(ws, 8, ["Payment", "Due date", "Amount", "Note"])
    label(ws, "A9", "Deposit")
    fx(ws, "B9", "=$B$5", DATE)
    fx(ws, "C9", "=ROUND($B$2*$B$3/100,2)", MONEY)
    label(ws, "D9", "Due before work starts.", muted=True)
    for i in range(1, 7):
        r = 9 + i
        fx(ws, f"A{r}", f'=IF({i}>$B$4,"","Milestone {i}")')
        fx(ws, f"B{r}", f'=IF({i}>$B$4,"",$B$5+ROUND($B$6*7*{i}/$B$4,0))', DATE)
        if i == 1:
            amt = f'=IF(1>$B$4,"",IF(1<$B$4,ROUNDDOWN(($B$2-$C$9)/$B$4,2),$B$2-$C$9))'
        else:
            amt = f'=IF({i}>$B$4,"",IF({i}<$B$4,ROUNDDOWN(($B$2-$C$9)/$B$4,2),($B$2-$C$9)-SUM($C$10:C{r-1})))'
        fx(ws, f"C{r}", amt, MONEY)
        fx(ws, f"D{r}", f'=IF({i}>$B$4,"","Invoice when the milestone deliverable is presented.")').font = MUTED
    label(ws, "A16", "Total", bold=True)
    fx(ws, "C16", "=SUM(C9:C15)", MONEY, bold=True)
    label(ws, "D16", "Should equal the project total. The last milestone absorbs cent rounding.", muted=True)


def build_expenses(ws):
    lo, hi = EXP_ROWS
    label(ws, "A1", "Expenses", bold=True).font = TITLE
    label(ws, "A2", "Total", bold=True)
    fx(ws, "B2", f"=SUM(D{lo}:D{hi})", MONEY, bold=True)
    head(ws, 4, ["Date", "Category", "Description", "Amount", "Month"], [13, 18, 40, 14, 10])
    for r in range(lo, hi + 1):
        for col in "ABCD":
            ws[f"{col}{r}"].font = BLUE
        ws[f"A{r}"].number_format = DATE
        ws[f"D{r}"].number_format = MONEY
        fx(ws, f"E{r}", f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))')
    inp(ws, f"A{lo}", dt.date.today() - dt.timedelta(days=10), DATE)
    inp(ws, f"B{lo}", "Software"); inp(ws, f"C{lo}", "Example: design tool subscription"); inp(ws, f"D{lo}", 24, MONEY)
    ws.freeze_panes = "A5"


def build_summary(ws):
    ilo, ihi = INV_ROWS
    elo, ehi = EXP_ROWS
    ws.column_dimensions["A"].width = 14
    for col in "BCDE":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 15
    label(ws, "A1", "Summary", bold=True).font = TITLE
    label(ws, "A2", "Start month")
    first = dt.date.today().replace(day=1)
    inp(ws, "B2", first, DATE, key=True, note="First day of the first month to show. Twelve months follow.")
    head(ws, 4, ["Month", "Invoiced", "Paid", "Expenses", "Net"])
    for i in range(12):
        r = 5 + i
        fx(ws, f"A{r}", f"=EDATE($B$2,{i})", "mmm yyyy")
        fx(ws, f"B{r}", f'=SUMIFS(Invoices!$D${ilo}:$D${ihi},Invoices!$C${ilo}:$C${ihi},">="&A{r},Invoices!$C${ilo}:$C${ihi},"<"&EDATE(A{r},1))', MONEY)
        fx(ws, f"C{r}", f'=SUMIFS(Invoices!$D${ilo}:$D${ihi},Invoices!$G${ilo}:$G${ihi},">="&A{r},Invoices!$G${ilo}:$G${ihi},"<"&EDATE(A{r},1))', MONEY)
        fx(ws, f"D{r}", f'=SUMIFS(Expenses!$D${elo}:$D${ehi},Expenses!$A${elo}:$A${ehi},">="&A{r},Expenses!$A${elo}:$A${ehi},"<"&EDATE(A{r},1))', MONEY)
        fx(ws, f"E{r}", f"=C{r}-D{r}", MONEY)
    label(ws, "A17", "12 months", bold=True)
    for col in "BCDE":
        fx(ws, f"{col}17", f"=SUM({col}5:{col}16)", MONEY, bold=True)
    head(ws, 4, ["", "", "", "", "", "", "Quarter", "Net"])
    ws["G4"].value = "Quarter"; ws["H4"].value = "Net"
    for q in range(4):
        r = 5 + q
        fx(ws, f"G{r}", f'="Q"&{q + 1}')
        fx(ws, f"H{r}", f"=SUM(E{5 + q * 3}:E{7 + q * 3})", MONEY)
    label(ws, "G10", "Outstanding now", bold=True)
    fx(ws, "H10", "=Invoices!B2", MONEY, bold=True)
    label(ws, "G11", "Overdue now", bold=True)
    fx(ws, "H11", "=Invoices!D2", MONEY, bold=True)


README = """# Freelancer Finance Kit

Thank you for buying the kit. Open `freelancer-finance-kit.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. Start on the **Settings** sheet and fill the yellow cells.
2. Add invoices on the **Invoices** sheet; due dates, overdue days, late fees and status update automatically.
3. The **Rate Calculator**, **Payment Schedule**, **Expenses** and **Summary** sheets are explained on the **Start Here** sheet.

Blue text = type here. Black text = formula, leave it alone. Each table has one example row to overwrite.

No macros, no tracking. Licensed for personal or single-business use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Start Here"
    build_start(ws)
    build_settings(wb.create_sheet("Settings"))
    build_invoices(wb.create_sheet("Invoices"))
    build_rate(wb.create_sheet("Rate Calculator"))
    build_schedule(wb.create_sheet("Payment Schedule"))
    build_expenses(wb.create_sheet("Expenses"))
    build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    xlsx = os.path.join(args.out, "freelancer-finance-kit.xlsx")
    wb.save(xlsx)
    readme = os.path.join(args.out, "README.md")
    with open(readme, "w") as f:
        f.write(README)
    with zipfile.ZipFile(os.path.join(args.out, "freelancer-finance-kit.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, "freelancer-finance-kit.xlsx")
        z.write(readme, "README.md")
    print(f"built {xlsx}")


if __name__ == "__main__":
    main()
