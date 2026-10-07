"""Build the Insurance Agency Numbers Workbook, README and zip.
Usage: build_kit.py --out DIR [--zip-only]  (--zip-only re-zips an existing, recalculated xlsx with the README)"""
import argparse, os, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SLUG = "insurance-agency-numbers-workbook"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF"); BLACK = Font(name=FONT); BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14); MUTED = Font(name=FONT, italic=True, color="666666")
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E7E6E6")
MONEY = '#,##0.00;(#,##0.00);"-"'; PCT = '0.00"%"'; INT0 = "0"; HRS = "0.0"; RATIO = '0.00"x"'
BOOK_ROWS = (5, 16); PROD_ROWS = (5, 16); ACQ_ROWS = (5, 16); SPLIT_ROWS = (5, 24); SP_ROWS = (28, 33)
S = "Settings!$B$"  # settings cell prefix
BK = "'Book by Line'!"; PD = "Producers!"; AQ = "Acquisition!"; CS = "'Commission Splits'!"
BOOK_TOT = BOOK_ROWS[1] + 1; PROD_TOT = PROD_ROWS[1] + 1; ACQ_TOT = ACQ_ROWS[1] + 1; SPLIT_TOT = SPLIT_ROWS[1] + 1
NAMES = f"{PD}$A${PROD_ROWS[0]}:$A${PROD_ROWS[1]}"
# Settings rows (column B): referenced by every other sheet.
R_NB = 2; R_REN = 3; R_SPLIT = 4; R_FEE = 5; R_TRET = 6; R_TCOMP = 7; R_COSTHR = 8; R_YEARS = 9; R_STAFF = 10
R_PRODS = 11; R_EMPS = 12; R_POLS = 13; R_BOOKREN = 14; R_BOOKRET = 15


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w


def inp(ws, ref, value, fmt=None, key=False, note=None):
    c = ws[ref]; c.value = value; c.font = BLUE
    if fmt: c.number_format = fmt
    if key: c.fill = YELLOW
    if note: c.comment = Comment(note, "Knackdesk")
    return c


def fx(ws, ref, formula, fmt=None, bold=False):
    c = ws[ref]; c.value = formula; c.font = BOLD if bold else BLACK
    if fmt: c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, muted=False):
    c = ws[ref]; c.value = text; c.font = BOLD if bold else (MUTED if muted else BLACK); return c


def blank(ws, r, cols, fmts=None):
    for col in cols:
        ws[f"{col}{r}"].font = BLUE
        if fmts and col in fmts: ws[f"{col}{r}"].number_format = fmts[col]


def rows_block(ws, rows, bold_rows=()):
    """Write (row, label, formula, format, note) lines as label / bold formula / muted note."""
    for r, name, f, fmt, note in rows:
        b = r in bold_rows
        label(ws, f"A{r}", name, bold=b); fx(ws, f"B{r}", f, fmt, bold=b)
        if note: label(ws, f"C{r}", note, muted=True)


def note(ws, ref, text):
    ws[ref].comment = Comment(text, "Knackdesk")


def list_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(dv)
    for ref in cells: dv.add(ref)


def build_start(ws):
    ws.column_dimensions["A"].width = 100
    label(ws, "A1", "Insurance Agency Numbers Workbook", bold=True).font = TITLE
    label(ws, "A2", "by Knackdesk · knackdesk.com", muted=True)
    lines = ["", "HOW TO USE",
        "1. Settings: default new-business and renewal commission %, default agent split %, agency fee per policy, your target retention % and target compensation ratio %, producer cost per hour, years for lifetime commission and support staff count. Producers, employees, policies on the book, the book's renewal commission and the book's retention % are calculated from the other sheets.",
        "2. Book by Line: one row per line of business with policies at the start of the year, policies lost, new policies and average premium; renewal and new-business commission % default from Settings (type your own to override). Policies at the end, retention %, lapse %, net growth %, renewal commission on the retained book, new-business commission, commission lost and retention against your target are calculated, with totals.",
        "3. Producers: one row per producer with status, policies written, premium written, the commission revenue credited to them, compensation paid and hours on prospecting. Revenue per policy, compensation ratio %, a new-business commission check at the Settings % and the ratio against your target are calculated, with totals and a per-producer average.",
        "4. Acquisition: one row per month with marketing spend, lead costs, producer hours (default from Producers, type your own to override), policies written and first-year and renewal commission per policy (defaults from Settings and the book's average premium). Total acquisition cost, cost per policy, lifetime commission per policy, value-to-cost ratio and payback policies are calculated.",
        "5. Commission Splits: one row per policy or client with producer, premium and carrier commission %; agent split % and agency fee default from Settings (type your own to override). Gross commission, agent share, agency share, agent net after fee, the agent's effective % of premium and agency net are calculated, with a per-producer block below.",
        "6. Summary: policies on the book and retention against your target, renewal commission and commission lost, lines below target, revenue per producer and per employee, compensation ratio, acquisition cost, lifetime commission and value-to-cost, the best line and best producer, and total agent and agency net from splits.",
        "", "LEGEND",
        "Blue text = a cell you type into.   Black text = a formula, leave it alone.   Yellow fill = key inputs.",
        "", "DEFINITIONS",
        "Producer: a licensed agent who writes new business for the agency and is usually paid a share of the commission it earns.",
        "Book of business: all the policies the agency services and earns commission on. Policies at the end = policies at the start − policies lost + new policies.",
        "Policy retention: the share of the policies on the book at the start of the year that are still on it at the end. Retention % = (policies at start − policies lost) ÷ policies at start. New policies are not counted in retention.",
        "Lapse: a policy that is cancelled or not renewed. Lapse % = policies lost ÷ policies at start = 100% − retention %.",
        "Renewal commission and residuals: the commission the carrier pays each year a policy renews, usually a lower % than in the first year. Annual renewal commission on the retained book = (policies at start − policies lost) × average premium × renewal commission %.",
        "Revenue per producer: agency commission revenue (renewal plus new business on Book by Line) ÷ number of producers. Revenue per employee divides the same revenue by producers plus support staff.",
        "Cost per policy acquired: marketing spend + lead costs + producer hours × producer cost per hour, divided by the policies written in the period.",
        "Commission split: how the commission a carrier pays on a policy is shared between the producer (agent split %) and the agency, with any per-policy agency fee charged to the agent. Agent net = agent share − agency fee; agency net = agency share + agency fee.",
        "The targets in this workbook are your own; it does not supply industry benchmarks or valuation multiples. Set them from your own history and plan.",
        "Commission structures are set by your carriers and producer contracts; the workbook applies the terms you enter and does not check them.",
        "Support: hello@knackdesk.com",
        "Arithmetic only; not financial, legal or licensing advice."]
    for i, t in enumerate(lines, start=3): label(ws, f"A{i}", t, bold=t.isupper() and t != "")


SETTINGS_INPUTS = [
    (R_NB, "Default new-business commission %", 15, PCT, True, "Default 15. Placeholder; use the first-year commission % in your carrier contracts. Each line can override it on Book by Line."),
    (R_REN, "Default renewal commission %", 10, PCT, True, "Default 10. Placeholder; use the renewal commission % in your carrier contracts. Each line can override it on Book by Line."),
    (R_SPLIT, "Default agent split %", 60, PCT, True, "Default 60. Placeholder; the producer's share of the commission in your producer agreements. Each split can override it."),
    (R_FEE, "Agency fee per policy", 10, MONEY, False, "Default 10. Placeholder; a per-policy fee charged to the producer on Commission Splits. Type 0 if you do not charge one."),
    (R_TRET, "Target retention %", 90, PCT, True, "Default 90. Your own placeholder, not a recommendation and not an industry benchmark: the retention you are aiming for."),
    (R_TCOMP, "Target compensation ratio %", 35, PCT, True, "Default 35. Your own placeholder, not a recommendation and not an industry benchmark: producer pay as a % of the commission revenue credited to them."),
    (R_COSTHR, "Producer cost per hour for acquisition time", 45, MONEY, False, "Default 45. Placeholder; what an hour of producer prospecting time costs you, used in Acquisition cost."),
    (R_YEARS, "Years for lifetime commission (1 to 5)", 3, INT0, True, "Default 3. How many years of commission (first year plus renewals) a new policy is valued at on Acquisition; up to 5."),
    (R_STAFF, "Support staff count", 3, INT0, False, "Default 3. Account managers, CSRs and admin who are not producers; used for revenue per employee.")]


def build_settings(ws):
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 90
    label(ws, "A1", "Settings", bold=True).font = TITLE
    for r, name, val, fmt, key, txt in SETTINGS_INPUTS:
        label(ws, f"A{r}", name); inp(ws, f"B{r}", val, fmt, key=key); label(ws, f"C{r}", txt, muted=True)
    note(ws, f"B{R_NB}", "Assumption: 15% is a placeholder. First-year commission varies by carrier and line; override it per line.")
    note(ws, f"B{R_REN}", "Assumption: 10% is a placeholder. Renewal commission varies by carrier and line; override it per line.")
    note(ws, f"B{R_SPLIT}", "Assumption: 60% is a placeholder. Use the split in each producer's agreement.")
    note(ws, f"B{R_TRET}", "Your own placeholder, not a recommendation and not an industry benchmark.")
    note(ws, f"B{R_TCOMP}", "Your own placeholder, not a recommendation and not an industry benchmark.")
    note(ws, f"B{R_COSTHR}", "Assumption: 45 an hour is a placeholder. Use producer pay plus payroll costs ÷ hours worked.")
    note(ws, f"B{R_YEARS}", "Whole years from 1 to 5. 1 = first-year commission only.")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"B{R_YEARS}")
    plo, phi = PROD_ROWS
    rows_block(ws, [(R_PRODS, "Producers", f"=COUNTA({PD}A{plo}:A{phi})", INT0,
                     "formula: names on the Producers sheet (including anyone who left during the year)"),
                    (R_EMPS, "Total employees", f"=B{R_PRODS}+B{R_STAFF}", INT0,
                     "formula: producers + support staff count"),
                    (R_POLS, "Policies on the book now", f"={BK}L{BOOK_TOT}", INT0,
                     "formula: policies at the end of the year, from the Book by Line totals"),
                    (R_BOOKREN, "Book annual renewal commission", f"={BK}P{BOOK_TOT}", MONEY,
                     "formula: annual renewal commission on the retained book, from the Book by Line totals"),
                    (R_BOOKRET, "Book retention %", f"={BK}M{BOOK_TOT}", PCT,
                     "formula: total retention from the Book by Line totals; Acquisition values renewals at this rate")],
               bold_rows=(R_PRODS, R_EMPS, R_POLS, R_BOOKREN, R_BOOKRET))


def build_book(ws):
    lo, hi = BOOK_ROWS; tot = BOOK_TOT
    label(ws, "A1", "Book by Line", bold=True).font = TITLE
    label(ws, "A2", "One row per line of business. Renewal and new-business commission % default from Settings; type your own to override them. Retention counts only the policies on the book at the start of the year.", muted=True)
    head(ws, 4, ["Line of business", "Policies at start of year", "Policies lost", "New policies", "Average premium",
                 "Default renewal commission % (Settings)", "or type your own renewal %", "Renewal % used",
                 "Default new-business commission % (Settings)", "or type your own new-business %",
                 "New-business % used", "Policies at end", "Retention %", "Lapse %", "Net growth %",
                 "Annual renewal commission (retained book)", "New-business commission", "Commission lost to lapses",
                 "Retention vs target (points)"],
         [18, 10, 9, 9, 11, 11, 10, 9, 11, 10, 9, 9, 9, 9, 9, 13, 12, 12, 11])
    fm = {"B": INT0, "C": INT0, "D": INT0, "E": MONEY, "G": PCT, "J": PCT}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCDEGJ", fm); book_row(ws, r)
    sample = [("Personal auto", 420, 46, 70, 1400, None, None), ("Homeowners", 310, 25, 40, 1600, None, None),
              ("Commercial", 85, 6, 12, 5200, 12, None), ("Life", 120, 5, 15, 900, 2, 50),
              ("Health", 60, 9, 10, 2400, None, None), ("Umbrella", 75, 4, 10, 450, None, None)]
    for i, (line, start, lost, new, prem, ren, nb) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", line); inp(ws, f"B{r}", start, INT0, key=True); inp(ws, f"C{r}", lost, INT0, key=True)
        inp(ws, f"D{r}", new, INT0, key=True); inp(ws, f"E{r}", prem, MONEY, key=True)
        inp(ws, f"G{r}", ren, PCT); inp(ws, f"J{r}", nb, PCT)
    book_notes(ws, lo)
    book_totals(ws, lo, hi, tot)
    ws.freeze_panes = "B5"


def book_row(ws, r):
    g = f'=IF(A{r}="","",{{}})'
    fx(ws, f"F{r}", g.format(f"{S}{R_REN}"), PCT)
    fx(ws, f"H{r}", g.format(f'IF(G{r}<>"",G{r},F{r})'), PCT)
    fx(ws, f"I{r}", g.format(f"{S}{R_NB}"), PCT)
    fx(ws, f"K{r}", g.format(f'IF(J{r}<>"",J{r},I{r})'), PCT)
    fx(ws, f"L{r}", g.format(f"B{r}-C{r}+D{r}"), INT0)
    fx(ws, f"M{r}", g.format(f"IFERROR(ROUND((B{r}-C{r})/B{r}*100,2),0)"), PCT)
    fx(ws, f"N{r}", g.format(f"IFERROR(ROUND(C{r}/B{r}*100,2),0)"), PCT)
    fx(ws, f"O{r}", g.format(f"IFERROR(ROUND((L{r}-B{r})/B{r}*100,2),0)"), PCT)
    fx(ws, f"P{r}", g.format(f"ROUND((B{r}-C{r})*E{r}*H{r}/100,2)"), MONEY)
    fx(ws, f"Q{r}", g.format(f"ROUND(D{r}*E{r}*K{r}/100,2)"), MONEY)
    fx(ws, f"R{r}", g.format(f"ROUND(C{r}*E{r}*H{r}/100,2)"), MONEY)
    fx(ws, f"S{r}", g.format(f"ROUND(M{r}-{S}{R_TRET},2)"), PCT)


def book_notes(ws, lo):
    note(ws, f"B{lo}", "Policies in force on this line on the first day of the year.")
    note(ws, f"C{lo}", "Policies from the start-of-year book that lapsed, were cancelled or were not renewed during the year.")
    note(ws, f"E{lo}", "Assumption: the average annual premium per policy on this line; commission is a % of it.")
    note(ws, f"G{lo + 2}", "Example override: a carrier that pays a higher renewal commission on commercial lines.")
    note(ws, f"J{lo + 3}", "Example override: life policies often pay a high first-year commission and a small renewal; use your contract terms.")
    note(ws, f"M{lo}", "(Policies at start − policies lost) ÷ policies at start. 0 when no start policies are typed.")
    note(ws, f"P{lo}", "Assumption: every retained policy renews at the average premium and renewal % on this row.")
    note(ws, f"R{lo}", "Renewal commission the lapsed policies would have paid this year at the same premium and renewal %.")
    note(ws, f"S{lo}", "Retention % − your target retention % on Settings. Negative = below your own target.")


def book_totals(ws, lo, hi, tot):
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "BCDL":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", INT0, bold=True)
    for col in "PQR":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"E{tot}", f"=IFERROR(ROUND(SUMPRODUCT(B{lo}:B{hi},E{lo}:E{hi})/B{tot},2),0)", MONEY, bold=True)
    fx(ws, f"M{tot}", f"=IFERROR(ROUND((B{tot}-C{tot})/B{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"N{tot}", f"=IFERROR(ROUND(C{tot}/B{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND((L{tot}-B{tot})/B{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"S{tot}", f"=ROUND(M{tot}-{S}{R_TRET},2)", PCT, bold=True)
    note(ws, f"E{tot}", "Average premium weighted by policies at the start of the year. Acquisition uses it for the default commission per policy.")
    note(ws, f"M{tot}", "From the totals, not an average of the rows.")


def build_producers(ws):
    lo, hi = PROD_ROWS; tot = PROD_TOT; avg = tot + 1
    label(ws, "A1", "Producers", bold=True).font = TITLE
    label(ws, "A2", "One row per producer for the year, including anyone who left. Commission revenue credited, compensation paid and hours come from your own records; the check column applies the Settings new-business % to premium written.", muted=True)
    head(ws, 4, ["Producer", "Status", "Policies written this year", "Premium written", "Commission revenue credited",
                 "Compensation paid", "Hours on prospecting", "Revenue per policy", "Compensation ratio %",
                 "New-business commission at Settings % (check)", "Credited − check", "Compensation ratio vs target (points)"],
         [18, 10, 10, 13, 13, 13, 10, 11, 11, 13, 12, 13])
    fm = {"C": INT0, "D": MONEY, "E": MONEY, "F": MONEY, "G": HRS}
    rows = range(lo, hi + 1)
    for r in rows:
        blank(ws, r, "ABCDEFG", fm)
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"H{r}", g.format(f"IFERROR(ROUND(E{r}/C{r},2),0)"), MONEY)
        fx(ws, f"I{r}", g.format(f"IFERROR(ROUND(F{r}/E{r}*100,2),0)"), PCT)
        fx(ws, f"J{r}", g.format(f"ROUND(D{r}*{S}{R_NB}/100,2)"), MONEY)
        fx(ws, f"K{r}", g.format(f"ROUND(E{r}-J{r},2)"), MONEY)
        fx(ws, f"L{r}", g.format(f"ROUND(I{r}-{S}{R_TCOMP},2)"), PCT)
    list_dv(ws, '"active,left"', [f"B{r}" for r in rows])
    sample = [("A. Rivera", "active", 140, 210000, 34000, 12000, 380), ("B. Cole", "active", 95, 160000, 25500, 9600, 300),
              ("C. Park", "active", 120, 150000, 23000, 7800, 340), ("D. Shah", "active", 60, 85000, 13500, 5400, 260),
              ("E. Grant", "left", 30, 40000, 6200, 3100, 120)]
    for i, (name, status, pols, prem, rev, comp, hrs) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", name); inp(ws, f"B{r}", status); inp(ws, f"C{r}", pols, INT0)
        inp(ws, f"D{r}", prem, MONEY); inp(ws, f"E{r}", rev, MONEY, key=True)
        inp(ws, f"F{r}", comp, MONEY, key=True); inp(ws, f"G{r}", hrs, HRS)
    note(ws, f"A{lo + 4}", "Example: a producer who left mid-year. Keep the row so the year's totals stay complete.")
    note(ws, f"E{lo}", "Commission revenue the agency credits to this producer for the year (new business and renewals), from your own records.")
    note(ws, f"F{lo}", "Everything paid to the producer for the year: salary, commission share and bonuses.")
    note(ws, f"G{lo}", "Hours spent prospecting for new business. Acquisition divides the total by 12 for its default monthly hours.")
    note(ws, f"I{lo}", "Compensation paid ÷ commission revenue credited. 0 when no revenue is typed.")
    note(ws, f"J{lo}", "Assumption: premium written × the default new-business % on Settings. A check only; the credited figure is yours.")
    note(ws, f"L{lo}", "Compensation ratio % − your target on Settings. Positive = above your own target.")
    label(ws, f"A{tot}", "Total", bold=True)
    fx(ws, f"B{tot}", f"=COUNTA(A{lo}:A{hi})", '0" producers"', bold=True)
    for col, fmt in (("C", INT0), ("D", MONEY), ("E", MONEY), ("F", MONEY), ("G", HRS), ("J", MONEY), ("K", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    fx(ws, f"H{tot}", f"=IFERROR(ROUND(E{tot}/C{tot},2),0)", MONEY, bold=True)
    fx(ws, f"I{tot}", f"=IFERROR(ROUND(F{tot}/E{tot}*100,2),0)", PCT, bold=True)
    fx(ws, f"L{tot}", f"=ROUND(I{tot}-{S}{R_TCOMP},2)", PCT, bold=True)
    label(ws, f"A{avg}", "Per-producer average", bold=True)
    for col, fmt in (("C", HRS), ("D", MONEY), ("E", MONEY), ("F", MONEY), ("G", HRS), ("J", MONEY)):
        fx(ws, f"{col}{avg}", f"=IFERROR(ROUND({col}{tot}/$B${tot},2),0)", fmt, bold=True)
    note(ws, f"I{tot}", "From the totals, not an average of the rows.")
    ws.freeze_panes = "B5"


def build_acquisition(ws):
    lo, hi = ACQ_ROWS; tot = ACQ_TOT
    label(ws, "A1", "Acquisition", bold=True).font = TITLE
    label(ws, "A2", "One row per month. Producer hours default from Producers ÷ 12; commission per policy defaults from the Settings % on the book's average premium. Type your own to override them.", muted=True)
    head(ws, 4, ["Month", "Marketing spend", "Lead costs", "Default producer hours (Producers ÷ 12)",
                 "or type your own hours", "Producer hours used", "Policies written",
                 "Default first-year commission per policy", "or type your own first-year", "First-year commission used",
                 "Default renewal commission per policy", "or type your own renewal", "Renewal commission used",
                 "Total acquisition cost", "Cost per policy", "Lifetime commission per policy", "Value-to-cost ratio",
                 "Payback policies"],
         [12, 11, 11, 11, 10, 10, 9, 12, 11, 11, 12, 11, 11, 12, 11, 12, 10, 10])
    fm = {"B": MONEY, "C": MONEY, "E": HRS, "G": INT0, "I": MONEY, "L": MONEY}
    for r in range(lo, hi + 1):
        blank(ws, r, "ABCEGIL", fm); acquisition_row(ws, r)
    sample = [("Jan", 800, 600, None, 22, None, None), ("Feb", 800, 450, None, 18, None, None),
              ("Mar", 1200, 700, 140, 30, None, None), ("Apr", 900, 500, None, 21, None, None),
              ("May", 900, 650, None, 24, 260, None), ("Jun", 600, 400, None, 15, None, None)]
    for i, (month, mkt, leads, hrs, pols, fy, ren) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", month); inp(ws, f"B{r}", mkt, MONEY, key=True); inp(ws, f"C{r}", leads, MONEY, key=True)
        inp(ws, f"E{r}", hrs, HRS); inp(ws, f"G{r}", pols, INT0, key=True); inp(ws, f"I{r}", fy, MONEY)
        inp(ws, f"L{r}", ren, MONEY)
    acquisition_notes(ws, lo)
    acquisition_totals(ws, lo, hi, tot)
    ws.freeze_panes = "B5"


def lifetime(fy, ren):
    """First-year + renewal × (r + r² + ... ) for the Settings years (up to 5), r = book retention."""
    y = f"{S}{R_YEARS}"; rr = f"{BK}$M${BOOK_TOT}/100"
    terms = "+".join([f"IF({y}>=2,{rr},0)"] + [f"IF({y}>={n + 1},POWER({rr},{n}),0)" for n in (2, 3, 4)])
    return f"ROUND({fy}+{ren}*({terms}),2)"


def acquisition_row(ws, r):
    g = f'=IF(A{r}="","",{{}})'
    fx(ws, f"D{r}", g.format(f"ROUND({PD}$G${PROD_TOT}/12,1)"), HRS)
    fx(ws, f"F{r}", g.format(f'IF(E{r}<>"",E{r},D{r})'), HRS)
    fx(ws, f"H{r}", g.format(f"ROUND({S}{R_NB}/100*{BK}$E${BOOK_TOT},2)"), MONEY)
    fx(ws, f"J{r}", g.format(f'IF(I{r}<>"",I{r},H{r})'), MONEY)
    fx(ws, f"K{r}", g.format(f"ROUND({S}{R_REN}/100*{BK}$E${BOOK_TOT},2)"), MONEY)
    fx(ws, f"M{r}", g.format(f'IF(L{r}<>"",L{r},K{r})'), MONEY)
    fx(ws, f"N{r}", g.format(f"ROUND(B{r}+C{r}+F{r}*{S}{R_COSTHR},2)"), MONEY)
    fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/G{r},2),0)"), MONEY)
    fx(ws, f"P{r}", g.format(lifetime(f"J{r}", f"M{r}")), MONEY)
    fx(ws, f"Q{r}", g.format(f"IFERROR(ROUND(P{r}/O{r},2),0)"), RATIO)
    fx(ws, f"R{r}", g.format(f"IFERROR(ROUND(N{r}/J{r},1),0)"), HRS)


def acquisition_notes(ws, lo):
    note(ws, f"B{lo}", "Advertising, website, mailers and sponsorships for the month.")
    note(ws, f"C{lo}", "Bought leads, quote-comparison sites and referral fees for the month.")
    note(ws, f"E{lo + 2}", "Example override: a month with an extra prospecting push.")
    note(ws, f"I{lo + 4}", "Example override: a month weighted to commercial policies with a higher first-year commission.")
    note(ws, f"D{lo}", "Assumption: the year's prospecting hours on Producers spread evenly over 12 months.")
    note(ws, f"H{lo}", "Assumption: default new-business % on Settings × the average premium on Book by Line totals.")
    note(ws, f"N{lo}", "Marketing spend + lead costs + producer hours × producer cost per hour on Settings.")
    note(ws, f"P{lo}", "Assumption: first-year commission + renewal commission × (r + r² + ...) for the years on Settings (up to 5), where r is the total retention % on Book by Line.")
    note(ws, f"Q{lo}", "Lifetime commission per policy ÷ cost per policy. Above 1 = a policy earns back more than it cost to acquire.")
    note(ws, f"R{lo}", "Total acquisition cost ÷ first-year commission per policy: how many new policies' first-year commission pays back the month's cost.")


def acquisition_totals(ws, lo, hi, tot):
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("B", MONEY), ("C", MONEY), ("F", HRS), ("G", INT0), ("N", MONEY)):
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", fmt, bold=True)
    for col in "JM":
        fx(ws, f"{col}{tot}", f"=IFERROR(ROUND(SUMPRODUCT(G{lo}:G{hi},{col}{lo}:{col}{hi})/G{tot},2),0)", MONEY, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(N{tot}/G{tot},2),0)", MONEY, bold=True)
    fx(ws, f"P{tot}", "=" + lifetime(f"J{tot}", f"M{tot}"), MONEY, bold=True)
    fx(ws, f"Q{tot}", f"=IFERROR(ROUND(P{tot}/O{tot},2),0)", RATIO, bold=True)
    fx(ws, f"R{tot}", f"=IFERROR(ROUND(N{tot}/J{tot},1),0)", HRS, bold=True)
    note(ws, f"J{tot}", "Average per policy, weighted by policies written each month.")
    note(ws, f"O{tot}", "From the totals, not an average of the months.")


def build_splits(ws):
    lo, hi = SPLIT_ROWS; tot = SPLIT_TOT
    label(ws, "A1", "Commission Splits", bold=True).font = TITLE
    label(ws, "A2", "One row per policy or client. Agent split % and agency fee default from Settings; type your own to override them. The carrier commission % is the one in your carrier contract for this policy.", muted=True)
    head(ws, 4, ["Policy or client", "Producer", "Premium", "Carrier commission %", "Default agent split % (Settings)",
                 "or type your own split %", "Agent split % used", "Default agency fee (Settings)",
                 "or type your own fee", "Agency fee used", "Gross commission", "Agent share", "Agency share",
                 "Agent net after fee", "Agent effective % of premium", "Agency net"],
         [20, 14, 11, 10, 10, 10, 9, 10, 10, 9, 11, 11, 11, 11, 10, 11])
    fm = {"C": MONEY, "D": PCT, "F": PCT, "I": MONEY}
    rows = range(lo, hi + 1)
    for r in rows:
        blank(ws, r, "ABCDFI", fm); split_row(ws, r)
    list_dv(ws, NAMES, [f"B{r}" for r in rows])
    sample = [("Miller auto", "A. Rivera", 1450, 15, None, None), ("Lopez home", "A. Rivera", 1700, 15, None, None),
              ("Harbor Cafe BOP", "B. Cole", 4800, 12, 50, None), ("Nguyen life", "C. Park", 1100, 60, None, None),
              ("Patel auto", "C. Park", 1300, 15, None, None), ("Okafor home", "B. Cole", 1550, 15, None, 0),
              ("Summit Builders GL", "D. Shah", 7200, 12, None, None), ("Reyes umbrella", "D. Shah", 480, 15, None, None),
              ("Kim health", "A. Rivera", 2600, 8, None, None), ("Baker auto", "E. Grant", 1250, 15, None, None)]
    for i, (pol, prod, prem, carrier, split, fee) in enumerate(sample):
        r = lo + i
        inp(ws, f"A{r}", pol); inp(ws, f"B{r}", prod); inp(ws, f"C{r}", prem, MONEY, key=True)
        inp(ws, f"D{r}", carrier, PCT, key=True); inp(ws, f"F{r}", split, PCT); inp(ws, f"I{r}", fee, MONEY)
    note(ws, f"D{lo}", "The commission % the carrier pays the agency on this policy, from your carrier contract.")
    note(ws, f"F{lo + 2}", "Example override: a lower split on a house account the producer took over.")
    note(ws, f"I{lo + 5}", "Example override: no agency fee on this policy.")
    note(ws, f"N{lo}", "Agent share − agency fee. Can be negative on a small policy where the fee is larger than the share.")
    note(ws, f"O{lo}", "Agent net ÷ premium: the share of the premium the producer keeps.")
    label(ws, f"A{tot}", "Total", bold=True)
    for col in "CJKLMNP":
        fx(ws, f"{col}{tot}", f"=SUM({col}{lo}:{col}{hi})", MONEY, bold=True)
    fx(ws, f"O{tot}", f"=IFERROR(ROUND(N{tot}/C{tot}*100,2),0)", PCT, bold=True)
    build_split_producers(ws)
    ws.freeze_panes = "B5"


def split_row(ws, r):
    g = f'=IF(A{r}="","",{{}})'
    fx(ws, f"E{r}", g.format(f"{S}{R_SPLIT}"), PCT)
    fx(ws, f"G{r}", g.format(f'IF(F{r}<>"",F{r},E{r})'), PCT)
    fx(ws, f"H{r}", g.format(f"{S}{R_FEE}"), MONEY)
    fx(ws, f"J{r}", g.format(f'IF(I{r}<>"",I{r},H{r})'), MONEY)
    fx(ws, f"K{r}", g.format(f"ROUND(C{r}*D{r}/100,2)"), MONEY)
    fx(ws, f"L{r}", g.format(f"ROUND(K{r}*G{r}/100,2)"), MONEY)
    fx(ws, f"M{r}", g.format(f"ROUND(K{r}-L{r},2)"), MONEY)
    fx(ws, f"N{r}", g.format(f"ROUND(L{r}-J{r},2)"), MONEY)
    fx(ws, f"O{r}", g.format(f"IFERROR(ROUND(N{r}/C{r}*100,2),0)"), PCT)
    fx(ws, f"P{r}", g.format(f"ROUND(M{r}+J{r},2)"), MONEY)


def build_split_producers(ws):
    lo, hi = SPLIT_ROWS; plo, phi = SP_ROWS
    label(ws, f"A{plo - 2}", "PER-PRODUCER TOTALS", bold=True)
    labels = ["Producer", "Policies", "Premium", "Gross commission", "Agent net after fee", "Agency net"]
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=plo - 1, column=i, value=text); c.font = BOLD; c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    names = f"$B${lo}:$B${hi}"
    for r in range(plo, phi + 1):
        blank(ws, r, "A")
        g = f'=IF(A{r}="","",{{}})'
        fx(ws, f"B{r}", g.format(f"COUNTIF({names},A{r})"), INT0)
        for dst, src in (("C", "C"), ("D", "K"), ("E", "N"), ("F", "P")):
            fx(ws, f"{dst}{r}", g.format(f"SUMIF({names},A{r},${src}${lo}:${src}${hi})"), MONEY)
    for i, name in enumerate(["A. Rivera", "B. Cole", "C. Park", "D. Shah", "E. Grant"]):
        inp(ws, f"A{plo + i}", name)
    list_dv(ws, NAMES, [f"A{r}" for r in range(plo, phi + 1)])
    note(ws, f"A{plo}", "Type or pick each producer once; their splits above are counted and summed by name.")


def build_summary(ws):
    blo, bhi = BOOK_ROWS; plo, phi = PROD_ROWS
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 80
    label(ws, "A1", "Summary", bold=True).font = TITLE
    growth = f"{BK}$O${blo}:$O${bhi}"; rev = f"{PD}$E${plo}:$E${phi}"
    rows = [(2, "Policies on the book now", f"={S}{R_POLS}", INT0, "Book by Line totals: start − lost + new"),
            (3, "Retention %", f"={BK}M{BOOK_TOT}", PCT, "Book by Line totals: (start − lost) ÷ start"),
            (4, "Target retention % (your own)", f"={S}{R_TRET}", PCT, "from Settings: your own placeholder, not a benchmark"),
            (5, "Retention vs target (points)", "=ROUND(B3-B4,2)", PCT, "negative = below your own target"),
            (6, "Annual renewal commission on the retained book", f"={BK}P{BOOK_TOT}", MONEY, "retained policies × average premium × renewal %"),
            (7, "Commission lost to lapses", f"={BK}R{BOOK_TOT}", MONEY, "lost policies × average premium × renewal %"),
            (8, "Lines below target retention", f'=COUNTIF({BK}$S${blo}:$S${bhi},"<0")', INT0, "Book by Line rows with retention below your target"),
            (9, "Agency commission revenue", f"=ROUND({BK}P{BOOK_TOT}+{BK}Q{BOOK_TOT},2)", MONEY, "Book by Line: renewal commission + new-business commission"),
            (10, "Producers", f"={S}{R_PRODS}", INT0, "names on the Producers sheet"),
            (11, "Revenue per producer", "=IFERROR(ROUND(B9/B10,2),0)", MONEY, "agency commission revenue ÷ producers"),
            (12, "Revenue per employee", f"=IFERROR(ROUND(B9/{S}{R_EMPS},2),0)", MONEY, "agency commission revenue ÷ (producers + support staff)"),
            (13, "Compensation ratio %", f"={PD}I{PROD_TOT}", PCT, "Producers totals: compensation paid ÷ commission revenue credited"),
            (14, "Compensation ratio vs target (points)", f"=ROUND(B13-{S}{R_TCOMP},2)", PCT, "positive = above your own target"),
            (15, "Acquisition cost year to date", f"={AQ}N{ACQ_TOT}", MONEY, "Acquisition totals: marketing + leads + producer time"),
            (16, "Cost per policy acquired", f"={AQ}O{ACQ_TOT}", MONEY, "Acquisition totals: cost ÷ policies written"),
            (17, "Lifetime commission per policy", f"={AQ}P{ACQ_TOT}", MONEY, "first-year + renewals over the Settings years at book retention"),
            (18, "Value-to-cost ratio", f"={AQ}Q{ACQ_TOT}", RATIO, "lifetime commission per policy ÷ cost per policy"),
            (19, "Best line by net growth", f'=IFERROR(INDEX({BK}$A${blo}:$A${bhi},MATCH(MAX({growth}),{growth},0)),"")', None, "highest net growth % on Book by Line"),
            (20, "Its net growth %", f"=IFERROR(MAX({growth}),0)", PCT, "policies at end vs start"),
            (21, "Best producer by commission revenue", f'=IFERROR(INDEX({PD}$A${plo}:$A${phi},MATCH(MAX({rev}),{rev},0)),"")', None, "highest commission revenue credited on Producers"),
            (22, "Their commission revenue", f"=IFERROR(MAX({rev}),0)", MONEY, "commission revenue credited"),
            (23, "Total agent net from splits", f"={CS}N{SPLIT_TOT}", MONEY, "Commission Splits totals: agent share − agency fees"),
            (24, "Total agency net from splits", f"={CS}P{SPLIT_TOT}", MONEY, "Commission Splits totals: agency share + agency fees")]
    rows_block(ws, rows, bold_rows=tuple(r for r, *_ in rows))
    note(ws, "B4", "Your own placeholder, not a recommendation and not an industry benchmark.")
    note(ws, "B11", "Arithmetic only. The workbook supplies no benchmarks or valuation multiples to compare it with.")


README = """# Insurance Agency Numbers Workbook

Thank you for buying the workbook. Open `insurance-agency-numbers-workbook.xlsx` in Excel (2010 or later), Google Sheets (File > Import) or Apple Numbers.

1. **Settings:** default new-business and renewal commission %, default agent split %, agency fee per policy, your own target retention % and target compensation ratio %, producer cost per hour, years for lifetime commission and support staff count; producers, employees, policies on the book, the book's renewal commission and retention % follow.
2. **Book by Line:** each line of business with policies at the start of the year, policies lost, new policies, average premium and renewal and new-business commission % (or your own); policies at the end, retention %, lapse %, net growth %, renewal commission on the retained book, new-business commission, commission lost and retention against your target follow.
3. **Producers:** each producer with status, policies and premium written, commission revenue credited, compensation paid and prospecting hours; revenue per policy, compensation ratio %, a new-business commission check and the ratio against your target follow, with totals and a per-producer average.
4. **Acquisition:** each month's marketing spend, lead costs, producer hours (or your own), policies written and commission per policy (or your own); total acquisition cost, cost per policy, lifetime commission per policy, value-to-cost ratio and payback policies follow.
5. **Commission Splits:** each policy or client with producer, premium, carrier commission %, agent split % and agency fee (or your own); gross commission, agent and agency shares, agent net after fee, the agent's effective % of premium and agency net follow, with a per-producer block.
6. **Summary:** policies on the book, retention against your target, renewal commission and commission lost, lines below target, revenue per producer and per employee, compensation ratio, acquisition cost and value-to-cost, the best line and producer, and agent and agency net from splits.

Blue text = type here. Black text = formula, leave it alone. Definitions are on the Start Here sheet.

No macros, no tracking. Arithmetic only, not financial, legal or licensing advice; the targets are your own, not industry benchmarks. Licensed for personal or single-agency use; please do not redistribute.

Questions: hello@knackdesk.com · Free tools: https://knackdesk.com
"""


def build_workbook(xlsx):
    wb = Workbook(); ws = wb.active; ws.title = "Start Here"; build_start(ws)
    build_settings(wb.create_sheet("Settings")); build_book(wb.create_sheet("Book by Line"))
    build_producers(wb.create_sheet("Producers")); build_acquisition(wb.create_sheet("Acquisition"))
    build_splits(wb.create_sheet("Commission Splits")); build_summary(wb.create_sheet("Summary"))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.save(xlsx)


def write_zip(out, xlsx):
    readme = os.path.join(out, "README.md")
    with open(readme, "w") as f: f.write(README)
    with zipfile.ZipFile(os.path.join(out, f"{SLUG}.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(xlsx, f"{SLUG}.xlsx"); z.write(readme, "README.md")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--zip-only", action="store_true", help="re-zip the existing (recalculated) xlsx with the README")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    xlsx = os.path.join(args.out, f"{SLUG}.xlsx")
    if args.zip_only:
        if not os.path.exists(xlsx): raise SystemExit(f"missing {xlsx}; build it first")
    else:
        build_workbook(xlsx)
    write_zip(args.out, xlsx)
    print(f"{'zipped' if args.zip_only else 'built'} {xlsx}")


if __name__ == "__main__":
    main()
