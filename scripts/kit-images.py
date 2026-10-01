"""Render a cover image and sheet screenshots for a kit from its recalculated workbook.
Usage: kit-images.py <slug> [--sheets N]   -> products/<slug>/assets/images/{cover.png, sheet-N-<name>.png}
Pure Pillow rendering of cell values (data_only), so the look is consistent and never paginated."""
import argparse, os, re, sys, datetime as dt
import openpyxl
from PIL import Image, ImageDraw, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONT_DIR = "/System/Library/Fonts/Supplemental"
REG = os.path.join(FONT_DIR, "Arial.ttf"); BOLD = os.path.join(FONT_DIR, "Arial Bold.ttf")
BG = (251, 250, 247); FG = (29, 29, 27); MUTED = (107, 107, 102); ACCENT = (15, 107, 92); LINE = (230, 227, 220); CARD = (255, 255, 255)
BLUE = (0, 0, 255); YELLOW = (255, 249, 196); HEAD = (231, 230, 230)
COVER_W, COVER_H = 1200, 675; SHOT_W = 1200; SCALE = 2  # render at 2x for crisp text


def font(path, size): return ImageFont.truetype(path, size)


def frontmatter(slug):
    text = open(os.path.join(ROOT, "products", slug, "PLAN.md"), encoding="utf8").read()
    fm = text.split("---")[1]
    data = {}
    for line in fm.strip().splitlines():
        if ":" in line:
            k, v = line.split(":", 1); data[k.strip()] = v.strip().strip('"')
    return data


def wrap(draw, text, fnt, max_w):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= max_w: cur = t
        else: lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines


def cover(slug, data, sheets, out):
    W, H = COVER_W * SCALE, COVER_H * SCALE
    im = Image.new("RGB", (W, H), BG); d = ImageDraw.Draw(im)
    pad = 72 * SCALE
    # brand mark
    d.rounded_rectangle([pad, pad, pad + 56 * SCALE, pad + 56 * SCALE], radius=12 * SCALE, fill=ACCENT)
    d.text((pad + 14 * SCALE, pad + 6 * SCALE), "K", font=font(BOLD, 40 * SCALE), fill=(255, 255, 255))
    d.text((pad + 72 * SCALE, pad + 12 * SCALE), "Knackdesk", font=font(BOLD, 30 * SCALE), fill=FG)
    d.text((W - pad - d.textlength("Spreadsheet kit", font=font(REG, 26 * SCALE)), pad + 14 * SCALE), "Spreadsheet kit", font=font(REG, 26 * SCALE), fill=MUTED)
    # title + tagline
    y = pad + 104 * SCALE
    tf = font(BOLD, 60 * SCALE)
    for line in wrap(d, data["name"], tf, W - 2 * pad)[:2]:
        d.text((pad, y), line, font=tf, fill=FG); y += 70 * SCALE
    y += 10 * SCALE
    gf = font(REG, 30 * SCALE)
    for line in wrap(d, data.get("tagline", ""), gf, W - 2 * pad)[:3]:
        d.text((pad, y), line, font=gf, fill=MUTED); y += 40 * SCALE
    # sheet pills
    y += 22 * SCALE; x = pad; pf = font(REG, 24 * SCALE)
    for s in sheets:
        tw = d.textlength(s, font=pf); bw = tw + 36 * SCALE
        if x + bw > W - pad: x = pad; y += 60 * SCALE
        d.rounded_rectangle([x, y, x + bw, y + 46 * SCALE], radius=23 * SCALE, fill=CARD, outline=LINE, width=2 * SCALE)
        d.text((x + 18 * SCALE, y + 9 * SCALE), s, font=pf, fill=FG); x += bw + 14 * SCALE
    # footer
    price = f"${int(data['price_cents']) // 100} one-time"
    d.text((pad, H - pad - 40 * SCALE), price, font=font(BOLD, 34 * SCALE), fill=ACCENT)
    note = "Excel · Google Sheets · Numbers · formulas only, no macros"
    nf = font(REG, 24 * SCALE)
    d.text((W - pad - d.textlength(note, font=nf), H - pad - 32 * SCALE), note, font=nf, fill=MUTED)
    im.resize((COVER_W, COVER_H), Image.LANCZOS).save(out, optimize=True)


def fmt_value(cell):
    v = cell.value
    if v is None or v == "": return ""
    nf = cell.number_format or ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%b %Y") if "mmm" in nf and "d" not in nf.replace("ddd", "") else v.strftime("%Y-%m-%d")
    if isinstance(v, bool): return "Yes" if v else "No"
    if isinstance(v, (int, float)):
        if '"%"' in nf or nf.endswith("%"): return f"{v:,.2f}%"
        if "0.00" in nf or "#,##0.00" in nf: return f"{v:,.2f}"
        if "0.0" in nf: return f"{v:,.1f}"
        if nf in ("0", "#,##0"): return f"{v:,.0f}"
        return f"{v:,.2f}" if isinstance(v, float) and v != int(v) else f"{v:,.0f}"
    return str(v)


def is_bold(cell): return bool(cell.font and cell.font.bold)
def is_blue(cell): return bool(cell.font and cell.font.color and getattr(cell.font.color, "rgb", None) in ("000000FF", "FF0000FF"))
def fill_of(cell):
    fc = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
    if fc in ("00FFFF00", "FFFFFF00"): return YELLOW
    if fc in ("00E7E6E6", "FFE7E6E6"): return HEAD
    return None


def sheet_shot(ws, title, out, max_rows=26, max_cols=9):
    # determine used range
    # trim to the used block: last row/column with a visible value, plus two blank rows for context
    used_r = max((c.row for row in ws.iter_rows(max_row=min(ws.max_row, 80)) for c in row if c.value not in (None, "")), default=1)
    used_c = max((c.column for row in ws.iter_rows(max_row=min(ws.max_row, 80)) for c in row if c.value not in (None, "")), default=1)
    rows = min(used_r + 2, max_rows); cols = min(used_c, max_cols)
    widths = []
    for c in range(1, cols + 1):
        letter = openpyxl.utils.get_column_letter(c)
        w = ws.column_dimensions[letter].width if letter in ws.column_dimensions and ws.column_dimensions[letter].width else 10
        widths.append(max(7, min(w, 34)))
    total_w = sum(widths)
    cell_px = (SHOT_W - 2 * 24) * SCALE / total_w
    col_px = [int(w * cell_px) for w in widths]
    row_h = 30 * SCALE; top = 70 * SCALE; left = 24 * SCALE
    H = top + rows * row_h + 24 * SCALE
    im = Image.new("RGB", (SHOT_W * SCALE, H), BG); d = ImageDraw.Draw(im)
    d.rectangle([0, 0, SHOT_W * SCALE, 48 * SCALE], fill=ACCENT)
    d.text((left, 10 * SCALE), f"{title}  ·  sheet “{ws.title}”", font=font(BOLD, 22 * SCALE), fill=(255, 255, 255))
    rf = font(REG, 16 * SCALE); bf = font(BOLD, 16 * SCALE)
    x0 = left
    for r in range(1, rows + 1):
        y = top + (r - 1) * row_h; x = x0
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c); w = col_px[c - 1]
            fill = fill_of(cell)
            if fill: d.rectangle([x, y, x + w, y + row_h], fill=fill)
            d.rectangle([x, y, x + w, y + row_h], outline=LINE, width=1)
            text = fmt_value(cell)
            if text:
                fnt = bf if is_bold(cell) else rf
                color = BLUE if is_blue(cell) and not is_bold(cell) else FG
                if cell.font and cell.font.italic: color = MUTED
                maxw = w - 12 * SCALE
                while d.textlength(text, font=fnt) > maxw and len(text) > 3: text = text[:-2] + "…"
                tx = x + 6 * SCALE
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool): tx = x + w - 6 * SCALE - d.textlength(text, font=fnt)
                d.text((tx, y + 6 * SCALE), text, font=fnt, fill=color)
            x += w
    im.resize((SHOT_W, H // SCALE), Image.LANCZOS).save(out, optimize=True)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("slug"); ap.add_argument("--sheets", type=int, default=3); args = ap.parse_args()
    data = frontmatter(args.slug)
    assets = os.path.join(ROOT, "products", args.slug, "assets")
    xlsx = os.path.join(assets, f"{args.slug}.xlsx")
    outdir = os.path.join(assets, "images"); os.makedirs(outdir, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    sheets = [s for s in wb.sheetnames if s != "Start Here"]
    cover(args.slug, data, sheets, os.path.join(outdir, "cover.png"))
    picked = [s for s in sheets if s != "Settings"][: args.sheets]
    for i, name in enumerate(picked, start=1):
        safe = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
        sheet_shot(wb[name], data["name"], os.path.join(outdir, f"sheet-{i}-{safe}.png"))
    print("wrote", outdir, ["cover.png"] + [f"sheet-{i}-{re.sub(r'[^a-z0-9]+', '-', n.lower()).strip('-')}.png" for i, n in enumerate(picked, start=1)])


if __name__ == "__main__":
    main()
